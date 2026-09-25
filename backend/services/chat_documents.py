"""Typed, local Word artifacts. Model output is data, never executable code."""
from __future__ import annotations

import asyncio
import hashlib
import io
import json
import math
import re
import shutil
import threading
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Annotated, Literal

from pydantic import BaseModel, ConfigDict, Field, ValidationError
from starlette.concurrency import run_in_threadpool

from config import settings
from services.app_logging import get_logger
from services import session_store, image_vault

ROOT = settings.data_dir / "artifacts"
logger = get_logger("backend.chat_documents")
MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class StrictModel(BaseModel):
    model_config = ConfigDict(extra="forbid")


class TextBlock(StrictModel):
    type: Literal["paragraph", "heading", "bullet", "numbered"]
    text: str = Field(min_length=1, max_length=12000)


class TableBlock(StrictModel):
    type: Literal["table"]
    headers: list[str] = Field(min_length=1, max_length=8)
    rows: list[list[str]] = Field(min_length=1, max_length=60)


class ImageBlock(StrictModel):
    type: Literal["image"]
    image_id: str = Field(max_length=40)
    caption: str = Field(default="", max_length=1000)


class FunctionPlot(StrictModel):
    type: Literal["function_plot"]
    function: Literal["sin", "cos"]
    caption: str = Field(default="", max_length=1000)
    x_min: float = Field(default=0, ge=-10000, le=10000, allow_inf_nan=False)
    x_max: float = Field(default=6.283185307179586, ge=-10000, le=10000, allow_inf_nan=False)


class DocumentSpec(StrictModel):
    title: str = Field(min_length=1, max_length=160)
    filename: str = Field(default="document.docx", min_length=1, max_length=100)
    blocks: list[Annotated[TextBlock | TableBlock | ImageBlock | FunctionPlot, Field(discriminator="type")]] = Field(min_length=1, max_length=80)


def wants_document(text: str) -> bool:
    """Explicit creation requests only; quoted snippets and how-to questions stay chat."""
    text = re.sub(r"```[\s\S]*?```", "", text).strip().lower()
    if text.startswith("/docx"): return True
    if re.search(r"\b(how (do|can|should|to)|don't|do not|explain|python (code|script)|code (to|for))\b", text): return False
    return bool(re.search(r"(?:\.docx\b|\bdocx\b|\bword (?:document|file)\b)", text)
                and re.search(r"\b(make|create|generate|export|save|convert|turn|put|give|prepare|produce|write|download)\b", text))


def image_inventory(session_id: str) -> dict:
    if not re.fullmatch(r"[A-Za-z0-9_-]{1,128}", session_id or ""):
        raise ValueError("Open or save a chat before creating a document.")
    session = session_store.get_session(session_id)
    if not session: raise ValueError("Open or save a chat before creating a document.")
    result = {}
    for message in session.get("messages", []):
        for record in session_store._message_image_records(message):
            # Do not expose locked images to the model or copy them into a document.
            try:
                raw, _ = session_store._decode_image(record)
                image_vault.require_public(hashlib.sha256(raw).hexdigest())
            except (ValueError, OSError):
                continue
            result[f"image-{len(result) + 1}"] = {"message_id": message["id"], "image_id": record["id"],
                                                     "name": str(record.get("name") or "Chat image")[:200]}
    # Recent images remain available without expanding an unbounded prompt.
    return dict(list(result.items())[-20:])


def document_instruction(inventory: dict) -> str:
    images = [{"image_id": key, "name": value["name"]} for key, value in inventory.items()]
    return (
        "The application will create a real downloadable Word document and show it in chat. "
        "Return ONLY a JSON document matching this schema. Write the actual requested document content, "
        "not Python code, setup steps, or instructions for making a file. Use the conversation as source material. "
        "Do not claim you cannot create files. Use plain text within each block, without Markdown syntax. "
        "Include a concise title and a descriptive .docx filename. Preserve requested facts and do not invent missing ones. "
        "For pictures use ONLY image_id values from the available chat images below. "
        "A filename mentioned in prose or a code snippet does not mean that file exists. "
        "If a requested image is unavailable, include an honest paragraph saying it must be attached. "
        "For a mathematical sine or cosine plot, a function_plot block creates an actual chart locally; "
        "use sin or cos and the requested x range (default 0 to 6.283185307179586). "
        "Do not invent plots for other functions. Keep the document concise enough to complete.\n"
        f"Available chat images: {json.dumps(images)}\nSchema: {json.dumps(DocumentSpec.model_json_schema())}"
    )


def _artifact_dir(artifact_id: str) -> Path:
    if not re.fullmatch(r"[a-f0-9]{32}", artifact_id): raise FileNotFoundError("Document not found")
    root = ROOT.resolve()
    target = ROOT / artifact_id
    if target.resolve().parent != root or target.is_symlink(): raise FileNotFoundError("Document not found")
    return target


def read_artifact(artifact_id: str) -> dict:
    directory = _artifact_dir(artifact_id)
    try: value = json.loads((directory / "document.json").read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc: raise FileNotFoundError("This document is unavailable. Try creating it again.") from exc
    if not session_store.get_session(value["session_id"]): raise FileNotFoundError("The source chat is deleted. Restore it to view this document.")
    for digest in value.get("source_hashes", []): image_vault.require_public(digest)
    return value


def file_path(artifact_id: str, name: str) -> Path:
    read_artifact(artifact_id)
    if name != "document.docx" and not re.fullmatch(r"image-[0-9]+\.png", name): raise FileNotFoundError("Document file not found")
    directory = _artifact_dir(artifact_id)
    path = directory / name
    if not path.is_file() or path.resolve().parent != directory.resolve(): raise FileNotFoundError("Document file is missing")
    return path


def _plot(block: FunctionPlot) -> bytes:
    """Small deterministic mathematical plot; no eval, generated scripts, or subprocess."""
    from PIL import Image, ImageDraw, ImageFont
    if block.x_max <= block.x_min: raise ValueError("The plot's maximum x must be greater than its minimum.")
    image = Image.new("RGB", (1200, 680), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=22)
    left, top, width, height = 105, 70, 1045, 500
    for i in range(5):
        y = top + i * height / 4
        draw.line((left, y, left+width, y), fill="#dddddd", width=1)
        draw.text((30, y-12), f"{1 - i/2:g}", fill="#222222", font=font)
    for i in range(6):
        x = left + width * i / 5
        draw.line((x, top, x, top+height), fill="#eeeeee", width=1)
        draw.text((x-18, top+height+12), f"{block.x_min+(block.x_max-block.x_min)*i/5:.2f}", fill="#222222", font=font)
    draw.line((left, top+height/2, left+width, top+height/2), fill="#777777", width=2)
    draw.rectangle((left, top, left+width, top+height), outline="#333333", width=2)
    function = math.sin if block.function == "sin" else math.cos
    points = [(left + width*i/1000, top+height*(1-function(block.x_min+(block.x_max-block.x_min)*i/1000))/2) for i in range(1001)]
    draw.line(points, fill="#245b9b", width=4)
    draw.text((left, 20), f"y = {block.function}(x)", fill="#111111", font=font)
    draw.text((570, 630), "x (radians)", fill="#222222", font=font)
    output = io.BytesIO(); image.save(output, format="PNG"); return output.getvalue()


def create_document(spec: DocumentSpec, session_id: str, inventory: dict, cancel: threading.Event) -> dict:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from PIL import Image

    if not re.fullmatch(r"[A-Za-z0-9_-]{1,128}", session_id or "") or not session_store.get_session(session_id):
        raise ValueError("The source chat no longer exists.")
    artifact_id = uuid.uuid4().hex
    ROOT.mkdir(parents=True, exist_ok=True)
    temporary = ROOT / (".partial-" + artifact_id)
    temporary.mkdir()
    try:
        doc = Document()
        section = doc.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(.8)
        section.left_margin = section.right_margin = Inches(.85)
        normal = doc.styles["Normal"]
        normal.font.name, normal.font.size = "Calibri", Pt(11)
        normal.paragraph_format.space_after = Pt(8)
        for style in ("Title", "Heading 1", "Heading 2"):
            doc.styles[style].font.color.rgb = RGBColor(0, 0, 0)
        doc.add_paragraph(spec.title, "Title")
        blocks, source_hashes = [], []
        for index, block in enumerate(spec.blocks):
            if cancel.is_set(): raise InterruptedError("Document creation cancelled")
            if index == 0 and isinstance(block, TextBlock) and block.type == "heading" and block.text.strip().casefold() == spec.title.strip().casefold():
                continue
            item = block.model_dump()
            if isinstance(block, TextBlock):
                style = {"heading": "Heading 1", "bullet": "List Bullet", "numbered": "List Number"}.get(block.type)
                doc.add_paragraph(block.text, style)
            elif isinstance(block, TableBlock):
                if any(len(row) != len(block.headers) for row in block.rows): raise ValueError("A document table has unequal column counts. Try again with a simpler table.")
                if any(len(cell) > 3000 for row in [block.headers, *block.rows] for cell in row): raise ValueError("A table cell is too long.")
                table = doc.add_table(rows=1, cols=len(block.headers)); table.style = "Table Grid"
                for cell, text in zip(table.rows[0].cells, block.headers):
                    cell.text = text
                    for run in cell.paragraphs[0].runs: run.bold = True
                    shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "E8EEF5"); cell._tc.get_or_add_tcPr().append(shade)
                for row in block.rows:
                    for cell, text in zip(table.add_row().cells, row): cell.text = text
                doc.add_paragraph()
            else:
                if isinstance(block, ImageBlock):
                    source = inventory.get(block.image_id)
                    if not source: raise ValueError("The requested image is not attached to this chat. Attach it and try again.")
                    found = session_store.get_session_image_by_id(session_id, source["message_id"], source["image_id"])
                    if not found: raise ValueError("A selected chat image is missing. Attach it again before creating the document.")
                    raw = found[0]; digest = hashlib.sha256(raw).hexdigest()
                    image_vault.require_public(digest); source_hashes.append(digest)
                    if len(raw) > 25*1024*1024: raise ValueError("The selected image is too large for this document.")
                    with Image.open(io.BytesIO(raw)) as image:
                        if image.width * image.height > 40_000_000: raise ValueError("The selected image has too many pixels.")
                        image.thumbnail((2000, 2000))
                        output = io.BytesIO(); image.convert("RGB").save(output, format="PNG"); raw = output.getvalue()
                else:
                    raw = _plot(block)
                source_hashes.append(hashlib.sha256(raw).hexdigest())
                name = f"image-{index}.png"
                (temporary / name).write_bytes(raw)
                with Image.open(io.BytesIO(raw)) as image:
                    width = min(6.6, 6.5 * image.width / image.height)
                doc.add_picture(io.BytesIO(raw), width=Inches(width))
                if block.caption: doc.add_paragraph(block.caption, "Caption")
                item = {"type": "image", "image_file": name, "caption": block.caption}
            blocks.append(item)
        filename = re.sub(r"[^a-zA-Z0-9 _-]", "", Path(spec.filename.replace("\\", "/")).stem)[:70].strip(" .") or "document"
        if re.fullmatch(r"(?i)(con|prn|aux|nul|com[1-9]|lpt[1-9])", filename): filename = "document-" + filename
        filename += ".docx"
        doc.save(temporary / "document.docx")
        descriptor = {"id": artifact_id, "kind": "docx", "name": filename, "title": spec.title,
                      "size": (temporary / "document.docx").stat().st_size}
        metadata = {**descriptor, "session_id": session_id, "blocks": blocks,
                    "source_hashes": source_hashes, "created_at": datetime.now(timezone.utc).isoformat()}
        (temporary / "document.json").write_text(json.dumps(metadata, ensure_ascii=False), encoding="utf-8")
        for digest in source_hashes: image_vault.require_public(digest)
        if cancel.is_set(): raise InterruptedError("Document creation cancelled")
        temporary.rename(_artifact_dir(artifact_id))
        logger.info("Created Word artifact %s for session %s (%d bytes)", artifact_id, session_id, descriptor["size"])
        return descriptor
    finally:
        # Only our newly-created partial directory can be removed here.
        if temporary.exists() and temporary.resolve().parent == ROOT.resolve() and temporary.name == ".partial-" + artifact_id:
            shutil.rmtree(temporary)


async def stream_document(client, payload, request, client_request):
    """Use the existing chat queue/cancellation task; publish only complete files."""
    cancel = threading.Event()
    worker = None
    def event(**value): return f"data: {json.dumps(value)}\n\n"
    try:
        inventory = await run_in_threadpool(image_inventory, request.session_id or "")
        context_notes = "\n\n".join(m["content"] for m in payload["messages"] if m["role"] == "system")
        payload = {**payload, "format": DocumentSpec.model_json_schema(),
                   "messages": [{"role": "system", "content": context_notes + "\n\n" + document_instruction(inventory)},
                                *[{"role": m["role"], "content": m["content"]} for m in payload["messages"] if m["role"] != "system"]],
                   "options": {**payload.get("options", {}), "temperature": 0, "num_predict": 4096}}
        yield event(document_status="Drafting your Word document…")
        content, completed = "", False
        async with client.stream("POST", f"{settings.ollama_base_url}/api/chat", json=payload) as response:
            if response.is_error:
                detail = (await response.aread()).decode("utf-8", errors="replace")[:1500]
                raise ValueError(f"The model could not draft the document ({response.status_code}): {detail}")
            async for line in response.aiter_lines():
                if await client_request.is_disconnected(): raise asyncio.CancelledError()
                if not line: continue
                chunk = json.loads(line)
                if chunk.get("error"): raise ValueError(str(chunk["error"]))
                content += chunk.get("message", {}).get("content", "")
                if len(content) > 180000: raise ValueError("This document is too large. Try creating it in smaller sections.")
                yield ": drafting document\n\n"
                if chunk.get("done"):
                    if chunk.get("done_reason") == "length": raise ValueError("The document draft reached the model's output limit. Ask for a shorter document and retry.")
                    completed = True; break
        if not completed: raise ValueError("The model stopped before completing the document. Try again.")
        try: spec = DocumentSpec.model_validate_json(content)
        except ValidationError as exc: raise ValueError("The model did not return a usable document. Try again or use a model that supports structured output.") from exc
        yield event(document_status="Creating the Word file and preview…")
        worker = asyncio.create_task(run_in_threadpool(create_document, spec, request.session_id, inventory, cancel))
        artifact = await asyncio.shield(worker)
        if await client_request.is_disconnected(): raise asyncio.CancelledError()
        summary = f"Created **{artifact['name']}**. Open the document below to view it or download the Word file."
        context_parts = [spec.title]
        for block in spec.blocks:
            if isinstance(block, TextBlock): context_parts.append(block.text)
            elif isinstance(block, TableBlock): context_parts.extend(" | ".join(row) for row in [block.headers, *block.rows])
            else: context_parts.append(block.caption)
        message = {"id": request.reply_message_id or uuid.uuid4().hex, "role": "assistant", "content": summary,
                   "artifacts": [artifact], "document_text": "\n".join(context_parts)}
        saved = await run_in_threadpool(session_store.append_messages, request.session_id, [message], model=request.model)
        if saved is None: raise ValueError("The source chat was deleted before the document could be attached.")
        yield event(token=summary, artifacts=[artifact], document_text=message["document_text"], done=True)
    except asyncio.CancelledError:
        cancel.set()
        if worker:
            try: await asyncio.shield(worker)
            except Exception: pass
        raise
    except Exception as exc:
        logger.exception("Word document creation failed for session %s", request.session_id)
        detail = str(exc) if isinstance(exc, (ValueError, ImportError, OSError)) else "Document creation failed. See the app logs and try again."
        if isinstance(exc, ImportError): detail = "Word document tools are unavailable. Run the app setup again to install python-docx and Pillow, then retry."
        yield event(token=f"[Error: {detail}]", error=detail, done=True)
