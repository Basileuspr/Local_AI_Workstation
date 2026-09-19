"""
File Parser Service
Reads text from multiple file formats.

This is the shared foundation for both:
- Direct file upload (full text into chat context)
- Knowledge base (text gets chunked and embedded for RAG)

Supported formats: .txt, .md, .pdf, .docx
"""

import base64
import io
import re
from pathlib import Path

import httpx

from config import settings
from services.gpu_coordination import gpu_coordinator


def parse_file(file_bytes: bytes, filename: str) -> dict:
    """
    Extract text from a file based on its extension.
    
    Args:
        file_bytes: Raw file content as bytes
        filename: Original filename (used to detect format)
    
    Returns:
        dict with:
            - text: Extracted text content
            - filename: Original filename
            - format: Detected format
            - char_count: Length of extracted text
            - error: Error message if extraction failed, None otherwise
    """
    ext = Path(filename).suffix.lower()

    try:
        if ext in (".txt", ".md"):
            text = _parse_text(file_bytes)
        elif ext == ".pdf":
            text, pdf_metadata = _parse_pdf(file_bytes)
        elif ext in (".docx", ".doc"):
            text = _parse_docx(file_bytes)
        else:
            return {
                "text": "",
                "filename": filename,
                "format": ext,
                "char_count": 0,
                "error": f"Unsupported file format: {ext}",
            }

        result = {
            "text": text,
            "filename": filename,
            "format": ext,
            "char_count": len(text),
            "error": None,
        }
        if ext == ".pdf":
            result.update(pdf_metadata)
        return result

    except Exception as e:
        return {
            "text": "",
            "filename": filename,
            "format": ext,
            "char_count": 0,
            "error": f"Failed to parse {filename}: {str(e)}",
        }


def _parse_text(file_bytes: bytes) -> str:
    """Plain text and markdown — just decode the bytes."""
    # Try UTF-8 first, fall back to latin-1 which never fails
    try:
        return file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        return file_bytes.decode("latin-1")


def _extract_pdf_page_texts(file_bytes: bytes) -> list[str]:
    """Return one native text-layer value per PDF page."""
    from PyPDF2 import PdfReader

    reader = PdfReader(io.BytesIO(file_bytes))
    return [(page.extract_text() or "").strip() for page in reader.pages]


def _render_pdf_page_png(file_bytes: bytes, page_index: int) -> bytes:
    """Rasterize one PDF page entirely in memory for local vision OCR."""
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(file_bytes)
    page = document[page_index]
    bitmap = None
    try:
        # 3x PDF scale is approximately 216 DPI: legible without creating the
        # very large bitmaps that can make a long scanned document unstable.
        bitmap = page.render(scale=3)
        image = bitmap.to_pil()
        output = io.BytesIO()
        image.save(output, format="PNG")
        return output.getvalue()
    finally:
        if bitmap is not None:
            bitmap.close()
        page.close()
        document.close()


def _strip_ocr_fence(text: str) -> str:
    text = text.strip()
    match = re.fullmatch(r"```(?:text)?\s*\n?(.*?)\n?```", text, flags=re.DOTALL | re.IGNORECASE)
    return match.group(1).strip() if match else text


def _transcribe_page_png(page_png: bytes, page_number: int, keep_alive: int | str) -> str:
    """Transcribe a page with the configured local Ollama vision model."""
    encoded = base64.b64encode(page_png).decode("ascii")
    response = httpx.post(
        f"{settings.ollama_base_url}/api/chat",
        json={
            "model": settings.ocr_model,
            "stream": False,
            "think": False,
            "keep_alive": keep_alive,
            "options": {"temperature": 0},
            "messages": [{
                "role": "user",
                "content": (
                    f"Transcribe page {page_number} exactly as visible. Preserve reading order, "
                    "paragraph breaks, headings, lists, and table rows. Do not summarize, explain, "
                    "or wrap the answer in Markdown fences. If a portion is unreadable, write [unreadable]."
                ),
                "images": [encoded],
            }],
        },
        timeout=settings.ocr_timeout_seconds,
    )
    response.raise_for_status()
    message = response.json().get("message") or {}
    text = _strip_ocr_fence(message.get("content") or "")
    if not text:
        raise RuntimeError(
            f"The local OCR model returned no visible transcription for page {page_number}."
        )
    return text


def _ocr_pdf_pages(file_bytes: bytes, page_indices: list[int]) -> dict[int, str]:
    """OCR selected zero-based pages while holding the shared GPU lease."""
    if not page_indices:
        return {}

    lease_owner = "pdf-ocr"
    if not gpu_coordinator.acquire(lease_owner):
        owner = gpu_coordinator.current_owner() or "another local task"
        raise RuntimeError(f"PDF OCR is waiting because {owner} owns the GPU. Stop or reset it, then retry.")

    try:
        # An idle Diffusers pipeline can still retain CPU/GPU allocations.
        from services.image_generation import manager as image_manager

        image_manager.unload_for_training()
        transcriptions: dict[int, str] = {}
        for position, page_index in enumerate(page_indices):
            page_png = _render_pdf_page_png(file_bytes, page_index)
            keep_alive = 0 if position == len(page_indices) - 1 else "5m"
            transcriptions[page_index] = _transcribe_page_png(
                page_png,
                page_number=page_index + 1,
                keep_alive=keep_alive,
            )
        return transcriptions
    finally:
        gpu_coordinator.release(lease_owner)


def _parse_pdf(file_bytes: bytes) -> tuple[str, dict]:
    """
    Extract text from PDF, using local vision OCR only where the native text
    layer is empty or too small to be useful.
    """
    native_pages = _extract_pdf_page_texts(file_bytes)
    ocr_indices = [
        index for index, text in enumerate(native_pages)
        if len(text.strip()) < settings.ocr_min_page_chars
    ]
    ocr_pages = _ocr_pdf_pages(file_bytes, ocr_indices) if ocr_indices else {}

    pages = []
    for index, native_text in enumerate(native_pages):
        text = ocr_pages.get(index, native_text).strip()
        if text:
            pages.append(f"--- Page {index + 1} ---\n{text}")

    if not pages:
        raise RuntimeError("No text could be extracted or transcribed from this PDF.")

    return "\n\n".join(pages), {
        "page_count": len(native_pages),
        "text_layer_pages": [
            index + 1 for index, text in enumerate(native_pages)
            if len(text.strip()) >= settings.ocr_min_page_chars
        ],
        "ocr_pages": [index + 1 for index in ocr_indices],
        "ocr_model": settings.ocr_model if ocr_indices else None,
    }


def _parse_docx(file_bytes: bytes) -> str:
    """Extract text from Word documents."""
    from docx import Document

    doc = Document(io.BytesIO(file_bytes))
    paragraphs = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)

    if not paragraphs:
        return "[No text found in this document.]"

    return "\n\n".join(paragraphs)
