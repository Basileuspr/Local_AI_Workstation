"""Bounded rectangular table grids, including merged cells, for DOCX conversion."""
import re

from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

TABLE_ALIGN = {'left': WD_TABLE_ALIGNMENT.LEFT, 'center': WD_TABLE_ALIGNMENT.CENTER, 'right': WD_TABLE_ALIGNMENT.RIGHT}
CELL_ALIGN = {'top': WD_CELL_VERTICAL_ALIGNMENT.TOP, 'center': WD_CELL_VERTICAL_ALIGNMENT.CENTER, 'bottom': WD_CELL_VERTICAL_ALIGNMENT.BOTTOM}


def _integer(value, low, high):
    return isinstance(value, int) and not isinstance(value, bool) and low <= value <= high


def table_grid(node):
    rows = node.get('content', [])
    if not isinstance(rows, list) or not 1 <= len(rows) <= 100:
        raise ValueError('Tables support 1–100 rows and 1–12 grid columns.')
    grid, anchors, widths, columns = [[None] * 12 for _ in rows], [], [0] * 12, 0
    attrs = node.get('attrs') or {}
    if not isinstance(attrs, dict) or attrs.get('tableAlignment', 'left') not in tuple(TABLE_ALIGN):
        raise ValueError('Invalid table alignment.')
    if attrs.get('borderPreset', 'grid') not in ('grid', 'outer', 'none') or not isinstance(attrs.get('repeatHeader', False), bool):
        raise ValueError('Invalid table border or header setting.')
    for r, row in enumerate(rows):
        if not isinstance(row, dict) or row.get('type') != 'tableRow' or not isinstance(row.get('content', []), list):
            raise ValueError('Invalid table row.')
        c = 0
        for cell in row.get('content', []):
            while c < 12 and grid[r][c] is not None: c += 1
            if not isinstance(cell, dict) or cell.get('type') not in ('tableCell', 'tableHeader'):
                raise ValueError('Invalid table cell.')
            ca = cell.get('attrs') or {}
            if not isinstance(ca, dict): raise ValueError('Invalid cell attributes.')
            colspan, rowspan = ca.get('colspan', 1), ca.get('rowspan', 1)
            if not _integer(colspan, 1, 12) or not _integer(rowspan, 1, 100) or c + colspan > 12 or r + rowspan > len(rows):
                raise ValueError('Cell spans must fit inside the 100-row / 12-column table grid.')
            if ca.get('verticalAlign', 'top') not in tuple(CELL_ALIGN):
                raise ValueError('Invalid cell vertical alignment.')
            fill = ca.get('backgroundColor')
            if fill is not None and (not isinstance(fill, str) or not re.fullmatch(r'#[0-9a-fA-F]{6}', fill)):
                raise ValueError('Cell shading must be a six-digit hexadecimal color.')
            colwidth = ca.get('colwidth')
            if colwidth is not None:
                if not isinstance(colwidth, list) or len(colwidth) != colspan or any(not _integer(w, 0, 1600) or 0 < w < 24 for w in colwidth):
                    raise ValueError('Column widths must be whole pixels from 24 to 1,600 (or 0 for automatic).')
                for offset, width in enumerate(colwidth):
                    if width and widths[c + offset] and widths[c + offset] != width:
                        raise ValueError('Column widths must agree across all cells in a column.')
                    if width: widths[c + offset] = width
            anchor = {'row': r, 'column': c, 'rowspan': rowspan, 'colspan': colspan, 'node': cell}
            for y in range(r, r + rowspan):
                for x in range(c, c + colspan):
                    if grid[y][x] is not None: raise ValueError('Merged table cells overlap.')
                    grid[y][x] = anchor
            anchors.append(anchor)
            c += colspan
            columns = max(columns, c)
    if not columns or any(cell is None for row in grid for cell in row[:columns]):
        raise ValueError('Table cells must cover a rectangular grid without gaps.')
    if attrs.get('repeatHeader') and any(anchor['row'] == 0 and anchor['rowspan'] > 1 for anchor in anchors):
        raise ValueError('Split cells spanning below the first row before repeating the header row.')
    return {'rows': len(rows), 'columns': columns, 'anchors': anchors, 'widths': widths[:columns]}


def _element(tag, **attrs):
    node = OxmlElement('w:' + tag)
    for key, value in attrs.items(): node.set(qn('w:' + key), str(value))
    return node


def write_table(parent, node, available, write):
    info, attrs = table_grid(node), node.get('attrs') or {}
    table = parent.add_table(rows=info['rows'], cols=info['columns'])
    table.style = 'Table Grid'
    table.autofit = False
    table.alignment = TABLE_ALIGN[attrs.get('tableAlignment', 'left')]
    widths = [width or available * 96 / info['columns'] for width in info['widths']]
    scale = min(1, available * 96 / sum(widths))
    widths = [width * scale / 96 for width in widths]
    for column, width in zip(table.columns, widths): column.width = Inches(width)
    for row in table.rows:
        for c, cell in enumerate(row.cells): cell.width = Inches(widths[c])
    props = table._tbl.tblPr
    borders = _element('tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        visible = attrs.get('borderPreset', 'grid') != 'none' and (not edge.startswith('inside') or attrs.get('borderPreset', 'grid') == 'grid')
        borders.append(_element(edge, val='single' if visible else 'nil', sz=4, color='8994A4'))
    props.insert_element_before(borders, 'w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook', 'w:tblCaption', 'w:tblDescription', 'w:tblPrChange')
    margins = _element('tblCellMar')
    for edge in ('top', 'left', 'bottom', 'right'): margins.append(_element(edge, w=90, type='dxa'))
    props.insert_element_before(margins, 'w:tblLook', 'w:tblCaption', 'w:tblDescription', 'w:tblPrChange')
    if attrs.get('repeatHeader'):
        table.rows[0]._tr.get_or_add_trPr().append(_element('tblHeader', val=1))
    for anchor in info['anchors']:
        r, c, rs, cs = (anchor[key] for key in ('row', 'column', 'rowspan', 'colspan'))
        cell_node, cell = anchor['node'], table.cell(r, c)
        if rs > 1 or cs > 1: cell = cell.merge(table.cell(r + rs - 1, c + cs - 1))
        ca = cell_node.get('attrs') or {}
        cell.vertical_alignment = CELL_ALIGN[ca.get('verticalAlign', 'top')]
        cell._tc.clear_content()
        write(cell, cell_node.get('content') or [{'type': 'paragraph'}], sum(widths[c:c + cs]))
        if cell._tc[-1].tag != qn('w:p'): cell.add_paragraph()
        is_header = cell_node['type'] == 'tableHeader'
        if is_header:
            cell._tc.get_or_add_tcPr().insert(0, _element('cnfStyle', firstRow=1))
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if run.bold is None: run.bold = True
        fill = ca.get('backgroundColor') or ('#E8EEF7' if is_header else None)
        if fill: cell._tc.get_or_add_tcPr().insert_element_before(_element('shd', val='clear', fill=fill[1:].upper()), 'w:noWrap', 'w:tcMar', 'w:textDirection', 'w:tcFitText', 'w:vAlign', 'w:hideMark', 'w:tcPrChange')


def read_table(table, blocks, warn):
    rows, columns = list(table.rows), len(table.columns)
    if not 1 <= len(rows) <= 100 or not 1 <= columns <= 12:
        raise ValueError('A table exceeds the 100-row / 12-column editing limit or has no layout grid.')
    if table._tbl.xpath('.//w:hMerge'):
        raise ValueError('Legacy horizontal table merges are unsupported. Resave a copy with standard grid-span merges before importing.')
    widths = []
    for column in table.columns:
        pixels = round(column.width.inches * 96) if column.width is not None else 0
        widths.append(min(1600, max(24, pixels)) if pixels else 0)
        if pixels and pixels != widths[-1]: warn('Some table column widths were adjusted to the supported range.')
    if table.autofit:
        warn('Table column widths use the source layout grid. Automatic content-based sizing is converted to fixed column widths.')
    positions, matrix = {}, []
    for r, row in enumerate(rows):
        before, after = row.grid_cols_before, row.grid_cols_after
        spans = [cell.grid_span for cell in row._tr.tc_lst]
        if not _integer(before, 0, 12) or not _integer(after, 0, 12) or len(spans) > 12 or any(not _integer(span, 1, 12) for span in spans) or before + sum(spans) + after != columns:
            raise ValueError('A table contains invalid or oversized grid spans.')
        if any(cell.vMerge == 'continue' and cell.xpath('.//w:t') for cell in row._tr.tc_lst):
            warn('Hidden content in a vertical-merge continuation cell was omitted; the merged origin cell’s content is retained.')
        cells = list(row.cells)
        if before + len(cells) + after != columns:
            raise ValueError('A table has an inconsistent layout grid. The original file is unchanged.')
        if before or after: warn('Omitted cells at the start or end of a table row were filled with empty cells.')
        line = [None] * columns
        for c, cell in enumerate(cells, before):
            line[c] = cell._tc
            if cell._tc not in positions: positions[cell._tc] = {'cell': cell, 'points': []}
            positions[cell._tc]['points'].append((r, c))
        matrix.append(line)
    output_rows = [{'type': 'tableRow', 'content': []} for _ in rows]
    anchors = {}
    for entry in positions.values():
        points = entry['points']
        r, c = min(points)
        rs, cs = max(y for y, _ in points) - r + 1, max(x for _, x in points) - c + 1
        if len(points) != rs * cs or set(points) != {(y, x) for y in range(r, r + rs) for x in range(c, c + cs)}:
            raise ValueError('A merged cell is not rectangular and cannot be edited safely.')
        cell = entry['cell']
        content, flat = blocks(cell), []
        for item in content:
            if item['type'] == 'table':
                warn('Nested tables were flattened into the containing cell; nested-table editing is not supported.')
                for nested_row in item['content']:
                    for nested_cell in nested_row['content']: flat.extend(nested_cell['content'])
            elif item['type'] == 'pageBreak': warn('Manual page breaks inside table cells were omitted.')
            else: flat.append(item)
        props = cell._tc.tcPr
        shading = props.find(qn('w:shd')) if props is not None else None
        fill = shading.get(qn('w:fill')) if shading is not None else None
        background = '#' + fill.upper() if fill and re.fullmatch(r'[0-9A-Fa-f]{6}', fill) else None
        if shading is not None and background is None: warn('Theme or patterned cell shading was reduced to the default cell background.')
        conditional = props.find(qn('w:cnfStyle')) if props is not None else None
        header = conditional is not None and conditional.get(qn('w:firstRow')) in ('1', 'true', 'on')
        header = header or bool(rows[r]._tr.xpath('./w:trPr/w:tblHeader[not(@w:val="0") and not(@w:val="false") and not(@w:val="off")]'))
        anchors[r, c] = {'type': 'tableHeader' if header else 'tableCell', 'attrs': {
            'colspan': cs, 'rowspan': rs, 'colwidth': widths[c:c + cs] if any(widths[c:c + cs]) else None,
            'verticalAlign': next((key for key, value in CELL_ALIGN.items() if value == cell.vertical_alignment), 'top'),
            'backgroundColor': background,
        }, 'content': flat or [{'type': 'paragraph'}]}
    for r, line in enumerate(matrix):
        for c, cell in enumerate(line):
            if cell is None:
                output_rows[r]['content'].append({'type': 'tableCell', 'attrs': {'colwidth': [widths[c]]} if widths[c] else {}, 'content': [{'type': 'paragraph'}]})
            elif (r, c) in anchors: output_rows[r]['content'].append(anchors[r, c])
    borders = table._tbl.tblPr.find(qn('w:tblBorders'))
    preset = 'grid'
    if borders is not None:
        values = {edge: borders.find(qn('w:' + edge)) for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV')}
        visible = {edge: node is None or node.get(qn('w:val')) not in ('none', 'nil') for edge, node in values.items()}
        if not any(visible.values()): preset = 'none'
        elif all(visible[edge] for edge in ('top', 'left', 'bottom', 'right')) and not (visible['insideH'] or visible['insideV']): preset = 'outer'
        elif not all(visible.values()): warn('Custom table borders were approximated with a grid.')
        if any(node is not None and (node.get(qn('w:val')) not in ('single', 'none', 'nil') or node.get(qn('w:color'), '8994A4') != '8994A4') for node in values.values()):
            warn('Custom table border styles and colors were reduced to the editor’s border presets.')
    if table._tbl.xpath('.//w:tcBorders') or (table.style and table.style.name not in ('Table Grid', 'Normal Table')):
        warn('Custom table styles and individual cell borders were approximated; direct cell shading is retained.')
    repeat = bool(rows[0]._tr.xpath('./w:trPr/w:tblHeader[not(@w:val="0") and not(@w:val="false") and not(@w:val="off")]'))
    if repeat and any(cell['attrs']['rowspan'] > 1 for (r, _), cell in anchors.items() if r == 0):
        repeat = False
        warn('Repeating the first row was disabled because it contains a cell merged into later rows.')
    if any(row._tr.xpath('./w:trPr/w:tblHeader') for row in rows[1:]): warn('Only the first table row can repeat as a header in this editor.')
    node = {'type': 'table', 'attrs': {'tableAlignment': next((key for key, value in TABLE_ALIGN.items() if value == table.alignment), 'left'),
                                    'borderPreset': preset, 'repeatHeader': repeat}, 'content': output_rows}
    table_grid(node)
    return node
