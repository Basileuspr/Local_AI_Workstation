"""Bounded, read-only extraction. Never execute code, render images, or OCR PDFs."""
from __future__ import annotations

import ast
import hashlib
import io
import os
from pathlib import Path
import re
import stat
from datetime import datetime, timezone

TEXT_EXTENSIONS = {'.py', '.txt', '.md', '.markdown', '.rst', '.csv', '.json', '.jsonl', '.yaml', '.yml', '.toml', '.ini', '.cfg', '.log', '.js', '.jsx', '.ts', '.tsx', '.mjs', '.cjs', '.css', '.html', '.xml', '.sql', '.sh', '.ps1', '.bat', '.c', '.cpp', '.h', '.rs', '.go', '.java'}
IMAGE_EXTENSIONS = {'.jpg', '.jpeg', '.png', '.gif', '.webp', '.bmp', '.tif', '.tiff', '.ico', '.avif', '.heic', '.svg'}
EXCLUDED_DIRS = {'.git', '.hg', '.svn', 'node_modules', 'venv', '.venv', '__pycache__', '.pytest_cache', '.codex', '.claude'}


def linked(info):
    return stat.S_ISLNK(info.st_mode) or bool(getattr(info, 'st_file_attributes', 0) & 0x400)


def offline(info):
    return bool(getattr(info, 'st_file_attributes', 0) & (0x1000 | 0x40000 | 0x400000))


def signature(info):
    # Python/Windows path stat and handle fstat expose different ctime meanings.
    # Identity, size and modification time agree across both interfaces.
    return info.st_dev, info.st_ino, info.st_size, info.st_mtime_ns, info.st_ctime_ns if os.name != 'nt' else 0


def root_path(value):
    path = Path(value.strip())
    if not path.is_absolute() or str(path).startswith(('\\\\', '//')):
        raise ValueError('Choose an absolute local folder path.')
    for parent in (path, *path.parents):
        if linked(parent.lstat()):
            raise ValueError('Choose the real folder instead of a link or junction.')
    if not path.is_dir() or offline(path.stat()):
        raise ValueError('Choose an available local folder.')
    return path.resolve()


def protected(name):
    value = Path(name).name.lower()
    return value == '.env' or value.startswith('.env.') or value in {'.npmrc', '.netrc', '.pypirc', 'id_rsa', 'id_ed25519'} or Path(value).suffix in {'.pem', '.key', '.p12', '.pfx'}


def read_source(path, root, max_bytes, cancel):
    before = path.lstat()
    if linked(before) or offline(before) or not stat.S_ISREG(before.st_mode) or not path.resolve().is_relative_to(root):
        raise ValueError('Linked, offline, external or non-regular file; content not opened.')
    if before.st_size > max_bytes:
        raise ValueError('File exceeds the configured byte limit; content not read.')
    flags = os.O_RDONLY | getattr(os, 'O_BINARY', 0) | getattr(os, 'O_NOFOLLOW', 0) | getattr(os, 'O_NONBLOCK', 0)
    chunks, size = [], 0
    with os.fdopen(os.open(path, flags), 'rb') as handle:
        if signature(os.fstat(handle.fileno())) != signature(before):
            raise ValueError('File changed before reading.')
        while True:
            cancel()
            chunk = handle.read(min(65536, max_bytes + 1 - size))
            if not chunk:
                break
            size += len(chunk)
            if size > max_bytes:
                raise ValueError('File grew past the byte limit while reading.')
            chunks.append(chunk)
        after = os.fstat(handle.fileno())
    latest = path.lstat()
    if linked(latest) or signature(before) != signature(after) or signature(before) != signature(latest) or size != before.st_size or not path.resolve().is_relative_to(root):
        raise ValueError('File changed while reading; review it again.')
    return b''.join(chunks)


def file_metadata(info):
    return {'bytes': info.st_size, 'modified_at': datetime.fromtimestamp(info.st_mtime, timezone.utc).isoformat()}


def python_metadata(text):
    try:
        tree = ast.parse(text)
    except (SyntaxError, RecursionError):
        return {'syntax': 'Could not parse Python structure; may be incomplete or use a different Python version.'}
    functions = [{'name': node.name, 'line': node.lineno} for node in ast.walk(tree) if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))]
    classes = [{'name': node.name, 'line': node.lineno} for node in ast.walk(tree) if isinstance(node, ast.ClassDef)]
    imports = sorted({node.module or '' for node in ast.walk(tree) if isinstance(node, ast.ImportFrom)} | {alias.name for node in ast.walk(tree) if isinstance(node, ast.Import) for alias in node.names})
    return {'module_docstring': (ast.get_docstring(tree) or '')[:2000], 'imports': imports[:100],
            'functions': functions[:100], 'classes': classes[:100], 'function_count': len(functions), 'class_count': len(classes)}


def extract(raw, name, max_chars, cancel=lambda: None):
    cancel()
    extension = Path(name).suffix.lower()
    metadata = {'sha256': hashlib.sha256(raw).hexdigest()}
    coverage = {'partial': False, 'reason': '', 'characters': 0}
    if extension in IMAGE_EXTENSIONS:
        metadata['content_policy'] = 'Metadata only; pixels, OCR and visual interpretation were not requested.'
        if extension == '.svg':
            metadata['format'] = 'SVG'
            opening = re.search(r'<svg\b[^>]*>', raw[:16384].decode('utf-8', errors='replace'), re.I)
            if opening:
                for key in ('width', 'height', 'viewBox'):
                    match = re.search(r'\b' + key + r'\s*=\s*["\x27]([^"\x27]{1,200})["\x27]', opening[0])
                    if match:
                        metadata[key] = match[1]
        else:
            from PIL import Image, ExifTags
            with Image.open(io.BytesIO(raw)) as image:
                metadata.update(format=image.format, width=image.width, height=image.height, mode=image.mode)
                # PNG's override calls load() to find late chunks. Only read the
                # metadata available at open; never decode pixels for metadata.
                exif = Image.Image.getexif(image)
                metadata['metadata_scope'] = 'Available header metadata; no pixel decoding or late image-chunk scan.'
                metadata['exif'] = {ExifTags.TAGS.get(key, str(key)): str(value)[:500] for key, value in exif.items()
                                    if not isinstance(value, bytes) and key not in {37500, 37510}}
        return {'kind': 'image', 'text': '', 'status': 'metadata_only', 'metadata': metadata, 'coverage': coverage}
    if extension == '.pdf':
        from PyPDF2 import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            return {'kind': 'pdf', 'text': '', 'status': 'unreadable', 'metadata': {**metadata, 'encrypted': True}, 'coverage': {**coverage, 'reason': 'Encrypted PDF; no password or OCR attempted.'}}
        metadata.update(page_count=len(reader.pages), title=str((reader.metadata or {}).get('/Title', ''))[:500])
        readable, unreadable, pages, used, capped, failed = [], [], [], 0, [], []
        for index, page in enumerate(reader.pages[:500]):
            cancel()
            number = index + 1
            try:
                text = (page.extract_text() or '').strip().replace('\r\n', '\n').replace('\r', '\n')
            except Exception:
                failed.append(number)
                continue
            if not text:
                unreadable.append(number)
                continue
            readable.append(number)
            prefix = f'--- Page {number} ---\n'
            available = max(0, max_chars - used - len(prefix))
            if len(text) > available:
                capped.append(number)
            if available:
                value = prefix + text[:available]
                pages.append(value); used += len(value) + 2
        metadata.update(text_layer_pages=readable, pages_without_extractable_text=unreadable, pages_limited_by_text_cap=capped,
                        pages_with_extraction_errors=failed, pages_not_inspected=max(0, len(reader.pages) - 500), ocr_attempted=False)
        text = '\n\n'.join(pages)
        partial = bool(unreadable or capped or failed or len(reader.pages) > 500)
        coverage.update(partial=partial, characters=len(text), reason='Only the available native text layer was reviewed; unreadable or capped pages are listed.' if partial else '')
        if not text:
            coverage['reason'] = 'No extractable text. Scanned/image-only pages were not OCRed or visually analyzed.'
        return {'kind': 'pdf', 'text': text, 'status': 'readable' if text else 'unreadable', 'metadata': metadata, 'coverage': coverage}
    if extension == '.docx':
        from docx import Document
        document = Document(io.BytesIO(raw))
        values = [paragraph.text for paragraph in document.paragraphs]
        values.extend(' | '.join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
        text = '\n'.join(values)
        metadata['content_policy'] = 'Extracted paragraphs and table text only; embedded images were not analyzed.'
    elif extension in TEXT_EXTENSIONS or Path(name).name.lower() in {'readme', 'license', 'makefile', 'dockerfile'}:
        encoding = 'utf-16' if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError:
            raise ValueError('Text encoding is not UTF-8 or BOM-marked UTF-16; content not guessed.')
        if '\0' in text:
            raise ValueError('Binary content; no text analysis attempted.')
        metadata['encoding'] = encoding
    else:
        return {'kind': 'other', 'text': '', 'status': 'metadata_only', 'metadata': metadata, 'coverage': {**coverage, 'reason': 'Unsupported content type; filesystem metadata only.'}}
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    length = len(text)
    text = text[:max_chars]
    coverage.update(characters=len(text), original_characters=length, partial=length > max_chars,
                    reason='Only the configured character limit was reviewed.' if length > max_chars else '')
    kind = 'python' if extension == '.py' else 'markdown' if extension in {'.md', '.markdown'} else 'document' if extension == '.docx' else 'text'
    if kind == 'python':
        metadata.update(python_metadata(text))
    if kind == 'markdown':
        metadata['headings'] = re.findall(r'^#{1,6}\s+(.+)$', text, re.M)[:100]
    return {'kind': kind, 'text': text, 'status': 'readable' if text.strip() else 'empty', 'metadata': metadata, 'coverage': coverage}


def text_batches(text, size=4000):
    """Contiguous source slices with line ranges; no omitted tail or overlap."""
    offset, line = 0, 1
    while offset < len(text):
        end = min(offset + size, len(text))
        if end < len(text):
            boundary = text.rfind('\n', offset + size // 2, end)
            if boundary >= 0:
                end = boundary + 1
        value = text[offset:end]
        last_line = line + value.count('\n') - int(value.endswith('\n'))
        yield {'text': value, 'first_line': line, 'last_line': max(line, last_line)}
        line += value.count('\n'); offset = end
