"""Point bookmarks, internal links and a locally generated, linked contents block.

The contents block is a readable snapshot in DOCX, tagged for this editor to
rebuild. It is deliberately not a Word TOC field and needs no Office process.
"""
import re
from uuid import uuid4

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAME = re.compile(r'[A-Za-z_][A-Za-z0-9_]{0,39}')
HEADING_ID = re.compile(r'LAW_H[0-9a-f]{32}')
TOC_TAG = 'LAW_TOC_V1:'


def valid_name(value):
    return isinstance(value, str) and bool(NAME.fullmatch(value))


def walk(node):
    yield node
    for child in node.get('content', []):
        yield from walk(child)


def contents_attrs(attrs):
    title, level = attrs.get('title', 'Contents'), attrs.get('maxLevel', 3)
    if not isinstance(title, str) or not title.strip() or len(title) > 100 or re.search(r'[\x00-\x1f]', title):
        raise ValueError('The contents title must contain 1–100 printable characters.')
    if type(level) is not int or level not in (1, 2, 3):
        raise ValueError('The contents can include heading levels 1–3.')
    return {'title': title.strip(), 'maxLevel': level}


def validate_references(model):
    names, bookmarks, contents = set(), 0, 0
    for node in walk(model):
        attrs = node.get('attrs') or {}
        name = None
        if node['type'] in ('paragraph', 'heading') and attrs.get('referenceId') is not None:
            name = attrs['referenceId']
            if not isinstance(name, str) or not HEADING_ID.fullmatch(name):
                raise ValueError('Invalid heading reference ID.')
        if node['type'] == 'bookmark':
            bookmarks += 1
            name = attrs.get('name')
            if not valid_name(name) or HEADING_ID.fullmatch(name):
                raise ValueError('Invalid or reserved bookmark name.')
        if name:
            if name.lower() in names:
                raise ValueError('Bookmark and heading reference names must be unique.')
            names.add(name.lower())
        if node['type'] == 'tableOfContents':
            contents += 1; contents_attrs(attrs)
    if bookmarks > 100 or contents > 1:
        raise ValueError('A document supports up to 100 bookmarks and one table of contents.')


def element(tag, **attrs):
    result = OxmlElement('w:' + tag)
    for key, value in attrs.items(): result.set(qn('w:' + key), str(value))
    return result


def internal_link(paragraph, run, name):
    link = element('hyperlink', anchor=name, history='1')
    paragraph._p.remove(run._r); link.append(run._r); paragraph._p.append(link)
    run.font.color.rgb = RGBColor.from_string('0563C1'); run.underline = True


class ReferenceWriter:
    def __init__(self, model, styles):
        self.ids, self.headings, self.serial = {}, [], 0
        for node in walk(model):
            if node['type'] not in ('paragraph', 'heading'): continue
            attrs = node.get('attrs') or {}
            sid = attrs.get('styleId') or (f'heading-{attrs.get("level", 1)}' if node['type'] == 'heading' else 'normal')
            level, name = styles[sid]['level'], attrs.get('referenceId')
            if level and not name: name = 'LAW_H' + uuid4().hex
            if name: self.ids[id(node)] = name
            if level:
                text = ''.join(child.get('text', '') if child['type'] == 'text' else '\n' if child['type'] == 'hardBreak' else '' for child in node.get('content', []))
                self.headings.append((name, text or 'Untitled heading', level))

    def bookmark(self, paragraph, name):
        self.serial += 1
        paragraph._p.append(element('bookmarkStart', id=self.serial, name=name))
        paragraph._p.append(element('bookmarkEnd', id=self.serial))

    def contents(self, parent, attrs):
        attrs = contents_attrs(attrs)
        sdt, props, content = element('sdt'), element('sdtPr'), element('sdtContent')
        props.append(element('alias', val=attrs['title'])); props.append(element('tag', val=TOC_TAG + str(attrs['maxLevel'])))
        sdt.extend([props, content])
        placeholder = parent.add_paragraph(); placeholder._p.addprevious(sdt); placeholder._p.getparent().remove(placeholder._p)

        def add_paragraph():
            paragraph = parent.add_paragraph()
            content.append(paragraph._p)
            return paragraph

        title = add_paragraph(); title.paragraph_format.keep_with_next = True
        run = title.add_run(attrs['title']); run.bold = True; run.font.size = Pt(18)
        for name, text, level in self.headings:
            if level > attrs['maxLevel']: continue
            paragraph = add_paragraph(); paragraph.paragraph_format.left_indent = Inches((level - 1) * .2)
            paragraph.paragraph_format.space_after = Pt(4)
            internal_link(paragraph, paragraph.add_run(text), name)


class ReferenceReader:
    def __init__(self, warn):
        self.warn, self.names, self.count, self.contents_seen = warn, set(), 0, False

    def bookmark(self, start):
        name = start.get(qn('w:name'))
        if not valid_name(name) or name.lower() in self.names:
            self.warn('Invalid or duplicate bookmarks were omitted. Review internal links.'); return None
        is_heading = bool(HEADING_ID.fullmatch(name))
        if not is_heading and self.count >= 100:
            self.warn('Bookmarks beyond the 100-bookmark limit were omitted.'); return None
        self.names.add(name.lower())
        if not is_heading: self.count += 1
        after = start.getnext()
        if after is None or after.tag != qn('w:bookmarkEnd') or after.get(qn('w:id')) != start.get(qn('w:id')):
            self.warn('Bookmark ranges were converted to point bookmarks at their starting positions.')
        return name

    def contents(self, sdt, top_level):
        tag = sdt.find('w:sdtPr/w:tag', sdt.nsmap)
        value = tag.get(qn('w:val'), '') if tag is not None else ''
        if not value.startswith(TOC_TAG): return None
        alias = sdt.find('w:sdtPr/w:alias', sdt.nsmap)
        try:
            attrs = contents_attrs({'title': alias.get(qn('w:val')) if alias is not None else 'Contents', 'maxLevel': int(value[len(TOC_TAG):])})
        except (ValueError, TypeError):
            self.warn('Invalid contents settings were reduced to their visible text.'); return None
        if not top_level or self.contents_seen:
            self.warn('Additional or nested contents blocks were reduced to visible text.'); return None
        self.contents_seen = True
        return {'type': 'tableOfContents', 'attrs': attrs}


def missing_links(model):
    names = set()
    for node in walk(model):
        attrs = node.get('attrs') or {}
        if node['type'] == 'bookmark': names.add(attrs['name'])
        elif node['type'] == 'documentCaption': names.add(attrs['id'])
        elif attrs.get('referenceId'): names.add(attrs['referenceId'])
    return any(mark.get('type') == 'link' and mark.get('attrs', {}).get('href', '').startswith('#') and
               mark['attrs']['href'][1:] not in names for node in walk(model) for mark in node.get('marks', []))
