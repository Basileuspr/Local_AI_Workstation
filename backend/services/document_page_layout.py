"""Single-section headers, footers and page fields for the document editor."""
import math
import re

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

STORIES = {
    'header': 'header', 'footer': 'footer',
    'firstHeader': 'first_page_header', 'firstFooter': 'first_page_footer',
    'evenHeader': 'even_page_header', 'evenFooter': 'even_page_footer',
}
ALIGNMENTS = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER, 'right': WD_ALIGN_PARAGRAPH.RIGHT}
NUMBER_FORMATS = ('decimal', 'lowerRoman', 'upperRoman', 'lowerLetter', 'upperLetter')
DEFAULT_PAGE_DETAILS = {
    **{key: '' for key in STORIES}, **{key + 'Alignment': 'left' for key in STORIES},
    'differentFirstPage': False, 'differentOddEven': False,
    'headerDistance': .5, 'footerDistance': .5,
    'pageNumbers': False, 'pageNumberPosition': 'footer', 'pageNumberAlignment': 'center',
    'pageNumberStyle': 'page', 'pageNumberFormat': 'decimal', 'pageNumberStart': 1,
    'pageNumberFirstPage': True,
}
PAGINATION = {'pageBreakBefore': 'page_break_before', 'keepWithNext': 'keep_with_next',
              'keepTogether': 'keep_together', 'widowControl': 'widow_control'}


def validate_page_details(layout):
    for key in STORIES:
        text = layout[key]
        if not isinstance(text, str) or len(text) > 2000 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', text):
            raise ValueError('Headers and footers must be plain text up to 2,000 characters.')
        if layout[key + 'Alignment'] not in ('left', 'center', 'right'):
            raise ValueError('Invalid header or footer alignment.')
    for key in ('differentFirstPage', 'differentOddEven', 'pageNumbers', 'pageNumberFirstPage'):
        if not isinstance(layout[key], bool):
            raise ValueError('Invalid header, footer or page number setting.')
    for key in ('headerDistance', 'footerDistance'):
        value = layout[key]
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or not 0 <= value <= 3:
            raise ValueError('Header and footer distances must be between 0 and 3 inches.')
    for key, options in {'pageNumberPosition': ('header', 'footer'), 'pageNumberAlignment': ('left', 'center', 'right'),
                         'pageNumberStyle': ('number', 'page', 'pageOf'), 'pageNumberFormat': NUMBER_FORMATS}.items():
        if layout[key] not in options:
            raise ValueError('Invalid page number format or position.')
    start = layout['pageNumberStart']
    if isinstance(start, bool) or not isinstance(start, int) or not 1 <= start <= 9999:
        raise ValueError('Page numbering must start at a whole number from 1 to 9,999.')


def element(tag, **attrs):
    node = OxmlElement('w:' + tag)
    for key, value in attrs.items():
        node.set(qn('w:' + key), str(value))
    return node


def write_page_details(document, layout):
    section = document.sections[0]
    section.different_first_page_header_footer = layout['differentFirstPage']
    document.settings.odd_and_even_pages_header_footer = layout['differentOddEven']
    section.header_distance = Inches(layout['headerDistance'])
    section.footer_distance = Inches(layout['footerDistance'])
    numbering = element('pgNumType', fmt=layout['pageNumberFormat'], start=layout['pageNumberStart'])
    section._sectPr.insert_element_before(numbering, 'w:cols', 'w:formProt', 'w:vAlign', 'w:noEndnote',
                                         'w:titlePg', 'w:textDirection', 'w:bidi', 'w:rtlGutter', 'w:docGrid',
                                         'w:printerSettings', 'w:sectPrChange')
    for key, attribute in STORIES.items():
        is_first, is_even = key.startswith('first'), key.startswith('even')
        active = (not is_first or layout['differentFirstPage']) and (not is_even or layout['differentOddEven'])
        position = 'header' if key.lower().endswith('header') else 'footer'
        numbered = (active and layout['pageNumbers'] and position == layout['pageNumberPosition']
                    and (not is_first or layout['pageNumberFirstPage']))
        # Retain text in disabled variants so re-enabling them does not destroy it.
        if not layout[key] and not numbered:
            continue
        story = getattr(section, attribute)
        paragraph = story.paragraphs[0]
        paragraph.text = layout[key]
        paragraph.alignment = ALIGNMENTS[layout[key + 'Alignment']]
        paragraph.paragraph_format.space_after = Pt(0)
        if numbered:
            if layout[key]:
                paragraph = story.add_paragraph()
            paragraph.alignment = ALIGNMENTS[layout['pageNumberAlignment']]
            paragraph.paragraph_format.space_after = Pt(0)
            if layout['pageNumberStyle'] != 'number':
                paragraph.add_run('Page ')
            paragraph._p.append(element('fldSimple', instr='PAGE', dirty='true'))
            if layout['pageNumberStyle'] == 'pageOf':
                paragraph.add_run(' of ')
                paragraph._p.append(element('fldSimple', instr='NUMPAGES', dirty='true'))
    if layout['pageNumbers']:
        document.settings.element.append(element('updateFields', val='true'))


def _paragraph_tokens(paragraph):
    """Read simple and complex fields; keep cached text separate from instructions."""
    output, stack = [], []

    def emit(value):
        if not stack:
            output.append(value)
        elif stack[-1]['result']:
            stack[-1]['tokens'].append(value)

    def walk(node):
        tag = node.tag
        if tag == qn('w:pPr'):
            return
        if tag == qn('w:fldSimple'):
            emit({'instruction': node.get(qn('w:instr'), ''), 'tokens': _paragraph_tokens(node)})
            return
        if tag == qn('w:fldChar'):
            kind = node.get(qn('w:fldCharType'))
            if kind == 'begin':
                stack.append({'instruction': '', 'result': False, 'tokens': []})
            elif kind == 'separate' and stack:
                stack[-1]['result'] = True
            elif kind == 'end' and stack:
                emit(stack.pop())
            return
        if tag == qn('w:instrText'):
            if stack and not stack[-1]['result']:
                stack[-1]['instruction'] += node.text or ''
            return
        if tag in (qn('w:t'), qn('w:tab'), qn('w:br'), qn('w:cr')):
            emit(node.text or '' if tag == qn('w:t') else '\t' if tag == qn('w:tab') else '\n')
            return
        for child in node:
            walk(child)

    for child in paragraph:
        walk(child)
    # A malformed field must not turn its instruction into visible user text.
    while stack:
        emit(stack.pop())
    return output


def _cached_text(tokens):
    return ''.join(token if isinstance(token, str) else _cached_text(token['tokens']) for token in tokens)


def _number_line(tokens):
    template, switches, count = '', [], 0
    for token in tokens:
        if isinstance(token, str):
            template += token
            continue
        match = re.fullmatch(r'\s*(PAGE|NUMPAGES)\s*((?:\\\*\s+(?:MERGEFORMAT|Arabic|roman|alphabetic)\s*)*)',
                             token['instruction'], re.I)
        if not match:
            return None
        count += 1
        template += '{' + match[1].upper() + '}'
        if match[1].upper() == 'PAGE':
            switches += re.findall(r'\\\*\s+(Arabic|roman|ROMAN|alphabetic|ALPHABETIC)\b', match[2])
    template = re.sub(r'\s+', ' ', template).strip().lower()
    styles = {'{page}': 'number', 'page {page}': 'page', 'page {page} of {numpages}': 'pageOf'}
    if template not in styles or count != (2 if styles[template] == 'pageOf' else 1):
        return None
    formats = {'Arabic': 'decimal', 'roman': 'lowerRoman', 'ROMAN': 'upperRoman',
               'alphabetic': 'lowerLetter', 'ALPHABETIC': 'upperLetter'}
    return {'style': styles[template], 'format': formats.get(switches[-1]) if switches else None}


def read_page_details(document, warn):
    layout = dict(DEFAULT_PAGE_DETAILS)
    section = document.sections[0]
    layout['differentFirstPage'] = section.different_first_page_header_footer
    layout['differentOddEven'] = document.settings.odd_and_even_pages_header_footer
    for key, attribute in [('headerDistance', 'header_distance'), ('footerDistance', 'footer_distance')]:
        value = getattr(section, attribute)
        if value is not None:
            layout[key] = round(min(3, max(0, value.inches)), 3)
            if abs(layout[key] - value.inches) > .001:
                warn('Header or footer distance was adjusted to the supported 0–3 inch range.')
    pgnum = section._sectPr.find(qn('w:pgNumType'))
    if pgnum is not None:
        form = pgnum.get(qn('w:fmt'), 'decimal')
        if form in NUMBER_FORMATS:
            layout['pageNumberFormat'] = form
        else:
            warn('An unsupported page-number format was changed to decimal numbering.')
        try:
            start = int(pgnum.get(qn('w:start'), '1'))
            layout['pageNumberStart'] = min(9999, max(1, start))
            if layout['pageNumberStart'] != start:
                warn('Page numbering start was adjusted to the supported 1–9,999 range.')
        except ValueError:
            warn('An invalid page numbering start was reset to 1.')

    found = []
    for key, attribute in STORIES.items():
        story = getattr(section, attribute)
        if story.is_linked_to_previous:
            continue  # First-section stories have no previous definition to inherit.
        warn('Headers and footers use plain text and supported page-number lines. Rich formatting, tables and pictures in these areas are not retained.')
        lines, alignments = [], []
        for paragraph in story.paragraphs:
            tokens = _paragraph_tokens(paragraph._p)
            number_line = _number_line(tokens)
            alignment = next((name for name, value in ALIGNMENTS.items() if value == paragraph.alignment), 'left')
            active = (not key.startswith('first') or layout['differentFirstPage']) and (not key.startswith('even') or layout['differentOddEven'])
            if number_line and active:
                found.append({**number_line, 'key': key, 'position': 'header' if key.lower().endswith('header') else 'footer', 'alignment': alignment})
            else:
                if any(not isinstance(token, str) for token in tokens):
                    warn('A custom or unsupported header/footer field was reduced to its cached text; it will not update on export.')
                lines.append(_cached_text(tokens))
                alignments.append(alignment)
        text = '\n'.join(lines)
        if len(text) > 2000:
            warn('Header or footer text beyond 2,000 characters was omitted.')
        layout[key] = text[:2000]
        if alignments:
            layout[key + 'Alignment'] = alignments[0]
            if len(set(alignments)) > 1:
                warn('Mixed header/footer paragraph alignment was simplified to the first text paragraph’s alignment.')
    if found:
        # Prefer the regular story's settings; all enabled variants share one number style.
        primary = next((item for item in found if item['key'] in ('header', 'footer')), found[0])
        layout.update(pageNumbers=True, pageNumberPosition=primary['position'], pageNumberStyle=primary['style'],
                      pageNumberAlignment=primary['alignment'], pageNumberFirstPage=any(item['key'].startswith('first') for item in found) if layout['differentFirstPage'] else True)
        if primary['format']:
            layout['pageNumberFormat'] = primary['format']
        expected = [layout['pageNumberPosition']]
        if layout['differentFirstPage'] and layout['pageNumberFirstPage']:
            expected.append('first' + layout['pageNumberPosition'].capitalize())
        if layout['differentOddEven']:
            expected.append('even' + layout['pageNumberPosition'].capitalize())
        if sorted(item['key'] for item in found) != sorted(expected) or any(
                (item['position'], item['style'], item['alignment'], item['format']) !=
                (primary['position'], primary['style'], primary['alignment'], primary['format']) for item in found):
            warn('Different page-number lines were normalized to one position, alignment and format across the enabled page variants.')
    return layout
