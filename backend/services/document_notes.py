"""Native DOCX footnote/endnote parts with bounded editable note bodies."""
import re
from types import SimpleNamespace
from uuid import uuid4

from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import XmlPart
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from services.document_references import element, walk

FORMATS = {'decimal', 'lowerRoman', 'upperRoman', 'lowerLetter', 'upperLetter'}
DEFAULTS = {'footnote': {'format': 'decimal', 'start': 1}, 'endnote': {'format': 'lowerRoman', 'start': 1}}


def note_settings(model):
    raw = (model.get('attrs') or {}).get('noteSettings')
    if raw is None: raw = {}
    if not isinstance(raw, dict): raise ValueError('Invalid note numbering settings.')
    result = {}
    for kind in DEFAULTS:
        value = raw.get(kind, {})
        if not isinstance(value, dict): raise ValueError('Invalid note numbering settings.')
        value = {**DEFAULTS[kind], **value}
        if not isinstance(value.get('format'), str) or value['format'] not in FORMATS:
            raise ValueError('Unsupported note number format.')
        if type(value['start']) is not int or not 1 <= value['start'] <= 9999:
            raise ValueError('Note starting numbers must be 1–9999.')
        result[kind] = {'format': value['format'], 'start': value['start']}
    return result


def validate_note_node(attrs):
    if attrs.get('kind') not in ('footnote', 'endnote') or not isinstance(attrs.get('id'), str) or not re.fullmatch(r'note-[0-9a-f]{32}', attrs['id']):
        raise ValueError('Invalid note kind or reference ID.')
    body = attrs.get('body')
    if not isinstance(body, dict) or body.get('type') != 'doc' or not isinstance(body.get('content'), list) or not 1 <= len(body['content']) <= 40:
        raise ValueError('A note needs 1–40 paragraphs.')
    text_length = 0
    for paragraph in body['content']:
        if not isinstance(paragraph, dict) or paragraph.get('type') != 'paragraph' or not isinstance(paragraph.get('content', []), list):
            raise ValueError('Notes support paragraphs and text formatting.')
        for child in paragraph.get('content', []):
            if not isinstance(child, dict) or child.get('type') not in ('text', 'hardBreak'):
                raise ValueError('Notes support text and line breaks; nested notes are unsupported.')
            if child['type'] == 'text':
                if not isinstance(child.get('text'), str): raise ValueError('Invalid note text.')
                try: text_length += len(child['text'].encode('utf-16-le')) // 2
                except UnicodeEncodeError as exc: raise ValueError('Invalid note text encoding.') from exc
            if not isinstance(child.get('marks', []), list) or len(child.get('marks', [])) > 10:
                raise ValueError('Unsupported note formatting.')
            for mark in child.get('marks', []):
                if not isinstance(mark, dict) or not isinstance(mark.get('attrs') or {}, dict): raise ValueError('Invalid note formatting attributes.')
                if mark.get('type') == 'link' and str((mark.get('attrs') or {}).get('href', '')).startswith('#'):
                    raise ValueError('Note links must use http, https or mailto.')
    if text_length > 10000: raise ValueError('A note can contain at most 10,000 characters.')


def validate_notes(model):
    note_settings(model)
    notes = [node for node in walk(model) if node['type'] == 'documentNote']
    if len(notes) > 200: raise ValueError('A document supports at most 200 notes.')
    if len({node['attrs']['id'] for node in notes}) != len(notes): raise ValueError('Note reference IDs must be unique.')


class NoteWriter:
    def __init__(self, document, model, write_run):
        self.references = {}
        settings = note_settings(model)
        notes = [node for node in walk(model) if node['type'] == 'documentNote']
        for kind in DEFAULTS:
            props = element(kind + 'Pr')
            props.extend([element('pos', val='pageBottom' if kind == 'footnote' else 'docEnd'),
                          element('numFmt', val=settings[kind]['format']), element('numStart', val=settings[kind]['start']), element('numRestart', val='continuous')])
            document.sections[0]._sectPr.insert_element_before(props, 'w:endnotePr', 'w:type', 'w:pgSz', 'w:pgMar')
            items = [node for node in notes if node['attrs']['kind'] == kind]
            if not items: continue
            title = 'Footnote' if kind == 'footnote' else 'Endnote'
            style_name = title + ' Text'
            style = document.styles[style_name] if style_name in document.styles else document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
            style.font.name = 'Calibri'; style.font.size = Pt(10); style.paragraph_format.space_after = Pt(0); style.paragraph_format.line_spacing = 1.0
            reference_name = title + ' Reference'
            reference_style = document.styles[reference_name] if reference_name in document.styles else document.styles.add_style(reference_name, WD_STYLE_TYPE.CHARACTER)
            reference_style.font.superscript = True
            root = element(kind + 's')
            for identifier, type_name, marker in [(-1, 'separator', 'separator'), (0, 'continuationSeparator', 'continuationSeparator')]:
                note = element(kind, id=identifier, type=type_name); paragraph = element('p'); run = element('r'); run.append(element(marker)); paragraph.append(run); note.append(paragraph); root.append(note)
            part = XmlPart(PackURI('/word/' + kind + 's.xml'), CT.WML_FOOTNOTES if kind == 'footnote' else CT.WML_ENDNOTES, root, document.part.package)
            document.part.relate_to(part, RT.FOOTNOTES if kind == 'footnote' else RT.ENDNOTES)
            parent = SimpleNamespace(part=part)
            for identifier, node in enumerate(items, 1):
                attrs = node['attrs']; self.references[attrs['id']] = (kind, identifier, reference_style.style_id)
                note = element(kind, id=identifier); root.append(note)
                for index, block in enumerate(attrs['body']['content']):
                    paragraph = element('p'); ppr = element('pPr'); ppr.append(element('pStyle', val=style.style_id)); paragraph.append(ppr); note.append(paragraph)
                    if (block.get('attrs') or {}).get('textAlign'): ppr.append(element('jc', val=block['attrs']['textAlign']))
                    p = Paragraph(paragraph, parent)
                    if not index:
                        run = p.add_run(); run._r.get_or_add_rPr().append(element('rStyle', val=reference_style.style_id)); run._r.append(element(kind + 'Ref')); p.add_run().add_tab()
                    for child in block.get('content', []):
                        if child['type'] == 'hardBreak': p.add_run().add_break()
                        else: write_run(p, child)

    def reference(self, paragraph, node):
        kind, identifier, style = self.references[node['attrs']['id']]
        run = paragraph.add_run(); run._r.get_or_add_rPr().append(element('rStyle', val=style)); run.font.superscript = True
        run._r.append(element(kind + 'Reference', id=identifier))


class NoteReader:
    def __init__(self, document, warn, run_marks, safe_link):
        self.document, self.warn, self.run_marks, self.safe_link = document, warn, run_marks, safe_link
        self.parts, self.definitions, self.used = {}, {}, set()
        self.settings = {}
        for kind in DEFAULTS:
            settings = dict(DEFAULTS[kind])
            for parent in [document.settings.element, document.sections[0]._sectPr]:
                props = parent.find(qn('w:' + kind + 'Pr'))
                if props is None: continue
                fmt, start, restart, position = [props.find(qn('w:' + tag)) for tag in ['numFmt', 'numStart', 'numRestart', 'pos']]
                if fmt is not None:
                    value = fmt.get(qn('w:val'))
                    if value in FORMATS: settings['format'] = value
                    else: warn('Unsupported note numbering formats were normalized to decimal.'); settings['format'] = 'decimal'
                if start is not None:
                    try: value = int(start.get(qn('w:val')))
                    except (ValueError, TypeError): value = 1
                    settings['start'] = max(1, min(9999, value))
                    if settings['start'] != value: warn('Note starting numbers were adjusted to 1–9999.')
                if restart is not None and restart.get(qn('w:val')) != 'continuous': warn('Note page/section restarts were normalized to continuous numbering.')
                if position is not None and position.get(qn('w:val')) != ('pageBottom' if kind == 'footnote' else 'docEnd'): warn('Note placement was normalized to page-bottom footnotes and document-end endnotes.')
            self.settings[kind] = settings
            relation_type = RT.FOOTNOTES if kind == 'footnote' else RT.ENDNOTES
            for relation in document.part.rels.values():
                if relation.reltype != relation_type: continue
                if relation.is_external: warn('External note parts are not fetched.'); continue
                part = relation.target_part
                root = parse_xml(part.blob)
                if root.tag != qn('w:' + kind + 's'): raise ValueError('The document has an invalid notes part.')
                if sum(1 for node in root.iter() if node.tag in (qn('w:p'), qn('w:r'))) > 20000: raise ValueError('The notes part exceeds 20,000 paragraphs/runs.')
                self.parts[kind] = part
                for note in root.findall(qn('w:' + kind)):
                    if note.get(qn('w:type'), 'normal') != 'normal': continue
                    key = (kind, note.get(qn('w:id')))
                    if key in self.definitions: warn('Duplicate note definitions were normalized to the first definition.'); continue
                    self.definitions[key] = note

    def read(self, reference, kind):
        key = (kind, reference.get(qn('w:id'))); self.used.add(key)
        definition = self.definitions.get(key)
        if reference.get(qn('w:customMarkFollows')) not in (None, '0', 'false'):
            self.warn('Custom note marks were normalized to automatic numbering. Review any adjacent source symbols.')
        if definition is None:
            self.warn('A note reference has no matching text. A visible placeholder was inserted for repair.')
            body = {'type': 'doc', 'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': '[Missing note text]'}]}]}
        else: body = self.body(definition, self.parts[kind])
        return {'type': 'documentNote', 'attrs': {'id': 'note-' + uuid4().hex, 'kind': kind, 'body': body}}

    def body(self, definition, part):
        paragraphs, budget = [], 10000
        source = list(definition.iter(qn('w:p')))
        if len(source) > 40: self.warn('Note text beyond 40 paragraphs was omitted.')
        if any(node.tag in {qn('w:' + tag) for tag in ('tbl', 'drawing', 'sdt', 'fldChar', 'fldSimple', 'ins', 'del')} for node in definition.iter()):
            self.warn('Notes were normalized to supported paragraphs and text formatting; tables/fields and other advanced note content are not retained.')
        for paragraph in source[:40]:
            content, attrs, prefix = [], {}, False
            properties = paragraph.find(qn('w:pPr'))
            if properties is not None and any(child.tag not in (qn('w:pStyle'), qn('w:jc')) for child in properties):
                self.warn('Note paragraph spacing, lists and other layout properties were normalized to plain paragraphs.')
            alignment = paragraph.find('w:pPr/w:jc', paragraph.nsmap)
            if alignment is not None and alignment.get(qn('w:val')) in ('left', 'center', 'right', 'justify'): attrs['textAlign'] = alignment.get(qn('w:val'))
            for child in paragraph:
                href = None
                if child.tag == qn('w:hyperlink'):
                    relation = part.rels.get(child.get(qn('r:id')))
                    if relation and relation.is_external and self.safe_link(relation.target_ref) and not relation.target_ref.startswith('#'): href = relation.target_ref
                    else: self.warn('Unsupported internal note links were reduced to visible text.')
                    runs = child.findall(qn('w:r'))
                elif child.tag == qn('w:r'): runs = [child]
                elif child.tag == qn('w:fldSimple'): runs = child.findall(qn('w:r'))
                else:
                    if child.tag not in (qn('w:pPr'), qn('w:bookmarkStart'), qn('w:bookmarkEnd'), qn('w:proofErr')): self.warn('Some unsupported note content was omitted. Review notes against the source.')
                    continue
                for run in runs:
                    marks = self.run_marks(Run(run, None))
                    if href: marks.append({'type': 'link', 'attrs': {'href': href}})
                    for item in run:
                        if item.tag in (qn('w:footnoteRef'), qn('w:endnoteRef')): prefix = True; continue
                        if item.tag == qn('w:tab') and prefix: prefix = False; continue
                        if item.tag in (qn('w:t'), qn('w:tab')):
                            value = item.text or '' if item.tag == qn('w:t') else '\t'
                            if prefix and value.startswith(' '): value = value[1:]
                            prefix = False
                            kept, units = [], 0
                            for character in value:
                                size = 2 if ord(character) > 0xFFFF else 1
                                if units + size > budget: break
                                kept.append(character); units += size
                            if len(kept) != len(value): self.warn('Note text beyond 10,000 characters was omitted.')
                            budget -= units
                            if kept: content.append({'type': 'text', 'text': ''.join(kept), 'marks': marks})
                        elif item.tag in (qn('w:br'), qn('w:cr')): content.append({'type': 'hardBreak'})
                        elif item.tag not in (qn('w:rPr'), qn('w:lastRenderedPageBreak')): self.warn('Some unsupported note content was omitted. Review notes against the source.')
            paragraphs.append({'type': 'paragraph', **({'attrs': attrs} if attrs else {}), 'content': content})
        return {'type': 'doc', 'content': paragraphs or [{'type': 'paragraph'}]}

    def finish(self):
        if set(self.definitions) - self.used: self.warn('Unreferenced note definitions were omitted from the editable copy.')
