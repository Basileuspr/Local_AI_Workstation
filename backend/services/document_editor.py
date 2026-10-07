"""Rich document conversion, entirely in memory using python-docx.

Imports are normalized editable copies, never edits to the original package.
No Word, COM, subprocesses, filesystem paths, external links or AI are used.
"""
import base64
import binascii
from io import BytesIO
import math
import re
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree
from PIL import Image
from services.document_page_layout import (DEFAULT_PAGE_DETAILS, PAGINATION, validate_page_details,
                                           write_page_details, read_page_details)
from services.document_tables import table_grid, write_table, read_table
from services.document_styles import FONTS, style_catalog, write_styles, StyleReader
from services.document_references import (valid_name, HEADING_ID, ReferenceWriter, ReferenceReader,
                                          validate_references, internal_link, missing_links)
from services.document_notes import NoteWriter, NoteReader, validate_notes, validate_note_node
from services.document_citations import CitationWriter, CitationReader, validate_citations

MAX_FILE = 24 * 1024**2
MAX_MODEL = 32 * 1024**2
PAPERS = {'Letter': (8.5, 11), 'A4': (8.2677, 11.6929), 'Legal': (8.5, 14)}
DEFAULT_LAYOUT = {'paper': 'Letter', 'orientation': 'portrait', 'top': 1, 'bottom': 1, 'left': 1, 'right': 1, **DEFAULT_PAGE_DETAILS}
ALIGN = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
         'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}
SUPPORTED = {'doc', 'paragraph', 'heading', 'text', 'hardBreak', 'pageBreak',
             'bulletList', 'orderedList', 'listItem', 'table', 'tableRow', 'tableCell', 'tableHeader', 'image', 'bookmark', 'tableOfContents', 'documentNote', 'documentCitation', 'documentBibliography'}
MARKS = {'bold', 'italic', 'underline', 'strike', 'subscript', 'superscript', 'textStyle', 'highlight', 'link'}
COPY_NOTICE = ('This is an editable copy. Export creates a new DOCX containing the supported content. '
               'The original file is unchanged. Named paragraph styles retain supported formatting; pagination, themes and advanced Word features are not reproduced exactly.')


def number(value, low, high, default):
    if value is None:
        return default
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or not low <= value <= high:
        raise ValueError(f'A document measurement must be between {low} and {high}.')
    return value


def color(value):
    if not isinstance(value, str):
        raise ValueError('Invalid text color.')
    if re.fullmatch(r'#[0-9a-fA-F]{6}', value):
        return value.upper()
    match = re.fullmatch(r'rgb\(\s*(\d{1,3})\s*,\s*(\d{1,3})\s*,\s*(\d{1,3})\s*\)', value)
    if match and all(int(v) <= 255 for v in match.groups()):
        return '#' + ''.join(f'{int(v):02X}' for v in match.groups())
    raise ValueError('Use an RGB or six-digit hexadecimal color.')


def font_size(value):
    match = re.fullmatch(r'(\d+(?:\.\d+)?)(pt|px)', str(value))
    if not match:
        raise ValueError('Font size must use points or pixels.')
    points = float(match[1]) * (0.75 if match[2] == 'px' else 1)
    return number(points, 6, 96, 12)


def safe_link(value):
    return isinstance(value, str) and len(value) <= 2048 and (bool(re.fullmatch(r'(?:https?://|mailto:)[^\s\x00-\x1f<>]+', value, re.I)) or
                                                           (value.startswith('#') and valid_name(value[1:])))


def layout_model(value):
    if not isinstance(value, dict):
        raise ValueError('Invalid page settings.')
    layout = {**DEFAULT_LAYOUT, **value}
    validate_page_details(layout)
    if not isinstance(layout['paper'], str) or layout['paper'] not in PAPERS or layout['orientation'] not in ('portrait', 'landscape'):
        raise ValueError('Unsupported paper size or orientation.')
    width, height = PAPERS[layout['paper']]
    if layout['orientation'] == 'landscape':
        width, height = height, width
    for key in ('top', 'bottom', 'left', 'right'):
        layout[key] = number(layout[key], 0.25, 3, 1)
    if width - layout['left'] - layout['right'] < 1 or height - layout['top'] - layout['bottom'] < 1:
        raise ValueError('Page margins must leave at least one inch of writing space.')
    return layout, width, height


def image_bytes(src):
    if not isinstance(src, str) or len(src) > 12 * 1024**2:
        raise ValueError('Pictures must be local PNG, JPEG or WebP images under 8 MiB.')
    match = re.fullmatch(r'data:image/(png|jpeg|webp);base64,([A-Za-z0-9+/=\r\n]+)', src)
    if not match:
        raise ValueError('Only embedded PNG, JPEG and WebP pictures are supported. Remote images are not loaded.')
    try:
        raw = base64.b64decode(match[2], validate=True)
        if len(raw) > 8 * 1024**2:
            raise ValueError('Picture exceeds 8 MiB.')
        with Image.open(BytesIO(raw)) as image:
            if image.width * image.height > 24_000_000 or image.format not in ('PNG', 'JPEG', 'WEBP'):
                raise ValueError('Pictures must be PNG, JPEG or WebP and at most 24 megapixels.')
            output = BytesIO()
            image.convert('RGBA' if 'A' in image.getbands() else 'RGB').save(output, format='PNG')
            encoded = output.getvalue()
            if len(encoded) > 8 * 1024**2:
                raise ValueError('The normalized picture exceeds 8 MiB. Resize it before inserting it.')
            return encoded, image.width, image.height
    except (OSError, binascii.Error, Image.DecompressionBombError) as exc:
        raise ValueError('The picture is damaged or unsupported.') from exc


def validate_model(doc):
    counts = {'nodes': 0, 'text': 0, 'images': 0}
    children = {'doc': {'paragraph', 'heading', 'bulletList', 'orderedList', 'table', 'image', 'pageBreak', 'tableOfContents', 'documentBibliography'},
                'paragraph': {'text', 'hardBreak', 'bookmark', 'documentNote', 'documentCitation'}, 'heading': {'text', 'hardBreak', 'bookmark', 'documentNote', 'documentCitation'},
                'bulletList': {'listItem'}, 'orderedList': {'listItem'},
                'listItem': {'paragraph', 'heading', 'bulletList', 'orderedList'},
                'table': {'tableRow'}, 'tableRow': {'tableCell', 'tableHeader'},
                'tableCell': {'paragraph', 'heading', 'bulletList', 'orderedList', 'image'},
                'tableHeader': {'paragraph', 'heading', 'bulletList', 'orderedList', 'image'}}

    def visit(node, depth=0):
        if not isinstance(node, dict) or not isinstance(node.get('type'), str) or node['type'] not in SUPPORTED or depth > 20:
            raise ValueError('Unsupported document structure or nesting.')
        counts['nodes'] += 1
        if counts['nodes'] > 20000:
            raise ValueError('Document exceeds 20,000 elements.')
        kind = node['type']
        attrs = node.get('attrs') or {}
        content = node.get('content', [])
        marks = node.get('marks', [])
        if not isinstance(attrs, dict) or not isinstance(content, list) or not isinstance(marks, list):
            raise ValueError('Invalid document attributes.')
        if kind == 'text':
            if not isinstance(node.get('text'), str):
                raise ValueError('Invalid document text.')
            counts['text'] += len(node['text'])
            if counts['text'] > 1_000_000 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', node['text']):
                raise ValueError('Document text is too long or contains unsupported control characters.')
        if kind == 'documentNote':
            validate_note_node(attrs)
            visit(attrs['body'], depth + 1)
        if kind in ('paragraph', 'heading'):
            sid = attrs.get('styleId')
            if sid is not None and (not isinstance(sid, str) or sid not in styles):
                raise ValueError('A paragraph refers to an unknown style.')
            if any(attrs.get(key) is not None and not isinstance(attrs[key], bool) for key in PAGINATION):
                raise ValueError('Paragraph page-break controls must be true, false or inherited.')
            if attrs.get('textAlign') not in (None, *ALIGN):
                raise ValueError('Invalid paragraph alignment.')
            for key, low, high, default in [('indent', 0, 6, 0), ('spaceBefore', 0, 72, 0),
                                             ('spaceAfter', 0, 72, 8), ('lineSpacing', 1, 3, 1.15)]:
                number(attrs.get(key), low, high, default)
        if kind == 'heading' and attrs.get('level', 1) not in (1, 2, 3):
            raise ValueError('Only heading levels 1–3 are supported.')
        if kind == 'orderedList':
            number(attrs.get('start'), 1, 10000, 1)
        if kind == 'table':
            table_grid(node)
        if kind == 'image':
            counts['images'] += 1
            if counts['images'] > 20:
                raise ValueError('A document can contain at most 20 pictures in this phase.')
            image_bytes(attrs.get('src'))
            number(attrs.get('width'), 24, 1600, 480)
            if not isinstance(attrs.get('alt', ''), str) or len(attrs.get('alt', '')) > 2000:
                raise ValueError('Picture description is too long.')
        for mark in marks:
            if not isinstance(mark, dict) or not isinstance(mark.get('type'), str) or mark['type'] not in MARKS:
                raise ValueError('Unsupported text formatting.')
            ma = mark.get('attrs') or {}
            if not isinstance(ma, dict):
                raise ValueError('Invalid formatting attributes.')
            if mark['type'] == 'textStyle':
                if ma.get('fontWeight') not in (None, 'normal', 'bold') or ma.get('fontStyle') not in (None, 'normal', 'italic'):
                    raise ValueError('Invalid direct text emphasis.')
                if ma.get('fontFamily') and ma['fontFamily'] not in FONTS:
                    raise ValueError('Choose a supported font from the ribbon.')
                if ma.get('fontSize'):
                    font_size(ma['fontSize'])
                if ma.get('color'):
                    color(ma['color'])
            if mark['type'] == 'highlight' and ma.get('color'):
                color(ma['color'])
            if mark['type'] == 'link' and not safe_link(ma.get('href')):
                raise ValueError('Links must use http, https, mailto or a valid document bookmark.')
        for child in content:
            if not isinstance(child, dict) or not isinstance(child.get('type'), str) or child['type'] not in children.get(kind, set()):
                raise ValueError('Invalid document nesting.')
            visit(child, depth + 1)
    if not isinstance(doc, dict) or doc.get('type') != 'doc':
        raise ValueError('The editor requires a document.')
    styles = style_catalog(doc)
    visit(doc)
    validate_references(doc)
    validate_notes(doc)
    validate_citations(doc)


def _element(tag, **attrs):
    node = OxmlElement('w:' + tag)
    for key, value in attrs.items():
        node.set(qn('w:' + key), str(value))
    return node


def write_text_run(p, child):
    run = p.add_run(child['text'])
    for mark in child.get('marks', []):
        name, ma = mark['type'], mark.get('attrs') or {}
        if name in ('bold', 'italic', 'underline'):
            setattr(run, name, True)
        elif name in ('strike', 'subscript', 'superscript'):
            setattr(run.font, name, True)
        elif name == 'textStyle':
            if ma.get('fontFamily'): run.font.name = ma['fontFamily']
            if ma.get('fontSize'): run.font.size = Pt(font_size(ma['fontSize']))
            if ma.get('color'): run.font.color.rgb = RGBColor.from_string(color(ma['color'])[1:])
        elif name == 'highlight':
            run._r.get_or_add_rPr().append(_element('shd', val='clear', fill=color(ma.get('color') or '#FFFF00')[1:]))
        elif name == 'link':
            if ma['href'].startswith('#'):
                internal_link(p, run, ma['href'][1:])
            else:
                relation = p.part.relate_to(ma['href'], RT.HYPERLINK, is_external=True)
                link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), relation)
                p._p.remove(run._r); link.append(run._r); p._p.append(link)
                run.font.color.rgb = RGBColor.from_string('0563C1'); run.underline = True
    # Explicit Off overrides inherited emphasis, independent of mark order.
    for mark in child.get('marks', []):
        if mark['type'] != 'textStyle': continue
        ma = mark.get('attrs') or {}
        if ma.get('fontWeight'): run.bold = ma['fontWeight'] == 'bold'
        if ma.get('fontStyle'): run.italic = ma['fontStyle'] == 'italic'


def export_docx(model, layout):
    validate_model(model)
    layout, width, height = layout_model(layout)
    document = Document()
    document.core_properties.author = ''
    document.core_properties.last_modified_by = ''
    document.core_properties.title = ''
    styles = write_styles(document, style_catalog(model))
    references = ReferenceWriter(model, style_catalog(model))
    section = document.sections[0]
    section.page_width, section.page_height = Inches(width), Inches(height)
    from docx.enum.section import WD_ORIENT
    section.orientation = WD_ORIENT.LANDSCAPE if layout['orientation'] == 'landscape' else WD_ORIENT.PORTRAIT
    for key in ('top', 'bottom', 'left', 'right'):
        setattr(section, key + '_margin', Inches(layout[key]))
    usable = width - layout['left'] - layout['right']
    write_page_details(document, layout)
    notes = NoteWriter(document, model, write_text_run)
    citations = CitationWriter(document, model, write_text_run)

    def numbering(kind, start):
        root = document.part.numbering_part.element
        abstract_id = max([int(n.get(qn('w:abstractNumId'))) for n in root.findall(qn('w:abstractNum'))] + [-1]) + 1
        abstract = _element('abstractNum', abstractNumId=abstract_id)
        for level in range(9):
            lvl = _element('lvl', ilvl=level)
            for child in [_element('start', val=start), _element('numFmt', val='bullet' if kind == 'bulletList' else 'decimal'),
                          _element('lvlText', val='•' if kind == 'bulletList' else f'%{level + 1}.'), _element('lvlJc', val='left')]:
                lvl.append(child)
            props = _element('pPr'); props.append(_element('ind', left=(level + 1) * 360, hanging=240)); lvl.append(props)
            abstract.append(lvl)
        # abstractNum must precede num in the numbering schema.
        first_num = root.find(qn('w:num'))
        root.insert(root.index(first_num) if first_num is not None else len(root), abstract)
        return root.add_num(abstract_id).numId

    def paragraph(parent, node, num=None, depth=0):
        attrs = node.get('attrs') or {}
        sid = attrs.get('styleId') or (f'heading-{attrs.get("level", 1)}' if node['type'] == 'heading' else 'normal')
        p = parent.add_paragraph(style=styles[sid])
        if id(node) in references.ids: references.bookmark(p, references.ids[id(node)])
        fmt = p.paragraph_format
        if attrs.get('textAlign') is not None: p.alignment = ALIGN[attrs['textAlign']]
        if attrs.get('spaceBefore') is not None: fmt.space_before = Pt(attrs['spaceBefore'])
        if attrs.get('spaceAfter') is not None: fmt.space_after = Pt(attrs['spaceAfter'])
        if attrs.get('lineSpacing') is not None: fmt.line_spacing = attrs['lineSpacing']
        for key, attribute in PAGINATION.items():
            if attrs.get(key) is not None:
                setattr(fmt, attribute, attrs[key])
        if num is not None:
            np = p._p.get_or_add_pPr().get_or_add_numPr()
            np.get_or_add_ilvl().val, np.get_or_add_numId().val = min(depth, 8), num
        elif attrs.get('indent') is not None:
            fmt.left_indent = Inches(attrs['indent'] * 0.25)
        for child in node.get('content', []):
            if child['type'] == 'bookmark':
                references.bookmark(p, child['attrs']['name']); continue
            if child['type'] == 'hardBreak':
                p.add_run().add_break(); continue
            if child['type'] == 'documentNote':
                notes.reference(p, child); continue
            if child['type'] == 'documentCitation':
                citations.citation(p, child['attrs']); continue
            write_text_run(p, child)
        return p

    def write(parent, nodes, available=usable, depth=0):
        for node in nodes:
            kind, attrs = node['type'], node.get('attrs') or {}
            if kind in ('paragraph', 'heading'):
                paragraph(parent, node)
            elif kind in ('bulletList', 'orderedList'):
                num = numbering(kind, int(attrs.get('start') or 1))
                for item in node.get('content', []):
                    first = True
                    for part in item.get('content', []):
                        if part['type'] in ('paragraph', 'heading'):
                            p = paragraph(parent, part, num if first else None, depth)
                            if not first: p.paragraph_format.left_indent = Inches((depth + 1) * 0.25)
                            first = False
                        else: write(parent, [part], available, depth + 1)
            elif kind == 'pageBreak':
                parent.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            elif kind == 'image':
                raw, _, _ = image_bytes(attrs['src'])
                p = parent.add_paragraph()
                picture = p.add_run().add_picture(BytesIO(raw), width=Inches(min(available, (attrs.get('width') or 480) / 96)))
                picture._inline.docPr.set('descr', attrs.get('alt') or '')
            elif kind == 'table':
                write_table(parent, node, available, lambda cell, content, cell_width: write(cell, content, cell_width, depth))
            elif kind == 'tableOfContents':
                references.contents(parent, attrs)
            elif kind == 'documentBibliography':
                citations.bibliography(parent, attrs)
    write(document, model.get('content', []))
    if not document.paragraphs and not document.tables:
        document.add_paragraph()
    result = BytesIO(); document.save(result)
    if result.tell() > MAX_FILE:
        raise ValueError('The exported document would exceed 24 MiB. Reduce the number or size of pictures.')
    return result.getvalue()


def _checked_package(raw):
    if not raw or len(raw) > MAX_FILE:
        raise ValueError('Open a DOCX file no larger than 24 MiB.')
    try:
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            entries = archive.infolist()
            names = [entry.filename for entry in entries]
            if len(entries) > 2048 or len(set(names)) != len(names) or sum(entry.file_size for entry in entries) > 64 * 1024**2:
                raise ValueError('The DOCX package exceeds the import limits.')
            if 'word/document.xml' not in names or '[Content_Types].xml' not in names:
                raise ValueError('This is not a DOCX document.')
            for entry in entries:
                if entry.flag_bits & 1 or '..' in entry.filename.split('/') or entry.filename.startswith('/') or '\\' in entry.filename:
                    raise ValueError('Unsupported document package.')
                if entry.filename.endswith(('.xml', '.rels')):
                    content = archive.read(entry)
                    tree = etree.fromstring(content, etree.XMLParser(resolve_entities=False, load_dtd=False, no_network=True))
                    if tree.getroottree().docinfo.doctype:
                        raise ValueError('Document XML declarations with DTDs are unsupported.')
            if any('vbaproject' in name.lower() for name in names):
                raise ValueError('Macro-enabled packages are unsupported. Open a plain .docx copy.')
    except (zipfile.BadZipFile, RuntimeError, etree.XMLSyntaxError) as exc:
        raise ValueError('The DOCX file is damaged, encrypted or unsupported.') from exc
    return names


def import_docx(raw):
    names = _checked_package(raw)
    try:
        document = Document(BytesIO(raw))
    except Exception as exc:
        raise ValueError('The DOCX file cannot be read.') from exc
    warnings = [COPY_NOTICE]
    def warn(message):
        if message not in warnings: warnings.append(message)
    if any(re.search(r'word/comments', name) for name in names):
        warn('Comments are not imported in this phase.')
    if document.element.xpath('.//w:ins | .//w:del | .//w:fldChar | .//w:pict | .//w:object | .//w:altChunk'):
        warn('Tracked changes, content controls, fields, legacy drawings and embedded objects may be omitted or reduced to visible text. Review the copy against the original.')
    if len(document.sections) > 1:
        warn('Only the first section’s page settings are used. Section breaks and columns are not retained.')
    section = document.sections[0]
    pw, ph = section.page_width.inches, section.page_height.inches
    landscape = pw > ph
    dims = (ph, pw) if landscape else (pw, ph)
    paper = min(PAPERS, key=lambda name: abs(PAPERS[name][0] - dims[0]) + abs(PAPERS[name][1] - dims[1]))
    if max(abs(PAPERS[paper][i] - dims[i]) for i in (0, 1)) > .03:
        warn('Custom paper size was approximated with the nearest supported paper size.')
    layout = {**read_page_details(document, warn), 'paper': paper, 'orientation': 'landscape' if landscape else 'portrait'}
    for key in ('top', 'bottom', 'left', 'right'):
        value = getattr(section, key + '_margin')
        original = value.inches if value is not None else 1
        layout[key] = round(min(3, max(.25, original)), 3)
        if layout[key] != round(original, 3): warn('Unsupported margins were adjusted to the editor’s range.')
    try: layout_model(layout)
    except ValueError:
        layout = {**layout, 'top': 1, 'bottom': 1, 'left': 1, 'right': 1}
        warn('Margins were reset to one inch to leave usable writing space.')

    style_reader = StyleReader(document, warn)
    references = ReferenceReader(warn)

    def run_marks(run):
        marks = []
        for name in ('bold', 'italic', 'underline', 'strike', 'subscript', 'superscript'):
            if getattr(run.font, name): marks.append({'type': name})
        attrs = {}
        if run.bold is False: attrs['fontWeight'] = 'normal'
        if run.italic is False: attrs['fontStyle'] = 'normal'
        if run.font.name:
            if run.font.name in FONTS: attrs['fontFamily'] = run.font.name
            else: warn('Some fonts are unavailable in the ribbon and use the document default.')
        if run.font.size: attrs['fontSize'] = f'{min(96, max(6, run.font.size.pt)):g}pt'
        if run.font.color.rgb: attrs['color'] = '#' + str(run.font.color.rgb)
        if attrs: marks.append({'type': 'textStyle', 'attrs': attrs})
        shading = run._r.find('w:rPr/w:shd', run._r.nsmap)
        fill = shading.get(qn('w:fill')) if shading is not None else None
        if fill and re.fullmatch('[0-9a-fA-F]{6}', fill):
            marks.append({'type': 'highlight', 'attrs': {'color': '#' + fill}})
        elif run.font.highlight_color:
            palette = {7: '#FFFF00', 4: '#00FF00', 3: '#00FFFF', 5: '#FF00FF', 2: '#0000FF', 6: '#FF0000'}
            marks.append({'type': 'highlight', 'attrs': {'color': palette.get(int(run.font.highlight_color), '#FFFF00')}})
        return marks

    notes = NoteReader(document, warn, run_marks, safe_link)
    citations = CitationReader(document, warn)

    def paragraph(p):
        style = style_reader.read(p.style)
        kind, attrs = 'paragraph', {'styleId': style['id']}
        if style['level']:
            kind, attrs['level'] = 'heading', style['level']
        fmt = p.paragraph_format
        for key, attribute in PAGINATION.items():
            value, ancestor = getattr(fmt, attribute), p.style
            visited = set()
            while value is None and ancestor is not None and ancestor.style_id not in visited:
                visited.add(ancestor.style_id)
                value = getattr(ancestor.paragraph_format, attribute)
                ancestor = ancestor.base_style
            attrs[key] = value
        attrs['textAlign'] = next((key for key, val in ALIGN.items() if val == p.alignment), None)
        attrs['indent'] = min(6, max(0, fmt.left_indent.inches / .25)) if fmt.left_indent is not None else None
        for key, val, default in [('spaceBefore', fmt.space_before, 0), ('spaceAfter', fmt.space_after, 8)]:
            attrs[key] = min(72, max(0, val.pt)) if val is not None else None
        attrs['lineSpacing'] = min(3, max(1, fmt.line_spacing)) if isinstance(fmt.line_spacing, float) else None
        if fmt.line_spacing is not None and not isinstance(fmt.line_spacing, float):
            attrs['lineSpacing'] = 1.15; warn('Fixed or minimum paragraph line spacing was approximated with multiple-line spacing.')
        output, content = [], []
        def flush(force=False):
            if content or force or attrs.get('referenceId'):
                output.append({'type': kind, 'attrs': dict(attrs), 'content': list(content)}); content.clear()
                attrs.pop('referenceId', None)
        for child in citations.inline_items(p._p):
            if isinstance(child, dict): content.append(child); continue
            if child.tag == qn('w:pPr'): continue
            if child.tag == qn('w:bookmarkStart'):
                name = references.bookmark(child)
                if name:
                    if HEADING_ID.fullmatch(name) and not content and not attrs.get('referenceId'):
                        attrs['referenceId'] = name
                    elif HEADING_ID.fullmatch(name):
                        warn('An extra heading destination was omitted. Review internal links.')
                    else: content.append({'type': 'bookmark', 'attrs': {'name': name}})
                continue
            link = None
            if child.tag == qn('w:hyperlink'):
                relation = p.part.rels.get(child.get(qn('r:id')))
                anchor = child.get(qn('w:anchor'))
                if relation and relation.is_external and safe_link(relation.target_ref):
                    link = relation.target_ref
                    if anchor and '#' not in link and safe_link(link + '#' + anchor): link += '#' + anchor
                elif not child.get(qn('r:id')) and valid_name(anchor): link = '#' + anchor
                else: warn('Unsupported hyperlinks were reduced to text.')
                runs = child.findall(qn('w:r'))
            elif child.tag == qn('w:r'): runs = [child]
            else:
                # Never silently present a copy with invisible unhandled structures as complete.
                if child.tag not in (qn('w:bookmarkStart'), qn('w:bookmarkEnd'), qn('w:proofErr')):
                    warn('Some unsupported paragraph content was omitted. Review imported text against the original.')
                continue
            for elem in runs:
                marks = run_marks(Run(elem, p))
                if link: marks.append({'type': 'link', 'attrs': {'href': link}})
                for part in elem:
                    if part.tag in (qn('w:footnoteReference'), qn('w:endnoteReference')):
                        content.append(notes.read(part, 'footnote' if part.tag == qn('w:footnoteReference') else 'endnote'))
                    elif part.tag in (qn('w:t'), qn('w:tab')):
                        text = '\t' if part.tag == qn('w:tab') else part.text or ''
                        if text: content.append({'type': 'text', 'text': text, 'marks': marks})
                    elif part.tag in (qn('w:br'), qn('w:cr')):
                        if part.get(qn('w:type')) == 'page': flush(); output.append({'type': 'pageBreak'})
                        else: content.append({'type': 'hardBreak'})
                    elif part.tag == qn('w:drawing'):
                        flush()
                        if part.xpath('.//wp:anchor'): warn('Floating pictures were converted to pictures between paragraphs; wrapping and positioning are not retained.')
                        for blip in part.xpath('.//a:blip'):
                            image_part = p.part.related_parts.get(blip.get(qn('r:embed')))
                            if image_part is None:
                                warn('Linked pictures are not fetched and were omitted.'); continue
                            source = 'data:' + image_part.content_type + ';base64,' + base64.b64encode(image_part.blob).decode()
                            try: png, iw, _ = image_bytes(source)
                            except ValueError:
                                warn('An oversized or unsupported picture was omitted.'); continue
                            extents = part.xpath('.//wp:extent')
                            width = int(extents[0].get('cx', 0)) / 9525 if extents else iw
                            props = part.xpath('.//wp:docPr')
                            output.append({'type': 'image', 'attrs': {'src': 'data:image/png;base64,' + base64.b64encode(png).decode(),
                                           'width': max(24, min(1600, round(width))), 'alt': props[0].get('descr', '') if props else ''}})
        flush(not output)
        return output

    numbering_root = document.part.numbering_part.element
    def list_info(p):
        props = p._p.pPr
        np = props.numPr if props is not None else None
        if np is None and p.style is not None:
            style_props = p.style.element.pPr
            np = style_props.numPr if style_props is not None else None
        if np is None or np.numId is None or np.numId.val == 0: return None
        level = min(8, np.ilvl.val if np.ilvl is not None else 0)
        num = next((n for n in numbering_root.findall(qn('w:num')) if n.get(qn('w:numId')) == str(np.numId.val)), None)
        if num is None: return None
        abstract_id = num.find(qn('w:abstractNumId')).get(qn('w:val'))
        abstract = next((a for a in numbering_root.findall(qn('w:abstractNum')) if a.get(qn('w:abstractNumId')) == abstract_id), None)
        lvl = next((l for l in abstract.findall(qn('w:lvl')) if l.get(qn('w:ilvl')) == str(level)), None) if abstract is not None else None
        fmt = lvl.find(qn('w:numFmt')) if lvl is not None else None
        start = lvl.find(qn('w:start')) if lvl is not None else None
        form = fmt.get(qn('w:val')) if fmt is not None else 'decimal'
        if form not in ('bullet', 'decimal'): warn('Custom numbering was converted to decimal numbering.')
        return ('bulletList' if form == 'bullet' else 'orderedList', level, np.numId.val,
                max(1, min(10000, int(start.get(qn('w:val'), '1')))) if start is not None else 1)

    def inner_content(parent, container=None, depth=0):
        if depth > 20: raise ValueError('Unsupported content-control nesting.')
        if container is None:
            container = document.element.body if parent is document else parent._element
        for child in container:
            if child.tag == qn('w:p'): yield Paragraph(child, parent)
            elif child.tag == qn('w:tbl'): yield Table(child, parent)
            elif child.tag == qn('w:sdt'):
                bibliography = citations.tagged(child, block=parent is document and depth == 0)
                if bibliography and bibliography['type'] == 'documentBibliography': yield bibliography; continue
                toc = references.contents(child, parent is document and depth == 0)
                if toc: yield toc
                else:
                    warn('Content controls and external tables of contents were reduced to supported visible content; their live fields are not retained.')
                    content = child.find(qn('w:sdtContent'))
                    if content is not None: yield from inner_content(parent, content, depth + 1)
            elif child.tag == qn('w:bookmarkStart'):
                warn('Bookmarks outside text paragraphs were omitted. Review internal links.')

    def blocks(parent):
        result, stack = [], []
        for obj in inner_content(parent):
            if isinstance(obj, Paragraph):
                parts, info = paragraph(obj), list_info(obj)
                if info and all(part['type'] in ('paragraph', 'heading') for part in parts):
                    if parts and parts[0]['type'] == 'heading':
                        parts[0]['type'] = 'paragraph'; parts[0]['attrs'].pop('level', None)
                    kind, level, identifier, start = info
                    level = min(level, len(stack))
                    stack = stack[:level + 1]
                    if len(stack) <= level or stack[level][0] != (kind, identifier):
                        stack = stack[:level]
                        target = result if not stack else stack[-1][1]['content'][-1]['content']
                        node = {'type': kind, 'attrs': {'start': start} if kind == 'orderedList' else {}, 'content': []}
                        target.append(node); stack.append(((kind, identifier), node))
                    stack[-1][1]['content'].append({'type': 'listItem', 'content': parts})
                else:
                    stack = []; result.extend(parts)
            elif isinstance(obj, Table):
                stack = []
                result.append(read_table(obj, blocks, warn))
            elif isinstance(obj, dict):
                stack = []; result.append(obj)
        return result
    if len(document.element.xpath('.//w:p | .//w:r | .//w:tc')) > 20000:
        raise ValueError('Document exceeds 20,000 paragraphs, runs or cells.')
    content = blocks(document) or [{'type': 'paragraph'}]
    notes.finish()
    model = {'type': 'doc', 'attrs': {'styles': list(style_reader.catalog.values()), 'noteSettings': notes.settings, 'sources': citations.sources, 'citationStyle': citations.style}, 'content': content}
    citations.finish(model)
    validate_model(model)
    if missing_links(model): warn('Some internal links have missing destinations. Use References → Check internal links to review them.')
    return {'document': model, 'layout': layout, 'warnings': warnings}
