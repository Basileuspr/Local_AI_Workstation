"""Numbered captions and local cross-references; locked SEQ/REF snapshots in DOCX."""
import json
import re
from uuid import uuid4
from lxml import etree
from docx.opc.part import XmlPart
from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from services.document_references import element, walk, valid_name
from services.document_citations import printable

CAPTION_ID = re.compile(r'LAW_C[0-9a-f]{32}')
LABEL = re.compile(r'[A-Za-z][A-Za-z0-9 _-]{0,29}')
FORMATS = {'decimal':'ARABIC', 'upperRoman':'ROMAN', 'lowerRoman':'roman', 'upperLetter':'ALPHABETIC', 'lowerLetter':'alphabetic'}
CAPTION_TAG, REFERENCE_TAG = 'LAW_CAPTION_V1:', 'LAW_XREF_V1:'
NAMESPACE = 'urn:local-ai-workstation:document-captions:v1'


def caption_attrs(attrs):
    if not isinstance(attrs.get('id'), str) or not CAPTION_ID.fullmatch(attrs['id']): raise ValueError('Invalid caption ID.')
    label = attrs.get('label', 'Figure')
    if not isinstance(label, str) or not LABEL.fullmatch(label) or label != label.strip(): raise ValueError('Invalid caption label.')
    if not printable(attrs.get('text', ''), 1000): raise ValueError('Caption text supports up to 1,000 printable characters.')
    return {'id': attrs['id'], 'label': label, 'text': attrs.get('text', '')}


def reference_attrs(attrs):
    if not isinstance(attrs.get('id'), str) or not re.fullmatch(r'xref-[0-9a-f]{32}', attrs['id']): raise ValueError('Invalid cross-reference ID.')
    if not valid_name(attrs.get('target')) or attrs.get('display', 'full') not in ('full', 'labelNumber', 'number', 'text') or type(attrs.get('hyperlink', True)) is not bool: raise ValueError('Invalid cross-reference destination or display.')
    return {'id': attrs['id'], 'target': attrs['target'], 'display': attrs.get('display', 'full'), 'hyperlink': attrs.get('hyperlink', True)}


def caption_settings(model):
    settings = (model.get('attrs') or {}).get('captionSettings', {})
    if not isinstance(settings, dict) or len(settings) > 20: raise ValueError('Use up to 20 caption labels.')
    for label, rule in settings.items():
        if not isinstance(label, str) or not LABEL.fullmatch(label) or label != label.strip() or not isinstance(rule, dict) or not isinstance(rule.get('format'), str) or rule['format'] not in FORMATS or type(rule.get('start')) is not int or not 1 <= rule['start'] <= 9999: raise ValueError('Invalid caption numbering; use a supported format and start from 1–9,999.')
    return {label: {'format': rule['format'], 'start': rule['start']} for label, rule in settings.items()}


def validate_captions(model):
    settings = caption_settings(model); labels, ids, references, caption_count = set(settings), set(), 0, 0
    reserved = {node['attrs']['name'].lower() for node in walk(model) if node['type'] == 'bookmark'}
    for node in walk(model):
        if node['type'] not in ('documentCaption', 'documentCrossReference'): continue
        attrs = caption_attrs(node.get('attrs') or {}) if node['type'] == 'documentCaption' else reference_attrs(node.get('attrs') or {})
        if attrs['id'] in ids or attrs['id'].lower() in reserved: raise ValueError('Caption and cross-reference IDs must be unique and not collide with bookmarks.')
        ids.add(attrs['id'])
        if node['type'] == 'documentCaption': labels.add(attrs['label']); caption_count += 1
        else: references += 1
    if len(labels) > 20 or len({label.lower() for label in labels}) != len(labels): raise ValueError('Use up to 20 different caption labels.')
    if caption_count > 200 or references > 1000: raise ValueError('Use up to 200 captions and 1,000 cross-references.')


def caption_number(number, format):
    if format.endswith('Letter'):
        text = ''
        while number: number -= 1; text = chr(65 + number % 26) + text; number //= 26
        return text.lower() if format == 'lowerLetter' else text
    if format.endswith('Roman'):
        text = ''
        for value, symbol in [(1000,'M'),(900,'CM'),(500,'D'),(400,'CD'),(100,'C'),(90,'XC'),(50,'L'),(40,'XL'),(10,'X'),(9,'IX'),(5,'V'),(4,'IV'),(1,'I')]:
            while number >= value: text += symbol; number -= value
        return text.lower() if format == 'lowerRoman' else text
    return str(number)


def caption_context(model):
    counters, captions = {}, {}; settings = caption_settings(model)
    for node in walk(model):
        if node['type'] != 'documentCaption': continue
        attrs = caption_attrs(node['attrs']); key = attrs['label'].lower(); rule = settings.get(attrs['label'], {'format':'decimal', 'start':1})
        ordinal = counters.get(key, rule['start'] - 1) + 1; counters[key] = ordinal
        number = caption_number(ordinal, rule['format']); label_number = attrs['label'] + ' ' + number
        captions[attrs['id']] = {**attrs, 'ordinal':ordinal, 'number':number, 'labelNumber':label_number, 'full':label_number + (': ' + attrs['text'] if attrs['text'] else '')}
    return captions


def reference_targets(model):
    targets = {}
    for node in walk(model):
        attrs = node.get('attrs') or {}
        if node['type'] in ('paragraph', 'heading') and attrs.get('referenceId'):
            text = ''.join(child.get('text', '') for child in node.get('content', [])) or 'Untitled heading'; targets[attrs['referenceId']] = {'full': text}
        if node['type'] == 'bookmark': targets[attrs['name']] = {'full':attrs['name']}
    targets.update(caption_context(model)); return targets


def reference_text(attrs, targets):
    item = targets.get(attrs['target'])
    return item.get(attrs.get('display', 'full'), item['full']) if item else '[Missing reference: ' + attrs['target'] + ']'


def control(tag, attrs):
    sdt, props, content = element('sdt'), element('sdtPr'), element('sdtContent')
    props.append(element('tag', val=tag + json.dumps(attrs, ensure_ascii=False, separators=(',', ':')))); sdt.extend([props,content]); return sdt, content


class CaptionWriter:
    def __init__(self, document, model, references):
        self.references, self.settings, self.captions, self.targets, self.seen = references, caption_settings(model), caption_context(model), reference_targets(model), set()
        if self.settings:
            root = etree.Element('{' + NAMESPACE + '}numbering'); root.text = json.dumps({'version':1, 'settings':self.settings})
            package = document.part.package; part = XmlPart(package.next_partname('/customXml/item%d.xml'), CT.XML, root, package); document.part.relate_to(part, RT.CUSTOM_XML)

    def caption(self, parent, attrs):
        attrs = caption_attrs(attrs); item = self.captions[attrs['id']]; sdt, content = control(CAPTION_TAG, attrs)
        p = parent.add_paragraph(style='Caption'); p._p.addprevious(sdt); content.append(p._p); p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(8)
        self.references.serial += 1; serial = self.references.serial; p._p.append(element('bookmarkStart', id=serial, name=attrs['id']))
        p.add_run(attrs['label'] + ' ')
        rule = self.settings.get(attrs['label'], {'format':'decimal', 'start':1})
        instruction = ' SEQ "' + attrs['label'] + '" \\* ' + FORMATS[rule['format']]
        if attrs['label'] not in self.seen: instruction += ' \\r ' + str(rule['start']); self.seen.add(attrs['label'])
        field = element('fldSimple', instr=instruction + ' ', fldLock='1'); p._p.append(field)
        run = p.add_run(item['number']); field.append(run._r)
        if attrs['text']: p.add_run(': ' + attrs['text'])
        p._p.append(element('bookmarkEnd', id=serial))

    def reference(self, paragraph, attrs):
        attrs = reference_attrs(attrs); sdt, content = control(REFERENCE_TAG, attrs); paragraph._p.append(sdt)
        field = element('fldSimple', instr=' REF ' + attrs['target'] + (' \\h ' if attrs['hyperlink'] else ' '), fldLock='1'); content.append(field)
        run = paragraph.add_run(reference_text(attrs, self.targets)); field.append(run._r)
        if attrs['hyperlink']: run.font.color.rgb = RGBColor.from_string('0563C1'); run.underline = True


class CaptionReader:
    def __init__(self, document, warn):
        self.warn, self.settings, self.ids = warn, {}, set()
        for rel in document.part.rels.values():
            if rel.reltype != RT.CUSTOM_XML or rel.is_external or len(rel.target_part.blob) > 1024**2: continue
            try:
                root = etree.fromstring(rel.target_part.blob, etree.XMLParser(resolve_entities=False, load_dtd=False, no_network=True))
                if root.tag != '{' + NAMESPACE + '}numbering': continue
                value = json.loads(root.text or '')
                if not isinstance(value, dict) or value.get('version') != 1: raise ValueError('Invalid numbering')
                self.settings = caption_settings({'attrs':{'captionSettings':value.get('settings')}})
            except (ValueError, TypeError, etree.XMLSyntaxError): warn('Invalid caption numbering metadata was omitted.')

    def tagged(self, child, block=False):
        if child.tag != qn('w:sdt'): return None
        tag = child.find('w:sdtPr/w:tag', child.nsmap); value = tag.get(qn('w:val'), '') if tag is not None else ''
        prefix = CAPTION_TAG if block else REFERENCE_TAG
        if not value.startswith(prefix): return None
        try:
            if len(value) > 4096: raise ValueError('Oversized metadata')
            attrs = json.loads(value[len(prefix):])
            if not isinstance(attrs, dict): raise ValueError('Invalid attributes')
            attrs = caption_attrs(attrs) if block else reference_attrs(attrs)
            if block:
                # Retain supported text/label edits made to our exported
                # caption in another DOCX reader. The SEQ result remains local.
                paragraph = child.find('w:sdtContent/w:p', child.nsmap)
                if paragraph is not None:
                    field = next((part for part in paragraph if part.tag == qn('w:fldSimple') and re.match(r'\s*SEQ\s', part.get(qn('w:instr'), ''), re.I)), None)
                    if field is not None:
                        parts = list(paragraph); offset = parts.index(field)
                        label = ''.join(text.text or '' for part in parts[:offset] for text in part.iter(qn('w:t'))).strip()
                        tail = ''.join(text.text or '' for part in parts[offset + 1:] for text in part.iter(qn('w:t')))
                        if LABEL.fullmatch(label) and (not tail or tail.startswith(':')):
                            value = caption_attrs({**attrs, 'label':label, 'text':tail[1:].lstrip() if tail else ''})
                            if value != attrs: self.warn('Native caption label/text edits were imported; numbering and cross-references were regenerated locally.')
                            attrs = value
                    else: self.warn('Caption metadata was retained because its native SEQ structure was changed. Review caption text against the original.')
                if attrs['id'] in self.ids: attrs['id'] = 'LAW_C' + uuid4().hex; self.warn('A copied caption received an independent ID; existing cross-references retain their destination.')
                self.ids.add(attrs['id'])
            else: attrs['id'] = 'xref-' + uuid4().hex
            return {'type':'documentCaption' if block else 'documentCrossReference', 'attrs':attrs}
        except ValueError: self.warn('Invalid caption/cross-reference metadata became supported visible text.'); return None

    def field(self, instruction):
        if not re.match(r'\s*REF\s', instruction or '', re.I) or len(instruction) > 4096: return None
        match = re.match(r'\s*REF\s+("[^"]+"|[^\s]+)', instruction, re.I)
        target = match.group(1).strip('"') if match else ''
        if not valid_name(target): return None
        self.warn('Native REF fields were normalized to full-text local cross-references. Page/relative-position and other field switches are not reproduced.')
        return {'type':'documentCrossReference', 'attrs':{'id':'xref-' + uuid4().hex, 'target':target, 'display':'full', 'hyperlink': bool(re.search(r'\\h(?:\s|$)', instruction, re.I))}}

    def finish(self, model):
        targets = reference_targets(model)
        if any(node['type'] == 'documentCrossReference' and node['attrs']['target'] not in targets for node in walk(model)): self.warn('Some cross-references have missing destinations. Use References → Check cross-references to repair them.')
