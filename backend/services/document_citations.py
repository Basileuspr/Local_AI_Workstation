"""Local source catalogs, deterministic citations and native DOCX source storage.

Citation fields have readable cached results and are locked snapshots. Tagged
controls restore live behavior here. No Office process or external fetch runs.
"""
import json
import re
from copy import deepcopy
from uuid import uuid4
from lxml import etree
from docx.oxml.ns import qn
from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
from docx.opc.part import XmlPart
from docx.shared import Inches, Pt
from services.document_references import element, walk

B = 'http://schemas.microsoft.com/office/word/2004/10/bibliography'
LAW = 'urn:local-ai-workstation:document-citations:v1'
CITE_TAG, BIB_TAG = 'LAW_CITATION_V1:', 'LAW_BIBLIOGRAPHY_V1:'
FIELDS = {'title': 500, 'year': 4, 'publisher': 200, 'journal': 200, 'volume': 40, 'issue': 40, 'pages': 80, 'siteName': 200, 'url': 2048, 'doi': 200, 'edition': 80}
NATIVE_FIELDS = {'title': 'Title', 'year': 'Year', 'publisher': 'Publisher', 'journal': 'JournalName', 'volume': 'Volume', 'issue': 'Issue', 'pages': 'Pages', 'siteName': 'InternetSiteTitle', 'url': 'URL', 'doi': 'DOI', 'edition': 'Edition'}
SOURCE_ID = re.compile(r'LAW_S[0-9a-f]{32}')


def printable(value, limit):
    if not isinstance(value, str) or re.search(r'[\x00-\x1f\ud800-\udfff]', value): return False
    return len(value.encode('utf-16-le')) // 2 <= limit


def canonical_source(source):
    if not isinstance(source, dict) or not isinstance(source.get('id'), str) or not SOURCE_ID.fullmatch(source['id']) or source.get('type') not in ('book', 'article', 'website'):
        raise ValueError('Invalid source ID or type.')
    for key, maximum in FIELDS.items():
        if not printable(source.get(key, ''), maximum): raise ValueError('Source fields exceed their printable text limits.')
    if not source.get('title', '').strip(): raise ValueError('Enter a source title.')
    if source.get('year') and not re.fullmatch(r'[0-9]{4}', source['year']): raise ValueError('Source years must contain four digits.')
    if source.get('url') and not re.fullmatch(r'https?://[^\s<>]+', source['url'], re.I): raise ValueError('Source URLs must use http or https.')
    if source.get('doi') and not re.fullmatch(r'10\.\d{4,9}/[^\s<>]+', source['doi']): raise ValueError('Invalid source DOI.')
    authors = source.get('authors')
    if not isinstance(authors, list) or len(authors) > 20: raise ValueError('Sources support up to 20 authors.')
    people = []
    for author in authors:
        if not isinstance(author, dict) or not all(printable(author.get(key, ''), 120) for key in ('family', 'given', 'literal')): raise ValueError('Invalid source author name.')
        if not (author.get('literal') or author.get('family') or '').strip() or (author.get('literal') and (author.get('family') or author.get('given'))): raise ValueError('Enter a family or organization name for each author.')
        people.append({'literal': author['literal'].strip()} if author.get('literal') else {'family': author['family'].strip(), 'given': author.get('given', '').strip()})
    return {'id': source['id'], 'type': source['type'], 'authors': people, **{key: source.get(key, '').strip() for key in FIELDS}}


def source_catalog(model):
    sources = (model.get('attrs') or {}).get('sources', [])
    if not isinstance(sources, list) or len(sources) > 100: raise ValueError('A document supports up to 100 sources.')
    result = [canonical_source(source) for source in sources]
    if len({source['id'] for source in result}) != len(result): raise ValueError('Source IDs must be unique.')
    return result


def citation_attrs(attrs, require_id=True):
    ids = attrs.get('sourceIds')
    if not isinstance(ids, list) or not 1 <= len(ids) <= 10 or any(not isinstance(value, str) or not SOURCE_ID.fullmatch(value) for value in ids) or len(set(ids)) != len(ids): raise ValueError('A citation needs 1–10 different source IDs.')
    if require_id and (not isinstance(attrs.get('id'), str) or not re.fullmatch(r'cite-[0-9a-f]{32}', attrs['id'])): raise ValueError('Invalid citation ID.')
    if attrs.get('mode', 'parenthetical') not in ('parenthetical', 'narrative'): raise ValueError('Invalid citation display mode.')
    for key in ('locator', 'prefix', 'suffix'):
        if not printable(attrs.get(key, ''), 120): raise ValueError('Citation details exceed 120 printable characters.')
    return {'id': attrs.get('id') or 'cite-' + uuid4().hex, 'sourceIds': ids, 'mode': attrs.get('mode', 'parenthetical'), **{key: attrs.get(key, '') for key in ('locator', 'prefix', 'suffix')}}


def bibliography_attrs(attrs):
    title = attrs.get('title', 'References'); include = attrs.get('includeUncited', False)
    if not printable(title, 100) or not title.strip() or type(include) is not bool: raise ValueError('Invalid bibliography title or source selection.')
    return {'title': title.strip(), 'includeUncited': include}


def validate_citations(model):
    source_catalog(model)
    if (model.get('attrs') or {}).get('citationStyle', 'apa') not in ('apa', 'numeric'): raise ValueError('Unsupported citation style.')
    ids, bibliography = set(), 0
    for node in walk(model):
        if node['type'] == 'documentCitation':
            attrs = citation_attrs(node.get('attrs') or {})
            if attrs['id'] in ids: raise ValueError('Citation IDs must be unique.')
            ids.add(attrs['id'])
        if node['type'] == 'documentBibliography': bibliography_attrs(node.get('attrs') or {}); bibliography += 1
    if len(ids) > 1000 or bibliography > 1: raise ValueError('A document supports 1,000 citations and one bibliography.')


def source_authors(source, reference=False, narrative=False):
    def name(author):
        if author.get('literal'): return author['literal']
        if not reference: return author['family']
        initials = ' '.join(word[0].upper() + '.' for word in re.findall(r'[^\W_]+', author.get('given', ''), re.U))
        return author['family'] + (', ' + initials if initials else '')
    names = [name(author) for author in source['authors']]
    if not names: return source['title']
    if not reference and len(names) >= 3: return names[0] + ' et al.'
    if len(names) == 1: return names[0]
    return ', '.join(names[:-1]) + (', & ' if reference else ' and ' if narrative else ' & ') + names[-1]


def context(model):
    def key(source): return '|'.join([';'.join((author.get('literal') or author.get('family', '')) + ',' + author.get('given', '') for author in source['authors']) or source['title'], source['year'], source['title']]).lower()
    sources = sorted(source_catalog(model), key=lambda source: (key(source), source['id']))
    nodes = list(walk(model)); citations = [node['attrs'] for node in nodes if node['type'] == 'documentCitation']
    used = {identifier for item in citations for identifier in item['sourceIds']}
    include = any(node['type'] == 'documentBibliography' and node.get('attrs', {}).get('includeUncited') for node in nodes)
    groups, years, numbers = {}, {}, {}
    for source in sources:
        if source['authors'] and (include or source['id'] in used): groups.setdefault(json.dumps([source['authors'], source['year']], sort_keys=True), []).append(source)
    for group in groups.values():
        for index, source in enumerate(group):
            ordinal, suffix = index + 1, ''
            while ordinal: ordinal -= 1; suffix = chr(97 + ordinal % 26) + suffix; ordinal //= 26
            years[source['id']] = (source['year'] or 'n.d.') + ((('' if source['year'] else '-') + suffix) if len(group) > 1 else '')
    for item in citations:
        for identifier in item['sourceIds']:
            if identifier not in numbers: numbers[identifier] = len(numbers) + 1
    for source in sources:
        if source['id'] not in numbers: numbers[source['id']] = len(numbers) + 1
    return {'sources': sources, 'map': {source['id']: source for source in sources}, 'years': years, 'numbers': numbers, 'used': used, 'style': (model.get('attrs') or {}).get('citationStyle', 'apa')}


def citation_label(attrs, info):
    locator = ', ' + attrs['locator'] if attrs.get('locator') else ''
    if info['style'] == 'numeric': label = '[' + ', '.join(str(info['numbers'][identifier]) if identifier in info['map'] else 'Missing source' for identifier in attrs['sourceIds']) + locator + ']'
    else:
        values = []
        for identifier in attrs['sourceIds']:
            source = info['map'].get(identifier)
            values.append([source_authors(source, narrative=attrs.get('mode') == 'narrative'), info['years'].get(identifier) or source['year'] or 'n.d.'] if source else ['Missing source', ''])
        label = values[0][0] + ' (' + values[0][1] + locator + ')' if attrs.get('mode') == 'narrative' and len(values) == 1 else '(' + '; '.join(', '.join(value for value in pair if value) for pair in values) + locator + ')'
    return (attrs['prefix'] + ' ' if attrs.get('prefix') else '') + label + (' ' + attrs['suffix'] if attrs.get('suffix') else '')


def bibliography_entries(model, include_uncited=False):
    info = context(model); sources = [source for source in info['sources'] if include_uncited or source['id'] in info['used']]
    if info['style'] == 'numeric': sources.sort(key=lambda source: info['numbers'][source['id']])
    result = []
    for source in sources:
        parts = []
        def add(text, italic=False, href=None):
            if text: parts.append({'text': text, 'italic': italic, 'href': href})
        names = source_authors(source, reference=True) if source['authors'] else ''
        authors = names + (' ' if names.endswith('.') else '. ') if names else ''
        if info['style'] == 'numeric': add('[' + str(info['numbers'][source['id']]) + '] ')
        add(authors); year = info['years'].get(source['id']) or source['year'] or 'n.d.'
        if info['style'] == 'apa' and authors: add('(' + year + '). ')
        add(source['title'], source['type'] != 'article'); add('. ')
        if info['style'] == 'apa' and not authors: add('(' + year + '). ')
        if source['type'] == 'book':
            if source['edition']: add('(' + source['edition'] + '). ')
            if source['publisher']: add(source['publisher'] + '. ')
        if source['type'] == 'article':
            add(source['journal'], True)
            if source['volume']: add(', '); add(source['volume'], True)
            if source['issue']: add('(' + source['issue'] + ')')
            if source['pages']: add(', ' + source['pages'])
            if source['journal'] or source['volume'] or source['pages']: add('. ')
        if source['type'] == 'website' and source['siteName'] and source['siteName'] != (source['authors'][0].get('literal') if source['authors'] else None): add(source['siteName'] + '. ')
        if info['style'] == 'numeric': add(year + '. ')
        href = 'https://doi.org/' + source['doi'] if source['doi'] else source['url']
        if href: add(href, href=href)
        elif parts: parts[-1]['text'] = parts[-1]['text'].rstrip()
        result.append({'source': source, 'parts': parts})
    return result


def b_element(tag, text=None):
    node = etree.Element('{' + B + '}' + tag)
    if text is not None: node.text = text
    return node


class CitationWriter:
    def __init__(self, document, model, write_run):
        self.model, self.write_run, self.info = model, write_run, context(model)
        if not self.info['sources'] and not any(node['type'] in ('documentCitation', 'documentBibliography') for node in walk(model)): return
        root = etree.Element('{' + B + '}Sources', nsmap={'b': B})
        if self.info['style'] == 'apa': root.set('SelectedStyle', '/APA.XSL'); root.set('StyleName', 'APA')
        for source in self.info['sources']:
            item = b_element('Source'); root.append(item); item.append(b_element('Tag', source['id'])); item.append(b_element('SourceType', {'book': 'Book', 'article': 'JournalArticle', 'website': 'InternetSite'}[source['type']]))
            for key, native in NATIVE_FIELDS.items():
                if source[key]: item.append(b_element(native, source[key]))
            if source['authors']:
                group, authors, names = b_element('Author'), b_element('Author'), b_element('NameList'); item.append(group); group.append(authors)
                if len(source['authors']) == 1 and source['authors'][0].get('literal'): authors.append(b_element('Corporate', source['authors'][0]['literal']))
                else:
                    authors.append(names)
                    for author in source['authors']:
                        person = b_element('Person'); person.append(b_element('Last', author.get('literal') or author['family'])); person.append(b_element('First', author.get('given', ''))); names.append(person)
        package = document.part.package
        part = XmlPart(package.next_partname('/customXml/item%d.xml'), CT.XML, root, package); document.part.relate_to(part, RT.CUSTOM_XML)
        metadata = etree.Element('{' + LAW + '}catalog', nsmap={'law': LAW}); metadata.text = json.dumps({'version': 1, 'sources': self.info['sources'], 'style': self.info['style']}, ensure_ascii=False)
        part = XmlPart(package.next_partname('/customXml/item%d.xml'), CT.XML, metadata, package); document.part.relate_to(part, RT.CUSTOM_XML)

    def citation(self, paragraph, attrs):
        attrs = citation_attrs(attrs); sdt, props, content = element('sdt'), element('sdtPr'), element('sdtContent')
        props.append(element('tag', val=CITE_TAG + json.dumps(attrs, ensure_ascii=False, separators=(',', ':')))); sdt.extend([props, content]); paragraph._p.append(sdt)
        instruction = ' CITATION ' + attrs['sourceIds'][0] + ''.join(' \\m ' + identifier for identifier in attrs['sourceIds'][1:])
        field = element('fldSimple', instr=instruction + ' ', fldLock='1'); content.append(field)
        run = paragraph.add_run(citation_label(attrs, self.info)); paragraph._p.remove(run._r); field.append(run._r)

    def bibliography(self, parent, attrs):
        attrs = bibliography_attrs(attrs); sdt, props, content = element('sdt'), element('sdtPr'), element('sdtContent')
        props.append(element('alias', val=attrs['title'])); props.append(element('tag', val=BIB_TAG + json.dumps(attrs, separators=(',', ':')))); sdt.extend([props, content])
        placeholder = parent.add_paragraph(); placeholder._p.addprevious(sdt); placeholder._p.getparent().remove(placeholder._p)
        def paragraph():
            p = parent.add_paragraph(); content.append(p._p); return p
        title = paragraph(); title.paragraph_format.keep_with_next = True; title.add_run(attrs['title']).bold = True
        for entry in bibliography_entries(self.model, attrs['includeUncited']):
            p = paragraph(); p.paragraph_format.left_indent = Inches(.5); p.paragraph_format.first_line_indent = Inches(-.5); p.paragraph_format.space_after = Pt(8)
            for part in entry['parts']:
                marks = ([{'type': 'italic'}] if part['italic'] else []) + ([{'type': 'link', 'attrs': {'href': part['href']}}] if part['href'] else [])
                self.write_run(p, {'type': 'text', 'text': part['text'], 'marks': marks})


class CitationReader:
    def __init__(self, document, warn, extra_reader=None):
        self.extra_reader = extra_reader
        self.warn, self.sources, self.tags, self.style, self.bibliography_seen = warn, [], {}, 'apa', False
        native_roots, metadata = [], None
        for rel in document.part.rels.values():
            if rel.reltype != RT.CUSTOM_XML: continue
            if rel.is_external: warn('External custom source parts are not fetched.'); continue
            raw = rel.target_part.blob
            if len(raw) > 1024 ** 2: continue
            try: root = etree.fromstring(raw, etree.XMLParser(resolve_entities=False, load_dtd=False, no_network=True))
            except etree.XMLSyntaxError: continue
            if root.tag == '{' + B + '}Sources': native_roots.append(root)
            if root.tag == '{' + LAW + '}catalog':
                try:
                    value = json.loads(root.text or '')
                    if not isinstance(value, dict) or value.get('version') != 1 or value.get('style') not in ('apa', 'numeric'): raise ValueError('Invalid catalog')
                    sources = source_catalog({'attrs': {'sources': value.get('sources')}}); metadata = (sources, value['style'])
                except (ValueError, TypeError): warn('Invalid editor source metadata was ignored; native source data was used where available.')
        if metadata: self.sources, self.style = metadata
        for root in native_roots:
            for node in root.findall('{' + B + '}Source'):
                tag = node.findtext('{' + B + '}Tag', '')
                if not tag or len(tag) > 100 or tag in self.tags: warn('Missing or duplicate source tags were omitted.'); continue
                existing = next((source for source in self.sources if source['id'] == tag), None)
                native_type = node.findtext('{' + B + '}SourceType')
                if native_type not in ('Book', 'JournalArticle', 'InternetSite'): warn('An unsupported source type was normalized to Book. Review its source fields.')
                source = {'id': tag if SOURCE_ID.fullmatch(tag) else 'LAW_S' + uuid4().hex, 'type': {'Book': 'book', 'JournalArticle': 'article', 'InternetSite': 'website'}.get(native_type, 'book'), 'authors': [], **{key: node.findtext('{' + B + '}' + native, '') for key, native in NATIVE_FIELDS.items()}}
                for author in node.findall('.//{' + B + '}Person'):
                    source['authors'].append({'family': author.findtext('{' + B + '}Last', ''), 'given': ' '.join(value for value in [author.findtext('{' + B + '}First', ''), author.findtext('{' + B + '}Middle', '')] if value)})
                corporate = node.find('.//{' + B + '}Corporate')
                if corporate is not None and corporate.text: source['authors'] = [{'literal': corporate.text}]
                if not source['title']: source['title'] = '[Untitled source: ' + tag[:60] + ']'; warn('A source with no title received a visible placeholder.')
                try: source = canonical_source(source)
                except ValueError: warn('An invalid or unsupported source record was omitted. Review imported citations.'); continue
                if existing:
                    # Mixed organization/person lists have a simpler native
                    # representation. Preserve our richer names unless the
                    # native author data was changed by another editor.
                    encoded = existing['authors'] if len(existing['authors']) == 1 and existing['authors'][0].get('literal') else [
                        {'family': author.get('literal') or author['family'], 'given': author.get('given', '')} for author in existing['authors']]
                    if source['authors'] == encoded: source['authors'] = existing['authors']
                    if source != existing: warn('Native source edits were imported. Citations and bibliography were regenerated with the editor’s basic format.')
                    self.sources[self.sources.index(existing)] = source; self.tags[tag] = source['id']; continue
                if len(self.sources) >= 100: raise ValueError('This document contains more than 100 bibliography sources.')
                self.sources.append(source); self.tags[tag] = source['id']
        for source in self.sources: self.tags[source['id']] = source['id']

    def tagged(self, child, block=False):
        if child.tag != qn('w:sdt'): return None
        tag = child.find('w:sdtPr/w:tag', child.nsmap); value = tag.get(qn('w:val'), '') if tag is not None else ''
        prefix = BIB_TAG if block else CITE_TAG
        if not value.startswith(prefix): return None
        try:
            if len(value) > 4096: raise ValueError('Tag is too long')
            attrs = json.loads(value[len(prefix):])
            if not isinstance(attrs, dict): raise ValueError('Invalid attributes')
            if block:
                if self.bibliography_seen: self.warn('An additional bibliography became visible text.'); return None
                attrs = bibliography_attrs(attrs); self.bibliography_seen = True
            else: attrs = citation_attrs(attrs); attrs['id'] = 'cite-' + uuid4().hex
            return {'type': 'documentBibliography' if block else 'documentCitation', 'attrs': attrs}
        except ValueError: self.warn('Invalid citation/bibliography metadata became visible text.'); return None

    def field(self, instruction):
        if not re.match(r'\s*CITATION\s', instruction or '', re.I): return None
        if len(instruction) > 4096: self.warn('An oversized citation field became visible text.'); return None
        tokens = re.findall(r'"[^"]*"|[^\s]+', instruction.strip()); tags = [tokens[1].strip('"')] if len(tokens) > 1 else []
        locator = ''
        for index, token in enumerate(tokens):
            if index + 1 >= len(tokens): continue
            if token.lower() == '\\m': tags.append(tokens[index + 1].strip('"'))
            if token.lower() == '\\p': locator = tokens[index + 1].strip('"')[:120]
        ids = []
        for tag in tags[:10]:
            if tag not in self.tags:
                identifier = tag if SOURCE_ID.fullmatch(tag) else 'LAW_S' + uuid4().hex; self.tags[tag] = identifier
                self.warn('A citation source is missing. Use Manage sources and Edit citation to repair it.')
            if self.tags[tag] not in ids: ids.append(self.tags[tag])
        if not ids: return None
        self.warn('Native citation fields were normalized to the editor’s supported source types and basic formats. Review citation display and locators.')
        return {'type': 'documentCitation', 'attrs': citation_attrs({'sourceIds': ids, 'locator': locator}, require_id=False)}

    def inline_items(self, paragraph):
        # Split marker runs to preserve ordinary text before/after a field,
        # including sources whose field markers share a run with text.
        children = []
        for child in paragraph:
            if child.tag == qn('w:r') and any(part.tag in (qn('w:fldChar'), qn('w:instrText')) for part in child):
                for part in child:
                    if part.tag == qn('w:rPr'): continue
                    run = element('r'); props = child.find(qn('w:rPr'))
                    if props is not None: run.append(deepcopy(props))
                    run.append(deepcopy(part)); children.append(run)
            else: children.append(child)
        index = 0
        while index < len(children):
            child = children[index]
            tagged = self.tagged(child) or (self.extra_reader.tagged(child) if self.extra_reader else None)
            if tagged: yield tagged; index += 1; continue
            if child.tag == qn('w:fldSimple'):
                instruction = child.get(qn('w:instr'), '')
                value = self.field(instruction) or (self.extra_reader.field(instruction) if self.extra_reader else None)
                if value: yield value
                else: yield from child.findall(qn('w:r'))
                index += 1; continue
            marker = child.find(qn('w:fldChar')) if child.tag == qn('w:r') else None
            if marker is not None and marker.get(qn('w:fldCharType')) == 'begin':
                end, depth, instruction = index, 0, ''
                while end < len(children):
                    for part in children[end]:
                        if part.tag == qn('w:fldChar'):
                            if part.get(qn('w:fldCharType')) == 'begin': depth += 1
                            if part.get(qn('w:fldCharType')) == 'end': depth -= 1
                        if part.tag == qn('w:instrText'): instruction += part.text or ''
                    if depth == 0: break
                    end += 1
                value = (self.field(instruction) or (self.extra_reader.field(instruction) if self.extra_reader else None)) if end < len(children) else None
                if value: yield value; index = end + 1; continue
            if child.tag == qn('w:sdt'):
                content = child.find(qn('w:sdtContent'))
                if content is not None:
                    self.warn('Inline content controls were reduced to supported visible text/citations.')
                    yield from self.inline_items(content)
            else: yield child
            index += 1

    def finish(self, model):
        known = {source['id'] for source in self.sources}
        if any(identifier not in known for node in walk(model) if node['type'] == 'documentCitation' for identifier in node['attrs']['sourceIds']): self.warn('Some citations have missing sources. Use Manage sources to locate and repair them.')
