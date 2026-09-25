import asyncio
import base64
import io
import json
import threading
from types import SimpleNamespace
from zipfile import ZipFile

import httpx
import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from PIL import Image

from services import chat_documents as docs, session_store, image_vault, image_store
from routes.artifacts import router


@pytest.fixture
def document_session(sessions_dir, tmp_path, monkeypatch):
    monkeypatch.setattr(docs, "ROOT", tmp_path / "artifacts")
    monkeypatch.setattr(image_vault, "ROOT", tmp_path / "vault")
    monkeypatch.setattr(image_store, "BLOBS_DIR", tmp_path / "blobs")
    return session_store.create_session()["id"]


@pytest.mark.parametrize("prompt", ["Okay, now make me a .docx", "Make it a Word document", "Can you export this as docx?", "/docx summarize this", "Save this in report.docx"])
def test_creation_intent(prompt):
    assert docs.wants_document(prompt)


@pytest.mark.parametrize("prompt", ["How do I create a docx in Python?", "Don't make a Word document", "Explain docx files", "Write python code for a docx", "hello", "```python\n# make a docx\n```"])
def test_code_and_explanation_requests_stay_normal_chat(prompt):
    assert not docs.wants_document(prompt)


def spec(blocks=None):
    return docs.DocumentSpec(title="Sine Wave Plot", filename="sine_wave.docx", blocks=blocks or [
        {"type": "paragraph", "text": "The plot shows one full sine wave cycle."},
        {"type": "function_plot", "function": "sin", "caption": "Figure 1 Sine wave from 0 to 2 pi"},
        {"type": "table", "headers": ["Parameter", "Value"], "rows": [["Amplitude", "1"]]},
    ])


def test_real_docx_contains_text_table_and_embedded_plot(document_session):
    artifact = docs.create_document(spec(), document_session, {}, threading.Event())
    path = docs.file_path(artifact["id"], "document.docx")
    doc = Document(path)
    assert doc.paragraphs[0].text == "Sine Wave Plot"
    assert len(doc.inline_shapes) == 1
    assert doc.tables[0].cell(1, 1).text == "1"
    with ZipFile(path) as archive:
        assert archive.testzip() is None
        assert "word/media/image1.png" in archive.namelist()
    metadata = docs.read_artifact(artifact["id"])
    assert metadata["blocks"][1]["image_file"] == "image-1.png"
    assert artifact["name"] == "sine_wave.docx"


def test_attached_image_is_embedded_and_future_lock_blocks_all_access(document_session, monkeypatch):
    output = io.BytesIO(); Image.new("RGB", (40, 30), "blue").save(output, "PNG")
    session_store.append_messages(document_session, [{"id": "image-message", "role": "user", "content": "Use this plot", "images": [base64.b64encode(output.getvalue()).decode()]}])
    inventory = docs.image_inventory(document_session)
    artifact = docs.create_document(spec([{"type": "image", "image_id": "image-1", "caption": "Attached plot"}]), document_session, inventory, threading.Event())
    metadata = docs.read_artifact(artifact["id"])
    monkeypatch.setattr(image_vault, "locked_hashes", lambda: set(metadata["source_hashes"]))
    for operation in [lambda: docs.read_artifact(artifact["id"]), lambda: docs.file_path(artifact["id"], "document.docx"), lambda: docs.file_path(artifact["id"], "image-0.png")]:
        with pytest.raises(image_vault.LockedImageError): operation()
    assert docs.image_inventory(document_session) == {}


@pytest.mark.parametrize("block", [
    {"type": "image", "image_id": "../../private.png"},
    {"type": "image", "image_id": "https://example.com/secret"},
    {"type": "table", "headers": ["a", "b"], "rows": [["a"]]},
    {"type": "function_plot", "function": "sin", "x_min": 3, "x_max": 1},
])
def test_invalid_assets_and_tables_do_not_publish_partial_files(document_session, block):
    with pytest.raises(ValueError): docs.create_document(spec([block]), document_session, {}, threading.Event())
    assert not list(docs.ROOT.iterdir())


def test_cancellation_leaves_no_completed_document(document_session):
    cancel = threading.Event(); cancel.set()
    with pytest.raises(InterruptedError): docs.create_document(spec(), document_session, {}, cancel)
    assert not list(docs.ROOT.iterdir())


def test_document_source_must_be_a_session_id(document_session):
    with pytest.raises(ValueError): docs.image_inventory("../../outside")
    with pytest.raises(ValueError): docs.create_document(spec(), "../../outside", {}, threading.Event())


def test_duplicate_title_and_windows_reserved_filename(document_session):
    value = spec([{"type": "heading", "text": "Sine Wave Plot"}, {"type": "paragraph", "text": "One cycle."}])
    value.filename = "CON.docx"
    artifact = docs.create_document(value, document_session, {}, threading.Event())
    assert artifact["name"] == "document-CON.docx"
    doc = Document(docs.file_path(artifact["id"], "document.docx"))
    assert [p.text for p in doc.paragraphs] == ["Sine Wave Plot", "One cycle."]


def test_artifact_routes_require_the_app_credential(document_session):
    from services.session_guard import SessionGuard
    artifact = docs.create_document(spec(), document_session, {}, threading.Event())
    app = FastAPI(); app.include_router(router); app.add_middleware(SessionGuard)
    with TestClient(app, base_url="http://127.0.0.1:8000") as client:
        for suffix in ("", "/download", "/images/image-1.png"):
            assert client.get(f"/artifacts/{artifact['id']}{suffix}").status_code == 403


def test_cancelling_during_build_waits_for_worker_and_does_not_attach(document_session, monkeypatch):
    started, finished = threading.Event(), threading.Event()
    def building(_spec, _session, _inventory, cancel):
        started.set()
        assert cancel.wait(5)
        finished.set()
        raise InterruptedError("Cancelled")
    monkeypatch.setattr(docs, "create_document", building)
    async def connected(): return False
    async def handler(_request):
        return httpx.Response(200, content=json.dumps({"message": {"content": spec().model_dump_json()}, "done": True}))
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            async def consume():
                async for _ in docs.stream_document(client, {"messages": []}, SimpleNamespace(session_id=document_session, reply_message_id="cancelled", model="test"), SimpleNamespace(is_disconnected=connected)): pass
            task = asyncio.create_task(consume())
            assert await asyncio.to_thread(started.wait, 5)
            task.cancel()
            with pytest.raises(asyncio.CancelledError): await task
    asyncio.run(run())
    assert finished.is_set()
    assert session_store.get_session(document_session)["messages"] == []


def test_preview_download_traversal_and_deleted_chat(document_session):
    artifact = docs.create_document(spec(), document_session, {}, threading.Event())
    app = FastAPI(); app.include_router(router)
    with TestClient(app) as client:
        assert client.get(f"/artifacts/{artifact['id']}").json()["title"] == "Sine Wave Plot"
        response = client.get(f"/artifacts/{artifact['id']}/download")
        assert response.status_code == 200 and response.content.startswith(b"PK")
        assert "sine_wave.docx" in response.headers["content-disposition"]
        assert client.get(f"/artifacts/{artifact['id']}/images/document.json").status_code == 404
        assert client.get("/artifacts/invalid").status_code == 404
        session_store.delete_session(document_session)
        assert client.get(f"/artifacts/{artifact['id']}/download").status_code == 404


def stream_result(session_id, output=None, end=True, error=None, reason="stop"):
    captured = []
    async def handler(request):
        captured.append(json.loads(request.content))
        chunks = [{"message": {"content": output if output is not None else spec().model_dump_json()}, "done": end, "done_reason": reason}]
        if error: chunks = [{"error": error}]
        return httpx.Response(200, content="\n".join(json.dumps(chunk) for chunk in chunks))
    async def connected(): return False
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            request = SimpleNamespace(session_id=session_id, reply_message_id="reply-1", model="test-model")
            payload = {"messages": [{"role": "system", "content": "Source fact: amplitude is 1"}, {"role": "user", "content": "Make me a docx"}]}
            return [event async for event in docs.stream_document(client, payload, request, SimpleNamespace(is_disconnected=connected))]
    events = [json.loads(line[6:]) for chunk in asyncio.run(run()) for line in chunk.splitlines() if line.startswith("data: ")]
    return events, captured


def test_stream_persists_artifact_before_completion_and_retains_context(document_session):
    events, captured = stream_result(document_session)
    result = events[-1]
    assert result["done"] and result["artifacts"][0]["kind"] == "docx"
    assert "Source fact" in captured[0]["messages"][0]["content"]
    assert captured[0]["format"]["type"] == "object"
    message = session_store.get_session(document_session)["messages"][0]
    assert message["id"] == "reply-1" and message["artifacts"] == result["artifacts"]
    assert "Amplitude | 1" in message["document_text"]


@pytest.mark.parametrize("kwargs", [{"output": "```python\nprint('fake file')\n```"}, {"end": False}, {"error": "CUDA failure"}, {"reason": "length"}])
def test_bad_or_incomplete_model_output_returns_visible_error_without_fake_attachment(document_session, kwargs):
    events, _ = stream_result(document_session, **kwargs)
    assert events[-1]["error"] and events[-1]["done"]
    assert not events[-1].get("artifacts")
    assert session_store.get_session(document_session)["messages"] == []
