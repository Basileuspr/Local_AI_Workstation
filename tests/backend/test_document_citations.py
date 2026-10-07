from copy import deepcopy
from io import BytesIO
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from services.document_editor import DEFAULT_LAYOUT, export_docx, import_docx, validate_model
from services.document_references import element, walk
from services.document_citations import B, canonical_source, context, citation_label, bibliography_entries


def source(digit='a', **changes):
    return canonical_source({'id': 'LAW_S' + digit * 32, 'type': 'book', 'title': 'Analytical notes', 'year': '1843', 'authors': [{'family': 'Lovelace', 'given': 'Ada'}], 'publisher': 'Example Press', **changes})


def cite(*ids, **changes):
    return {'type': 'documentCitation', 'attrs': {'id': 'cite-' + 'b' * 32, 'sourceIds': list(ids), 'mode': 'parenthetical', 'locator': 'p. 12', 'prefix': 'See', 'suffix': 'for details', **changes}}


def model(*sources, style='apa'):
    return {'type': 'doc', 'attrs': {'sources': list(sources), 'citationStyle': style}, 'content': [
        {'type': 'paragraph', 'content': [{'type': 'text', 'text': 'Before'}, cite(sources[0]['id']), {'type': 'text', 'text': 'After'}]},
        {'type': 'documentBibliography', 'attrs': {'title': 'References', 'includeUncited': False}}]}


def save(document):
    stream = BytesIO(); document.save(stream); return stream.getvalue()


@pytest.mark.parametrize('style', ['apa', 'numeric'])
def test_native_parts_cached_citation_and_live_bibliography_survive_three_round_trips(style):
    data = model(source(), source('c', type='website', title='Project guide', authors=[{'literal': 'Example Lab'}], url='https://example.org/guide'), style=style)
    data['content'][0]['content'][1]['attrs']['sourceIds'].append(data['attrs']['sources'][1]['id'])
    for _ in range(3):
        package = export_docx(data, DEFAULT_LAYOUT)
        with zipfile.ZipFile(BytesIO(package)) as archive:
            xml = etree.fromstring(archive.read('word/document.xml'))
            field = xml.find('.//w:fldSimple', xml.nsmap)
            assert field.get(qn('w:fldLock')) == '1' and 'CITATION LAW_S' in field.get(qn('w:instr')) and '\\m LAW_S' in field.get(qn('w:instr'))
            assert ''.join(field.itertext()) == citation_label(data['content'][0]['content'][1]['attrs'], context(data))
            native = next(root for name in archive.namelist() if name.startswith('customXml/item') and name.endswith('.xml') for root in [etree.fromstring(archive.read(name))] if root.tag == '{' + B + '}Sources')
            assert native.tag == '{' + B + '}Sources' and len(native) == 2
            assert native.find('.//{' + B + '}Last').text == 'Lovelace'
            assert native.find('.//{' + B + '}Corporate').text == 'Example Lab'
            assert 'https://example.org/guide' in archive.read('word/_rels/document.xml.rels').decode()
            assert xml.find('.//w:ind', xml.nsmap).get(qn('w:hanging')) == '720'
        result = import_docx(package); data = result['document']
        assert not any('citation' in warning.lower() or 'bibliography' in warning.lower() or 'source edits' in warning.lower() for warning in result['warnings'])
        assert data['attrs']['citationStyle'] == style and len(data['attrs']['sources']) == 2
        found = next(node for node in walk(data) if node['type'] == 'documentCitation')
        assert found['attrs']['locator'] == 'p. 12' and found['attrs']['prefix'] == 'See'
        assert [node['type'] for node in data['content']] == ['paragraph', 'documentBibliography']
        assert ''.join(node.get('text', '') for node in walk(data)) == 'BeforeAfter'


def test_basic_formats_match_expected_text_italics_links_and_year_suffixes():
    a = source(title='Alpha'); z = source('c', title='Zeta'); data = model(a, z)
    data['content'][0]['content'][1]['attrs'].update(sourceIds=[z['id'], a['id']], prefix='', suffix='', locator='')
    assert citation_label(data['content'][0]['content'][1]['attrs'], context(data)) == '(Lovelace, 1843b; Lovelace, 1843a)'
    data['attrs']['citationStyle'] = 'numeric'
    assert citation_label(data['content'][0]['content'][1]['attrs'], context(data)) == '[1, 2]'
    assert [entry['source']['title'] for entry in bibliography_entries(data)] == ['Zeta', 'Alpha']
    article = source(type='article', title='A study', journal='Example Journal', volume='2', issue='1', pages='10–20', doi='10.1234/example')
    data = model(article)
    assert ''.join(part['text'] for part in bibliography_entries(data)[0]['parts']) == 'Lovelace, A. (1843). A study. Example Journal, 2(1), 10–20. https://doi.org/10.1234/example'
    assert bibliography_entries(data)[0]['parts'][-1]['href'] == 'https://doi.org/10.1234/example'


def test_native_source_edits_import_and_mixed_author_metadata_is_preserved():
    mixed = source(authors=[{'literal': 'Example Lab'}, {'family': 'Lovelace', 'given': 'Ada'}])
    doc = Document(BytesIO(export_docx(model(mixed), DEFAULT_LAYOUT)))
    part = next(rel.target_part for rel in doc.part.rels.values() if not rel.is_external and b'Analytical notes' in rel.target_part.blob and b'b:Sources' in rel.target_part.blob)
    root = etree.fromstring(part.blob); root.find('.//{' + B + '}Title').text = 'Changed outside this editor'; part._blob = etree.tostring(root)
    result = import_docx(save(doc))
    assert result['document']['attrs']['sources'][0]['title'] == 'Changed outside this editor'
    assert result['document']['attrs']['sources'][0]['authors'] == mixed['authors']
    assert any('Native source edits' in warning for warning in result['warnings'])


@pytest.mark.parametrize('complex_field', [False, True])
def test_foreign_citation_fields_become_editable_without_losing_adjacent_text(complex_field):
    original = source(); doc = Document(BytesIO(export_docx(model(original), DEFAULT_LAYOUT)))
    p = doc.paragraphs[0]._p
    p.remove(p.find(qn('w:sdt')))
    if complex_field:
        run = element('r'); run.append(element('t')); run[-1].text = 'Left'
        run.append(element('fldChar', fldCharType='begin')); run.append(element('instrText')); run[-1].text = ' CITATION ' + original['id'] + ' \\p "p. 9" '
        run.append(element('fldChar', fldCharType='separate')); run.append(element('t')); run[-1].text = 'Old cached label'
        run.append(element('fldChar', fldCharType='end')); run.append(element('t')); run[-1].text = 'Right'; p.insert(list(p).index(p.find(qn('w:r'))) + 1, run)
    else:
        field = element('fldSimple', instr=' CITATION ' + original['id'] + ' \\p "p. 9" '); run = element('r'); run.append(element('t')); run[-1].text = 'Old cached label'; field.append(run); p.insert(1, field)
    result = import_docx(save(doc)); nodes = list(walk(result['document']))
    found = next(node for node in nodes if node['type'] == 'documentCitation')
    assert found['attrs']['sourceIds'] == [original['id']] and found['attrs']['locator'] == 'p. 9'
    visible = ''.join(node.get('text', '') for node in nodes)
    assert visible == ('BeforeLeftRightAfter' if complex_field else 'BeforeAfter')
    assert any('Native citation fields' in warning for warning in result['warnings'])


def test_missing_source_remains_visible_and_additional_bibliography_becomes_text():
    data = model(source()); data['content'][0]['content'][1]['attrs']['sourceIds'] = ['LAW_S' + 'f' * 32]
    doc = Document(BytesIO(export_docx(data, DEFAULT_LAYOUT)))
    bibliography = doc.element.body.findall(qn('w:sdt'))[0]; bibliography.addnext(deepcopy(bibliography))
    result = import_docx(save(doc))
    assert any('missing sources' in warning for warning in result['warnings'])
    assert any('additional bibliography' in warning for warning in result['warnings'])
    assert sum(node['type'] == 'documentBibliography' for node in walk(result['document'])) == 1
    found = next(node for node in walk(result['document']) if node['type'] == 'documentCitation')
    assert 'Missing source' in citation_label(found['attrs'], context(result['document']))


@pytest.mark.parametrize('changes', [{'year': '202'}, {'authors': [{'given': 'Ada'}]}, {'url': 'javascript:alert(1)'}, {'doi': 'bad'}, {'title': '😀' * 251}, {'title': '\ud800'}, {'type': 'script'}, {'authors': [{'literal': 'Lab', 'family': 'Name'}]}])
def test_invalid_sources_rejected(changes):
    with pytest.raises(ValueError): source(**changes)


def test_duplicate_ids_limits_invalid_details_and_nested_bibliography_rejected():
    original = source(); data = model(original, original)
    with pytest.raises(ValueError, match='unique'): validate_model(data)
    data = model(original); data['content'][0]['content'][1]['attrs']['locator'] = 'x' * 121
    with pytest.raises(ValueError): validate_model(data)
    data = model(original); data['content'].append(deepcopy(data['content'][1]))
    with pytest.raises(ValueError): validate_model(data)
    data = model(original); data['content'][1] = {'type': 'blockquote', 'content': [data['content'][1]]}
    with pytest.raises(ValueError): validate_model(data)
    data = model(original); data['attrs']['sources'] = [source(f'{index:032x}'[-1]) for index in range(101)]
    with pytest.raises(ValueError, match='100 sources'): validate_model(data)
    data = model(original); data['content'][0]['content'] = [cite(original['id'], id=f'cite-{index:032x}') for index in range(1001)]
    with pytest.raises(ValueError, match='1,000 citations'): validate_model(data)
