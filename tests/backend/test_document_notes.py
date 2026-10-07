from copy import deepcopy
from io import BytesIO
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

from services.document_editor import DEFAULT_LAYOUT, export_docx, import_docx, validate_model
from services.document_references import walk


def note(kind='footnote', digit='a', text='A note'):
    return {'type': 'documentNote', 'attrs': {'id': 'note-' + digit * 32, 'kind': kind,
        'body': {'type': 'doc', 'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': text}]}]}}}


def model(*notes):
    return {'type': 'doc', 'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': 'Before'}, *notes, {'type': 'text', 'text': 'After'}]}]}


def save(document):
    stream = BytesIO(); document.save(stream); return stream.getvalue()


def test_native_notes_parts_numbering_formatted_text_and_links_survive_three_round_trips():
    foot = note()
    foot['attrs']['body']['content'] = [{'type': 'paragraph', 'attrs': {'textAlign': 'right'}, 'content': [
        {'type': 'text', 'text': 'Rich note', 'marks': [{'type': 'bold'}, {'type': 'textStyle', 'attrs': {'fontFamily': 'Georgia', 'fontSize': '14pt', 'color': '#123456'}}]},
        {'type': 'hardBreak'}, {'type': 'text', 'text': 'Source', 'marks': [{'type': 'link', 'attrs': {'href': 'https://example.org/source'}}]}]},
        {'type': 'paragraph', 'content': [{'type': 'text', 'text': 'Another paragraph'}]}]
    source = model(foot, note('endnote', 'b', 'End note'))
    source['attrs'] = {'noteSettings': {'footnote': {'format': 'upperLetter', 'start': 3}, 'endnote': {'format': 'upperRoman', 'start': 4}}}
    for _ in range(3):
        package = export_docx(source, DEFAULT_LAYOUT)
        native = Document(BytesIO(package))
        assert native.paragraphs[0].text == 'BeforeAfter'
        assert len(native.element.xpath('.//w:footnoteReference')) == 1
        assert len(native.element.xpath('.//w:endnoteReference')) == 1
        props = native.sections[0]._sectPr
        assert props.find('w:footnotePr/w:numStart', props.nsmap).get(qn('w:val')) == '3'
        assert props.find('w:endnotePr/w:numFmt', props.nsmap).get(qn('w:val')) == 'upperRoman'
        with zipfile.ZipFile(BytesIO(package)) as archive:
            for kind in ('footnote', 'endnote'):
                root = etree.fromstring(archive.read(f'word/{kind}s.xml'))
                assert root.tag == qn(f'w:{kind}s')
                assert [child.get(qn('w:id')) for child in root] == ['-1', '0', '1']
                assert root.findall(f'.//w:{kind}Ref', root.nsmap)
                assert f'{kind}s' in archive.read('[Content_Types].xml').decode()
            assert 'https://example.org/source' in archive.read('word/_rels/footnotes.xml.rels').decode()
        result = import_docx(package)
        assert not any('unsupported note' in warning.lower() for warning in result['warnings'])
        source = result['document']
        found = [node for node in walk(source) if node['type'] == 'documentNote']
        assert [node['attrs']['kind'] for node in found] == ['footnote', 'endnote']
        blocks = found[0]['attrs']['body']['content']
        assert blocks[0]['attrs']['textAlign'] == 'right'
        assert blocks[0]['content'][0]['text'] == 'Rich note'
        assert {'type': 'bold'} in blocks[0]['content'][0]['marks']
        assert blocks[0]['content'][0]['marks'][1]['attrs']['fontFamily'] == 'Georgia'
        assert blocks[1]['content'][0]['text'] == 'Another paragraph'
        assert blocks[0]['content'][2]['marks'][-1]['attrs']['href'] == 'https://example.org/source'
        assert source['attrs']['noteSettings'] == {'footnote': {'format': 'upperLetter', 'start': 3}, 'endnote': {'format': 'upperRoman', 'start': 4}}


def test_missing_definition_is_visible_and_source_numbering_normalization_is_reported():
    native = Document(BytesIO(export_docx(model(note()), DEFAULT_LAYOUT)))
    native.element.xpath('.//w:footnoteReference')[0].set(qn('w:id'), '42')
    restart = native.sections[0]._sectPr.find('w:footnotePr/w:numRestart', native.element.nsmap)
    restart.set(qn('w:val'), 'eachPage')
    result = import_docx(save(native))
    found = next(node for node in walk(result['document']) if node['type'] == 'documentNote')
    assert found['attrs']['body']['content'][0]['content'][0]['text'] == '[Missing note text]'
    assert any('continuous numbering' in warning for warning in result['warnings'])
    assert any('Unreferenced' in warning for warning in result['warnings'])


def test_repeated_source_reference_becomes_independent_editable_notes():
    native = Document(BytesIO(export_docx(model(note()), DEFAULT_LAYOUT)))
    native.paragraphs[0]._p.append(deepcopy(native.paragraphs[0]._p[2]))
    result = import_docx(save(native))
    found = [node for node in walk(result['document']) if node['type'] == 'documentNote']
    assert len(found) == 2 and found[0]['attrs']['id'] != found[1]['attrs']['id']
    assert found[0]['attrs']['body'] == found[1]['attrs']['body']


@pytest.mark.parametrize('body', [
    {'type': 'doc', 'content': []},
    {'type': 'doc', 'content': [{'type': 'heading'}]},
    {'type': 'doc', 'content': [{'type': 'paragraph', 'content': [note()]}]},
    {'type': 'doc', 'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': 'x' * 10001}]}]},
    {'type': 'doc', 'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': '\ud800'}]}]},
    {'type': 'doc', 'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': 'link', 'marks': [{'type': 'link', 'attrs': {'href': 'file:///private'}}]}]}]},
    {'type': 'doc', 'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': 'link', 'marks': [{'type': 'link', 'attrs': {'href': '#Destination'}}]}]}]},
    {'type': 'doc', 'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': 'link', 'marks': [{'type': 'link', 'attrs': [1]}]}]}]},
    {'type': 'doc', 'content': [{'type': 'paragraph'}] * 41},
])
def test_invalid_note_bodies_are_rejected(body):
    item = note(); item['attrs']['body'] = body
    with pytest.raises(ValueError): export_docx(model(item), DEFAULT_LAYOUT)


@pytest.mark.parametrize('settings', [
    {'footnote': {'format': 'symbols'}}, {'endnote': {'start': 0}}, {'footnote': {'start': True}},
    {'endnote': {'start': 1.5}}, {'footnote': {'format': []}},
])
def test_invalid_note_numbering_is_rejected(settings):
    source = model(note()); source['attrs'] = {'noteSettings': settings}
    with pytest.raises(ValueError): validate_model(source)


def test_duplicate_ids_and_note_limit_are_rejected_and_emoji_uses_same_browser_character_budget():
    with pytest.raises(ValueError, match='unique'): validate_model(model(note(), note()))
    items = [note(digit='a') for _ in range(201)]
    for index, item in enumerate(items): item['attrs']['id'] = f'note-{index:032x}'
    with pytest.raises(ValueError, match='200 notes'): validate_model(model(*items))
    validate_model(model(note(text='😀' * 5000)))
    with pytest.raises(ValueError, match='10,000'): validate_model(model(note(text='😀' * 5001)))


def test_oversized_import_is_explicitly_truncated():
    native = Document(BytesIO(export_docx(model(note()), DEFAULT_LAYOUT)))
    part = next(rel.target_part for rel in native.part.rels.values() if rel.reltype == RT.FOOTNOTES)
    root = etree.fromstring(part.blob)
    root.findall('.//w:t', root.nsmap)[0].text = 'x' * 10001
    # Generic parts loaded by python-docx store their XML as bytes.
    part._blob = etree.tostring(root)
    result = import_docx(save(native))
    assert any('10,000 characters' in warning for warning in result['warnings'])
    source = next(node for node in walk(result['document']) if node['type'] == 'documentNote')
    assert len(source['attrs']['body']['content'][0]['content'][0]['text']) == 10000


def test_external_note_part_is_not_fetched_and_the_reference_is_repairable():
    package = export_docx(model(note()), DEFAULT_LAYOUT)
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(package)) as original, zipfile.ZipFile(output, 'w') as altered:
        for entry in original.infolist():
            content = original.read(entry)
            if entry.filename == 'word/_rels/document.xml.rels':
                root = etree.fromstring(content)
                relation = next(child for child in root if child.get('Type') == RT.FOOTNOTES)
                relation.set('TargetMode', 'External'); relation.set('Target', 'http://127.0.0.1:1/private')
                content = etree.tostring(root)
            altered.writestr(entry, content)
    result = import_docx(output.getvalue())
    assert any('not fetched' in warning for warning in result['warnings'])
    item = next(node for node in walk(result['document']) if node['type'] == 'documentNote')
    assert item['attrs']['body']['content'][0]['content'][0]['text'] == '[Missing note text]'
