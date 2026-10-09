from copy import deepcopy
from io import BytesIO
import zipfile
import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from services.document_editor import DEFAULT_LAYOUT, export_docx, import_docx, validate_model
from services.document_references import walk, element
from services.document_captions import caption_context, reference_targets, reference_text


def caption(digit='a', label='Figure', text='Original'):
    return {'type':'documentCaption','attrs':{'id':'LAW_C' + digit * 32,'label':label,'text':text}}


def ref(target=None, display='full', hyperlink=True):
    return {'type':'documentCrossReference','attrs':{'id':'xref-' + 'b' * 32,'target':target or 'LAW_C' + 'a' * 32,'display':display,'hyperlink':hyperlink}}


def model():
    return {'type':'doc','attrs':{'captionSettings':{'Figure':{'format':'upperRoman','start':4}}},'content':[caption(),
        {'type':'paragraph','content':[{'type':'text','text':'Before'},ref(display='labelNumber'),{'type':'text','text':'After'}]},caption('c','Table','Data')]}


def save(doc):
    stream = BytesIO(); doc.save(stream); return stream.getvalue()


def test_native_sequence_and_reference_fields_with_bookmark_range_survive_three_round_trips():
    data = model()
    for _ in range(3):
        raw = export_docx(data,DEFAULT_LAYOUT)
        with zipfile.ZipFile(BytesIO(raw)) as package:
            assert len(package.namelist()) == len(set(package.namelist()))
            root = etree.fromstring(package.read('word/document.xml'))
            fields = root.findall('.//w:fldSimple',root.nsmap)
            assert [field.get(qn('w:instr')).strip().split()[0] for field in fields] == ['SEQ','REF','SEQ']
            assert all(field.get(qn('w:fldLock')) == '1' for field in fields)
            assert ''.join(fields[0].itertext()) == 'IV' and ''.join(fields[1].itertext()) == 'Figure IV'
            assert root.find('.//w:bookmarkStart',root.nsmap).get(qn('w:name')) == 'LAW_C' + 'a' * 32
            caption_paragraph = root.find('.//w:sdtContent/w:p',root.nsmap)
            assert list(caption_paragraph)[1].tag == qn('w:bookmarkStart') and list(caption_paragraph)[-1].tag == qn('w:bookmarkEnd')
        result = import_docx(raw); data = result['document']
        assert not any('missing' in warning.lower() or 'caption' in warning.lower() or 'cross-reference' in warning.lower() for warning in result['warnings'])
        assert [node['type'] for node in data['content']] == ['documentCaption','paragraph','documentCaption']
        assert data['attrs']['captionSettings'] == {'Figure':{'format':'upperRoman','start':4}}
        found = next(node for node in walk(data) if node['type'] == 'documentCrossReference')
        assert found['attrs']['display'] == 'labelNumber' and reference_text(found['attrs'],reference_targets(data)) == 'Figure IV'
        assert ''.join(node.get('text','') for node in walk(data)) == 'BeforeAfter'


@pytest.mark.parametrize('display,expected',[('full','Figure IV: Original'),('labelNumber','Figure IV'),('number','IV'),('text','Original')])
def test_all_displays_and_source_edits_regenerate_cached_fields(display,expected):
    data = model(); data['content'][1]['content'][1]['attrs'].update(display=display,hyperlink=False)
    assert reference_text(data['content'][1]['content'][1]['attrs'],reference_targets(data)) == expected
    assert import_docx(export_docx(data,DEFAULT_LAYOUT))['document']['content'][1]['content'][1]['attrs']['hyperlink'] is False
    data['content'][0]['attrs']['text'] = 'Revised'; data['attrs']['captionSettings']['Figure']['start'] = 5
    assert caption_context(data)['LAW_C' + 'a' * 32]['full'] == 'Figure V: Revised'


def test_missing_destinations_are_visible_and_duplicate_captions_get_independent_ids():
    data = model(); data['content'][0] = caption('f')
    result = import_docx(export_docx(data,DEFAULT_LAYOUT))
    assert any('missing destinations' in warning for warning in result['warnings'])
    assert 'Missing reference' in reference_text(result['document']['content'][1]['content'][1]['attrs'],reference_targets(result['document']))
    doc = Document(BytesIO(export_docx(model(),DEFAULT_LAYOUT))); first = doc.element.body.find(qn('w:sdt')); first.addnext(deepcopy(first))
    result = import_docx(save(doc)); captions = [node for node in walk(result['document']) if node['type']=='documentCaption']
    assert len({node['attrs']['id'] for node in captions}) == 3
    assert any('independent ID' in warning for warning in result['warnings'])


def test_native_caption_text_edits_are_imported_and_references_regenerated():
    doc = Document(BytesIO(export_docx(model(),DEFAULT_LAYOUT)))
    root = doc.element.body.find(qn('w:sdt'))
    texts = root.findall('.//w:t',root.nsmap); texts[-1].text = ': Revised externally'
    result = import_docx(save(doc)); data = result['document']
    assert data['content'][0]['attrs']['text'] == 'Revised externally'
    assert any('Native caption label/text edits' in warning for warning in result['warnings'])
    assert reference_targets(data)['LAW_C' + 'a' * 32]['full'] == 'Figure IV: Revised externally'


@pytest.mark.parametrize('complex_field',[False,True])
def test_native_ref_import_keeps_adjacent_text_without_duplicate_cached_labels(complex_field):
    doc = Document(BytesIO(export_docx(model(),DEFAULT_LAYOUT))); p = doc.paragraphs[0]._p; control = p.find(qn('w:sdt')); index = list(p).index(control); p.remove(control)
    instruction = ' REF ' + 'LAW_C' + 'a' * 32 + ' \\h '
    if complex_field:
        run = element('r'); run.append(element('t')); run[-1].text='Left'; run.append(element('fldChar',fldCharType='begin'))
        run.append(element('instrText')); run[-1].text=instruction; run.append(element('fldChar',fldCharType='separate')); run.append(element('t')); run[-1].text='Old label'
        run.append(element('fldChar',fldCharType='end')); run.append(element('t')); run[-1].text='Right'; p.insert(index,run)
    else:
        field = element('fldSimple',instr=instruction); run=element('r'); run.append(element('t')); run[-1].text='Old label'; field.append(run); p.insert(index,field)
    result = import_docx(save(doc)); found = next(node for node in walk(result['document']) if node['type']=='documentCrossReference')
    assert found['attrs']['hyperlink'] and found['attrs']['display']=='full'
    assert ''.join(node.get('text','') for node in walk(result['document'])) == ('BeforeLeftRightAfter' if complex_field else 'BeforeAfter')
    assert any('Native REF fields' in warning for warning in result['warnings'])


@pytest.mark.parametrize('changes',[{'label':'x"bad'}, {'label':' Figure'}, {'text':'😀'*501}, {'text':'\ud800'}, {'id':'bad'}])
def test_invalid_captions_rejected(changes):
    data=model(); data['content'][0]['attrs'].update(changes)
    with pytest.raises(ValueError): validate_model(data)


@pytest.mark.parametrize('rule',[{'format':[],'start':1},{'format':'decimal','start':False},{'format':'decimal','start':0},{'format':'decimal','start':10000}])
def test_invalid_numbering_rejected(rule):
    data=model(); data['attrs']['captionSettings']['Figure']=rule
    with pytest.raises(ValueError): validate_model(data)


def test_limits_duplicate_ids_invalid_references_and_nested_captions_rejected():
    data=model(); data['content'].append(deepcopy(data['content'][0]))
    with pytest.raises(ValueError,match='unique'): validate_model(data)
    data=model(); data['content'][0]={'type':'bulletList','content':[{'type':'listItem','content':[{'type':'paragraph'},caption()]}]}
    with pytest.raises(ValueError): validate_model(data)
    data=model(); data['content'][1]['content'][1]['attrs']['target']='javascript:bad'
    with pytest.raises(ValueError): validate_model(data)
    data=model(); data['content']=[caption() for _ in range(201)]
    for index,node in enumerate(data['content']): node['attrs']['id']=f'LAW_C{index:032x}'
    with pytest.raises(ValueError,match='200 captions'): validate_model(data)
    data=model(); data['attrs']['captionSettings']={f'Label{index}':{'format':'decimal','start':1} for index in range(21)}
    with pytest.raises(ValueError,match='20 caption'): validate_model(data)
