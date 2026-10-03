from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches

from services.document_editor import DEFAULT_LAYOUT, export_docx, import_docx
from services.document_tables import table_grid, _element


def cell(text, **attrs):
    return {'type': 'tableCell', 'attrs': attrs, 'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': text}]}]}


def table(*rows, **attrs):
    return {'type': 'table', 'attrs': attrs, 'content': [{'type': 'tableRow', 'content': row} for row in rows]}


def exported(node):
    return export_docx({'type': 'doc', 'content': [node]}, DEFAULT_LAYOUT)


def saved(doc):
    buffer = BytesIO(); doc.save(buffer)
    return buffer.getvalue()


def text_content(node):
    return node.get('text', '') + ''.join(text_content(child) for child in node.get('content', []))


@pytest.mark.parametrize('border', ['grid', 'outer', 'none'])
@pytest.mark.parametrize('alignment', ['left', 'center', 'right'])
def test_two_dimensional_merge_widths_shading_and_alignment_survive_repeated_round_trips(border, alignment):
    merged = cell('Merged only once', colspan=2, rowspan=2, colwidth=[96, 144], backgroundColor='#FFE0B2', verticalAlign='center')
    model = table([merged, cell('C1', colwidth=[192])], [cell('C2', colwidth=[192])],
                  [cell('A3', colwidth=[96]), cell('B3', colwidth=[144]), cell('C3', colwidth=[192])],
                  tableAlignment=alignment, borderPreset=border)
    raw = exported(model)
    original = Document(BytesIO(raw)).tables[0]
    assert original.cell(0, 0)._tc is original.cell(1, 1)._tc
    assert original.cell(0, 0).vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER
    assert [column.width.inches for column in original.columns] == [1, 1.5, 2]
    assert original.cell(0, 0).width.inches == 2.5
    for _ in range(3):
        result = import_docx(raw)
        node = result['document']['content'][0]
        assert node['attrs']['borderPreset'] == border
        assert node['attrs']['tableAlignment'] == alignment
        assert node['content'][0]['content'][0]['attrs'] == merged['attrs']
        assert len(node['content'][1]['content']) == 1
        assert text_content(node).count('Merged only once') == 1
        assert table_grid(node)['columns'] == 3
        raw = exported(node)


def test_imports_external_vertical_merges_and_preserves_each_visible_cell_once():
    doc = Document(); tbl = doc.add_table(rows=3, cols=3)
    tbl.cell(0, 1).merge(tbl.cell(2, 2)).text = 'Tall merged cell'
    for row in range(3): tbl.cell(row, 0).text = str(row)
    tbl.alignment = WD_TABLE_ALIGNMENT.RIGHT
    result = import_docx(saved(doc))['document']['content'][0]
    assert [len(row['content']) for row in result['content']] == [2, 1, 1]
    merged = result['content'][0]['content'][1]
    assert merged['attrs']['rowspan'] == 3 and merged['attrs']['colspan'] == 2
    assert text_content(result) == '0Tall merged cell12'
    assert result['attrs']['tableAlignment'] == 'right'
    assert Document(BytesIO(exported(result))).tables[0].cell(2, 2).text == 'Tall merged cell'


def test_full_table_merge_allows_empty_continuation_rows():
    node = table([cell('One cell', colspan=2, rowspan=3)], [], [])
    result = import_docx(exported(node))['document']['content'][0]
    assert [len(row['content']) for row in result['content']] == [1, 0, 0]
    assert text_content(result) == 'One cell'


def test_header_cell_semantics_and_repeated_first_row_survive():
    heading = cell('Heading', colspan=2, colwidth=[96, 144]); heading['type'] = 'tableHeader'
    node = table([heading], [cell('A', colwidth=[96]), cell('B', colwidth=[144])], repeatHeader=True)
    raw = exported(node)
    tbl = Document(BytesIO(raw)).tables[0]
    assert tbl.rows[0]._tr.xpath('./w:trPr/w:tblHeader')
    result = import_docx(raw)['document']['content'][0]
    assert result['attrs']['repeatHeader']
    assert result['content'][0]['content'][0]['type'] == 'tableHeader'
    assert result['content'][1]['content'][0]['type'] == 'tableCell'


def test_wide_tables_scale_to_the_page_while_preserving_column_proportions():
    node = table([cell('A', colwidth=[768]), cell('B', colwidth=[384])])
    widths = [column.width.inches for column in Document(BytesIO(exported(node))).tables[0].columns]
    assert sum(widths) == pytest.approx(6.5, abs=.002)
    assert widths[0] / widths[1] == pytest.approx(2, abs=.002)


def test_nested_tables_are_flattened_with_notice_and_omitted_grid_cells_are_filled():
    doc = Document(); outer = doc.add_table(rows=2, cols=2)
    outer.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = 'Nested text'
    row = outer.rows[1]._tr
    row.remove(row.tc_lst[0]); row.get_or_add_trPr().append(_element('gridBefore', val=1))
    result = import_docx(saved(doc))
    assert any('Nested tables' in warning for warning in result['warnings'])
    assert any('Omitted cells' in warning for warning in result['warnings'])
    node = result['document']['content'][0]
    assert text_content(node).count('Nested text') == 1
    assert len(node['content'][1]['content']) == 2
    exported(node)


@pytest.mark.parametrize('patch', [{'colspan': 0}, {'colspan': True}, {'colspan': 13}, {'rowspan': 2},
    {'colwidth': [12]}, {'colwidth': [96, 96]}, {'colwidth': [float('nan')]},
    {'backgroundColor': 'url(secret)'}, {'verticalAlign': {}}, {'rowspan': '2'}])
def test_bad_cell_geometry_and_formats_are_rejected(patch):
    with pytest.raises(ValueError): exported(table([cell('A', **patch)]))


def test_grid_rejects_overlap_gaps_inconsistent_widths_and_repeating_vertical_merge():
    invalid = [table([cell('A'), cell('B')], [cell('C')]),
               table([cell('A'), cell('B', rowspan=2)], [cell('C', colspan=2)]),
               table([cell('A', colwidth=[96])], [cell('B', colwidth=[144])]),
               table([cell('A', rowspan=2)], [], repeatHeader=True)]
    for node in invalid:
        with pytest.raises(ValueError): exported(node)


def test_oversized_external_grid_span_is_rejected_before_expanding_cells():
    doc = Document(); tbl = doc.add_table(rows=1, cols=1)
    tbl.cell(0, 0)._tc.get_or_add_tcPr().append(_element('gridSpan', val=1000000000))
    with pytest.raises(ValueError, match='grid spans'): import_docx(saved(doc))
