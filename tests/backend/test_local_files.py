from contextlib import closing
import hashlib
import os
from pathlib import Path
import sqlite3
import stat
import threading
import zipfile

import pytest
from docx import Document as Word
from fastapi import FastAPI
from fastapi.testclient import TestClient

from services import local_documents as documents, local_database as database, local_video as video, local_files as files
from routes.local_files import router


@pytest.fixture
def docx(tmp_path):
    path = tmp_path / 'document.docx'
    doc = Word(); doc.add_heading('Heading', 1)
    p = doc.add_paragraph(); p.add_run('Plain '); p.add_run('bold').bold = True; p.add_run('italic').italic = True
    doc.add_paragraph('Another paragraph')
    table = doc.add_table(rows=2, cols=2); table.cell(0, 0).text = 'Cell'; table.cell(1, 1).text = 'Last'
    doc.sections[0].header.paragraphs[0].text = 'Preserve header'
    doc.save(path)
    return path


def change(doc, text='Edited'):
    run = doc.model()['blocks'][1]['runs'][0]
    return {key: (text if key == 'text' else run[key]) for key in ('id', 'text', 'bold', 'italic')}


def test_docx_model_and_save_as_preserves_package(docx, tmp_path):
    before = docx.read_bytes(); doc = documents.Document(docx)
    model = doc.model()
    assert model['blocks'][0]['style'] == 'Heading1'
    assert model['blocks'][1]['runs'][1]['bold']
    assert model['blocks'][1]['runs'][2]['italic']
    assert model['blocks'][3]['rows'][0][0]['paragraphs'][0]['runs'][0]['text'] == 'Cell'
    assert any('headers' in warning for warning in model['warnings'])
    target = tmp_path / 'copy.docx'
    doc.save(target, [change(doc)], None, True)
    assert docx.read_bytes() == before
    assert Word(target).paragraphs[1].text.startswith('Edited')
    with zipfile.ZipFile(docx) as original, zipfile.ZipFile(target) as saved:
        assert original.namelist() == saved.namelist()
        for name in original.namelist():
            if name != 'word/document.xml': assert original.read(name) == saved.read(name)
    run = doc.model()['blocks'][3]['rows'][0][0]['paragraphs'][0]['runs'][0]
    doc.save(target, [{'id':run['id'], 'text':'New cell', 'bold':True, 'italic':True}], None, True)
    reopened = Word(target)
    assert reopened.tables[0].cell(0,0).text == 'New cell'
    assert reopened.tables[0].cell(0,0).paragraphs[0].runs[0].bold
    assert reopened.sections[0].header.paragraphs[0].text == 'Preserve header'


def test_save_conflicts_readonly_and_atomic_failure(docx, monkeypatch):
    doc = documents.Document(docx); before = docx.read_bytes()
    with pytest.raises(ValueError, match='acknowledge'): doc.save(docx, [change(doc)], None)
    monkeypatch.setattr(documents.os, 'replace', lambda *_: (_ for _ in ()).throw(PermissionError('locked')))
    with pytest.raises(ValueError, match='another application'): doc.save(docx, [change(doc)], None, True)
    assert docx.read_bytes() == before
    assert not list(docx.parent.glob('.law-docx-*'))
    docx.chmod(stat.S_IREAD)
    try:
        with pytest.raises(ValueError, match='read-only'): doc.save(docx, [], None, True)
    finally: docx.chmod(stat.S_IWRITE | stat.S_IREAD)
    docx.write_bytes(before + b'changed')
    with pytest.raises(ValueError, match='changed'): doc.save(docx, [], None, True)


@pytest.mark.parametrize('xml', [b'<broken', b'<!DOCTYPE w:document [<!ENTITY x SYSTEM "file:///secret">]><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>&x;</w:body></w:document>'])
def test_bad_document_xml(tmp_path, xml):
    path = tmp_path / 'bad.docx'
    with zipfile.ZipFile(path, 'w') as z:
        z.writestr('[Content_Types].xml', '<Types/>'); z.writestr('word/document.xml', xml)
    with pytest.raises(ValueError): documents.Document(path)


def test_bad_document_container_and_traversal(tmp_path):
    path = tmp_path / 'bad.docx'; path.write_text('not zip')
    with pytest.raises(ValueError): documents.Document(path)
    with zipfile.ZipFile(path,'w') as z: z.writestr('../escape','bad')
    with pytest.raises(ValueError, match='Unsafe'): documents.Document(path)


@pytest.fixture
def db(tmp_path):
    path = tmp_path / 'example.db'
    with closing(sqlite3.connect(path)) as connection:
        connection.executescript('CREATE TABLE messages(id INTEGER PRIMARY KEY, content TEXT); CREATE TABLE other(value BLOB); CREATE VIEW names AS SELECT content FROM messages;')
        connection.executemany('INSERT INTO messages VALUES(?,?)', ((i, f'row {i}') for i in range(100050)))
        connection.commit()
    return path


def test_database_snapshot_schema_pagination_filter_count(db, tmp_path):
    before = documents.digest(db); folder=tmp_path/'snapshot'; folder.mkdir()
    target = database.snapshot(db, folder)
    schema = database.schema(target)
    assert {row['type'] for row in schema} == {'table','view'}
    assert schema[0]['columns'][0]['primary_key'] == 1
    page = database.query(target,db.name,table='messages',offset=100000)
    assert len(page['rows']) == 50 and page['offset'] == 100000 and not page['has_more']
    assert database.query(target,db.name,table='messages',search='row 100049')['rows'][0][0] == 100049
    assert database.query(target,db.name,table='messages',count=True)['row_count'] == 100050
    assert database.query(target,db.name,sql='WITH x AS (SELECT 3 v) SELECT v FROM x')['rows'] == [[3]]
    assert documents.digest(db) == before
    assert not Path(str(db)+'-wal').exists()


@pytest.mark.parametrize('sql', ['INSERT INTO messages VALUES(3,4)','UPDATE messages SET content=1','DELETE FROM messages','DROP TABLE messages','ALTER TABLE messages ADD x','CREATE TABLE x(y)','VACUUM','PRAGMA user_version=4','PRAGMA journal_mode=WAL','ATTACH DATABASE \'x\' AS x','SELECT 1; DROP TABLE messages','WITH x AS (SELECT 1) DELETE FROM messages','SELECT load_extension(\'evil\')','SELECT * FROM pragma_wal_checkpoint','SELECT missing FROM messages'])
def test_blocked_sql(db,sql):
    before = documents.digest(db)
    with pytest.raises(ValueError): database.query(db,db.name,sql=sql)
    assert documents.digest(db) == before


def test_query_cost_and_result_limits(db):
    with pytest.raises(ValueError): database.query(db,db.name,sql='WITH RECURSIVE x(v) AS (SELECT 1 UNION ALL SELECT v+1 FROM x) SELECT sum(v) FROM x')
    with pytest.raises(ValueError): database.query(db,db.name,sql='SELECT randomblob(4000000)')


def test_non_sqlite_corrupt_and_locked(db,tmp_path):
    other = tmp_path/'other.db'; other.write_bytes(b'another database format')
    with pytest.raises(ValueError,match='not recognized'): database.snapshot(other,tmp_path)
    other.write_bytes(database.HEADER + b'\xff'*300)
    corrupt=database.snapshot(other,tmp_path)
    with pytest.raises(sqlite3.DatabaseError): database.schema(corrupt)
    with closing(sqlite3.connect(db)) as connection:
        connection.execute('BEGIN EXCLUSIVE'); connection.execute("UPDATE messages SET content='uncommitted' WHERE id=1")
        with pytest.raises(ValueError,match='journal'): database.snapshot(db,tmp_path)
        connection.rollback()


def test_wal_snapshot_includes_commits_without_touching_source(tmp_path):
    path=tmp_path/'live.db'; folder=tmp_path/'copy';folder.mkdir()
    with closing(sqlite3.connect(path)) as db:
        db.execute('PRAGMA journal_mode=WAL');db.execute('CREATE TABLE x(n)');db.execute('INSERT INTO x VALUES(42)');db.commit()
        source=[path,Path(str(path)+'-wal'),Path(str(path)+'-shm')]
        before=[(p.stat().st_mtime_ns,documents.digest(p)) for p in source]
        target=database.snapshot(path,folder)
        assert database.query(target,path.name,table='x')['rows']==[[42]]
        assert before==[(p.stat().st_mtime_ns,documents.digest(p)) for p in source]


def make_video(path, frames=30, fps=10, width=160, height=96, audio=False, gop=10):
    import av
    import numpy as np
    with av.open(str(path),'w') as output:
        stream=output.add_stream('libx264',rate=fps);stream.width=width;stream.height=height;stream.pix_fmt='yuv420p';stream.codec_context.gop_size=gop
        stream.options={'sc_threshold':'0'}
        sound=output.add_stream('aac',rate=16000) if audio else None
        if sound: sound.layout='mono'
        for i in range(frames):
            pixels=np.full((height,width,3),40+i%190,dtype=np.uint8);pixels[:,:,2]=180
            picture=av.VideoFrame.from_ndarray(pixels,format='rgb24')
            picture.pts=i
            for packet in stream.encode(picture):output.mux(packet)
        for packet in stream.encode(None):output.mux(packet)
        if sound:
            samples=int(frames/fps*16000)
            for start in range(0,samples,1024):
                frame=av.AudioFrame.from_ndarray(np.zeros((1,min(1024,samples-start)),dtype=np.float32),format='flt',layout='mono');frame.sample_rate=16000;frame.pts=start
                for packet in sound.encode(frame):output.mux(packet)
            for packet in sound.encode(None):output.mux(packet)


@pytest.mark.parametrize('suffix,audio,width,height', [('.mp4',False,160,96),('.mkv',True,320,180),('.mov',False,640,360)])
def test_video_metadata_frames_keyframes_audio(tmp_path,suffix,audio,width,height):
    path=tmp_path/('video'+suffix);make_video(path,width=width,height=height,audio=audio)
    before=documents.digest(path); info=video.metadata(path)
    assert info['width']==width and info['height']==height and info['audio']==audio
    assert 2.8 < info['duration'] < 3.3
    quick=video.sample(path,tmp_path,info,'quick')
    detailed=video.sample(path,tmp_path,info,'detailed',keyframes=True)
    assert 1<=len(quick['frames'])<=12
    assert all(frame['keyframe'] for frame in detailed['frames'])
    assert all(0<=frame['time']<=info['duration'] for frame in quick['frames'])
    if audio:
        target=video.extract_audio(path,tmp_path,threading.Event());assert target.stat().st_size>0;target.unlink()
    else:
        with pytest.raises(Exception,match='audio track'):video.extract_audio(path,tmp_path,threading.Event())
    assert documents.digest(path)==before


def test_long_video_bounded_sampling_and_cleanup(tmp_path):
    path=tmp_path/'long.mp4';make_video(path,frames=1200,fps=1)
    info=video.metadata(path);assert info['duration']>=1199
    assert len(video.sample(path,tmp_path,info,'quick')['frames'])<=12
    detailed=video.sample(path,tmp_path,info,'detailed')
    assert 90<=len(detailed['frames'])<=120
    custom=video.sample(path,tmp_path,info,'custom',interval=.1)
    assert len(custom['frames'])<=120 and custom['interval']>=10
    for _ in range(10):
        opened=files.open_file(path);directory=files.get(opened['id']).directory
        files.close(opened['id']);assert not directory.exists()
    assert not files._sessions


def test_video_corruption_and_cancel(tmp_path):
    path=tmp_path/'corrupt.mp4';path.write_text('not video')
    with pytest.raises(Exception):files.open_file(path)
    assert not files._sessions
    make_video(path);info=video.metadata(path);cancel=threading.Event();cancel.set()
    with pytest.raises(Exception):video.sample(path,tmp_path,info,cancel=cancel)
    assert not list(tmp_path.glob('*.jpg'))


def test_api_native_grant_save_and_readonly_query(docx,db,monkeypatch):
    app=FastAPI();app.include_router(router);monkeypatch.setenv('LAW_LOCAL_FILES_TOKEN','native-secret')
    with TestClient(app) as client:
        assert client.post('/local-files/open',json={'path':str(docx)}).status_code==403
        headers={'x-local-files':'native-secret'}
        opened=client.post('/local-files/open',headers=headers,json={'path':str(docx)}).json()
        identifier=opened['id']
        assert client.get(f'/local-files/{identifier}/asset/../document.xml').status_code==404
        assert client.post(f'/local-files/{identifier}/save',json={'target':str(docx),'changes':[]}).status_code==403
        assert client.delete(f'/local-files/{identifier}').status_code==200
        opened=client.post('/local-files/open',headers=headers,json={'path':str(db)}).json();identifier=opened['id']
        assert client.post(f'/local-files/{identifier}/query',json={'table':'messages'}).json()['has_more']
        assert client.post(f'/local-files/{identifier}/query',json={'sql':'DROP TABLE messages'}).status_code==400
        assert client.delete(f'/local-files/{identifier}').status_code==200


def test_transcription_adapter_uses_existing_service(monkeypatch,tmp_path):
    from services import audio
    monkeypatch.setattr(audio,'status',lambda:{'models':{'base':{'ready':True}}})
    calls=[]
    monkeypatch.setattr(audio,'transcribe',lambda path,**kw:calls.append((path,kw)) or {'segments':[{'start':.2,'text':'sample'}]})
    result=video.transcribe(tmp_path/'sample.m4a')
    assert result['segments'][0]['start']==.2
    assert calls[0][1]=={'model_size':'base','acceleration':'cpu'}


def test_vision_selected_frames_limit_release_and_cancel(monkeypatch,tmp_path):
    import asyncio
    import httpx
    from services import video_analysis
    from types import SimpleNamespace
    from services.request_queue import RequestQueue, QueueCancelled
    class Coordinator:
        def reserve(self,owner): return True
        def release(self,owner): pass
    work= RequestQueue(Coordinator()); monkeypatch.setattr(video,'queue',work)
    calls=[];unloaded=[]
    async def prepare(kind): calls.append(('prepare',kind))
    async def describe(raw,model,cancel,timestamp,focus): calls.append((raw,model));return 'A red object is visible near the left side.'
    async def summarize(observations,model,cancel,focus,transcript):return 'The sampled video shows a red object. Movement between frames is uncertain.'
    class Client:
        def __init__(self,**kw): pass
        async def __aenter__(self):return self
        async def __aexit__(self,*args):pass
        async def post(self,url,json):
            if url.endswith('/api/show'):return SimpleNamespace(raise_for_status=lambda:None,json=lambda:{'capabilities':['vision']})
            unloaded.append(json);return None
    monkeypatch.setattr(video,'prepare_runtime',prepare)
    monkeypatch.setattr(video_analysis,'describe',describe)
    monkeypatch.setattr(video_analysis,'summarize',summarize)
    monkeypatch.setattr(httpx,'AsyncClient',Client)
    frames=[{'id':f'{i}.jpg','time':float(i),'observation':''} for i in range(5)]
    for frame in frames:(tmp_path/frame['id']).write_bytes(b'fixture')
    result=asyncio.run(video.vision(frames,tmp_path,'test-vision',3,threading.Event(),lambda _:None))
    assert result['status']=='complete' and result['summary'].startswith('The sampled video')
    assert [row['time'] for row in result['observations']]==[0,2,4]
    assert sum(bool(f['observation']) for f in frames)==3
    assert frames[0]['time']==0 and frames[3]['observation']==''
    assert unloaded==[{'model':'test-vision','keep_alive':0}]
    assert work.jobs[-1].status=='completed' and work.active is None
    work.paused=True;cancel=threading.Event();cancel.set()
    with pytest.raises(QueueCancelled):asyncio.run(video.vision(frames,tmp_path,'test-vision',3,cancel,lambda _:None))
    assert work.active is None and work.jobs[-1].status=='cancelled'


def test_video_pipeline_partial_results_and_temporary_audio_cleanup(tmp_path,monkeypatch):
    path=tmp_path/'audio.mp4';make_video(path,audio=True)
    opened=files.open_file(path);item=files.get(opened['id'])
    app=FastAPI();app.include_router(router)
    def unavailable(*args):raise ValueError('Transcription unavailable in test')
    monkeypatch.setattr(video,'transcribe',unavailable)
    async def no_vision(*args):raise ValueError('vision model unavailable in test')
    monkeypatch.setattr(video,'vision',no_vision)
    try:
        with TestClient(app) as client:
            response=client.post(f'/local-files/{item.id}/video',json={'operation':'full','model':'test-model'})
            assert response.status_code==200
            data=response.json()['data']
            assert len(data['frames'])>1
            assert 'unavailable' in data['transcript_warning']
            assert 'vision model' in data['vision_warning']
            assert not list(item.directory.glob('*.m4a'))
            with files.operation(item.id):
                assert client.delete(f'/local-files/{item.id}').status_code==400
                assert item.cancel.is_set()
    finally:files.close(item.id)


def test_document_protected_structures_survive(docx,tmp_path):
    from lxml import etree as ET
    members,root=documents.read_package(docx)
    body=root.find('w:body',documents.NS)
    tracked=ET.SubElement(body[1],documents.W+'ins');run=ET.SubElement(tracked,documents.W+'r');ET.SubElement(run,documents.W+'t').text='Tracked text'
    protected=ET.tostring(tracked)
    source=tmp_path/'tracked.docx'
    with zipfile.ZipFile(source,'w') as archive:
        for entry,data in members:archive.writestr(entry,ET.tostring(root) if entry.filename=='word/document.xml' else data)
    doc=documents.Document(source)
    assert doc.model()['blocks'][1]['runs'][-1]['editable'] is False
    doc.save(source,[change(doc)],None,True)
    _,saved=documents.read_package(source)
    assert ET.tostring(saved.find('.//w:ins',documents.NS))==protected


def test_emphasis_keeps_word_property_order_and_existing_color(docx):
    from docx.shared import RGBColor
    original=Word(docx);original.paragraphs[1].runs[0].font.color.rgb=RGBColor(12,34,56);original.save(docx)
    doc=documents.Document(docx);edit=change(doc);edit.update(bold=True,italic=True)
    doc.save(docx,[edit],None,True)
    saved=Word(docx).paragraphs[1].runs[0]
    assert saved.bold and saved.italic and str(saved.font.color.rgb)=='0C2238'
    assert [node.tag for node in saved._r.rPr]==[documents.W+'b',documents.W+'i',documents.W+'color']
