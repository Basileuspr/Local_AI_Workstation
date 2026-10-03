from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from services.document_editor import export_docx, import_docx, DEFAULT_LAYOUT
from services.document_styles import BASE, DEFAULT_STYLES


def style(**patch):
    return dict(BASE, id='callout', name='Callout', **patch)


def paragraph(text, **attrs):
    return {'type': 'paragraph', 'attrs': {'styleId': 'callout', **attrs}, 'content': [{'type': 'text', 'text': text}]}


def model(spec, *nodes):
    return {'type': 'doc', 'attrs': {'styles': [spec]}, 'content': list(nodes or [paragraph('Styled text')])}


def save(document):
    out = BytesIO(); document.save(out); return out.getvalue()


def test_named_styles_and_explicit_emphasis_off_survive_repeated_docx_round_trips():
    spec = style(fontFamily='Georgia', fontSize=15, bold=True, italic=True, color='#345678', spaceAfter=18, lineSpacing=1.5)
    first, second = paragraph('First'), paragraph('Direct override', spaceAfter=3)
    second['content'][0]['marks'] = [{'type': 'textStyle', 'attrs': {'fontWeight': 'normal', 'fontStyle': 'normal', 'fontSize': '10pt'}}]
    source = model(spec, first, second)
    for _ in range(3):
        raw = export_docx(source, DEFAULT_LAYOUT)
        doc = Document(BytesIO(raw))
        assert doc.paragraphs[0].style.name == 'Callout'
        assert doc.paragraphs[0].style.font.name == 'Georgia'
        assert doc.paragraphs[0].style.font.bold and doc.paragraphs[0].style.font.italic
        assert doc.paragraphs[0].runs[0].font.size is None
        assert doc.paragraphs[0].paragraph_format.space_after is None
        assert doc.paragraphs[1].paragraph_format.space_after.pt == 3
        assert doc.paragraphs[1].runs[0].bold is False and doc.paragraphs[1].runs[0].italic is False
        source = import_docx(raw)['document']
        assert next(item for item in source['attrs']['styles'] if item['id'] == 'callout') == spec
        assert source['content'][0]['attrs']['styleId'] == 'callout'
        direct = source['content'][1]['content'][0]['marks']
        assert next(item for item in direct if item['type'] == 'textStyle')['attrs']['fontWeight'] == 'normal'


def test_editing_one_definition_updates_all_uses_without_flattening_direct_overrides():
    source = model(style(fontSize=14), paragraph('A'), paragraph('B', spaceAfter=4))
    source['attrs']['styles'][0].update(fontSize=18, spaceAfter=20)
    doc = Document(BytesIO(export_docx(source, DEFAULT_LAYOUT)))
    assert all(p.style.font.size.pt == 18 for p in doc.paragraphs)
    assert doc.paragraphs[0].style.paragraph_format.space_after.pt == 20
    assert doc.paragraphs[1].paragraph_format.space_after.pt == 4


def test_import_resolves_supported_style_inheritance_and_keeps_direct_overrides():
    doc = Document()
    base = doc.styles.add_style('Base author style', WD_STYLE_TYPE.PARAGRAPH)
    base.font.name, base.font.size, base.font.bold = 'Arial', Pt(17), True
    base.paragraph_format.space_after = Pt(21)
    base.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    child = doc.styles.add_style('Note & reminder', WD_STYLE_TYPE.PARAGRAPH)
    child.base_style, child.font.italic = base, True
    child.font.color.rgb = RGBColor.from_string('226644')
    p = doc.add_paragraph('Inherited note', child); p.runs[0].bold = False
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    imported = import_docx(save(doc))['document']
    spec = next(item for item in imported['attrs']['styles'] if item['name'] == 'Note & reminder')
    assert (spec['fontFamily'], spec['fontSize'], spec['bold'], spec['italic'], spec['spaceAfter'], spec['textAlign']) == ('Arial', 17, True, True, 21, 'center')
    assert imported['content'][0]['attrs']['textAlign'] == 'right'
    exported = Document(BytesIO(export_docx(imported, DEFAULT_LAYOUT))).paragraphs[0]
    assert exported.style.name == 'Note & reminder' and exported.runs[0].bold is False
    assert exported.style.font.italic and exported.alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_custom_heading_and_unused_custom_styles_survive_and_do_not_change_header_styles():
    unused = style(); unused.update(id='unused', name='Unused style')
    heading = style(level=2); heading['name'] = 'Header'
    source = model(heading, {'type': 'heading', 'attrs': {'styleId': 'callout', 'level': 2}, 'content': [{'type': 'text', 'text': 'Custom section'}]})
    source['attrs']['styles'].append(unused)
    raw = export_docx(source, {**DEFAULT_LAYOUT, 'header': 'Page header'})
    doc = Document(BytesIO(raw))
    assert doc.paragraphs[0].style.style_id != doc.sections[0].header.paragraphs[0].style.style_id
    result = import_docx(raw)['document']
    assert result['content'][0]['type'] == 'heading' and result['content'][0]['attrs']['level'] == 2
    assert any(item['name'] == 'Unused style' for item in result['attrs']['styles'])


@pytest.mark.parametrize('patch', [dict(fontSize=float('nan')), dict(color='red;url(secret)'), dict(bold='true'), dict(level=True), dict(textAlign={}), dict(fontFamily=[]), dict(name='\x00bad'), dict(id='../../x')])
def test_invalid_style_definitions_are_rejected(patch):
    spec = style(); spec.update(patch)
    with pytest.raises(ValueError): export_docx(model(spec), DEFAULT_LAYOUT)


def test_duplicate_names_unknown_references_and_style_limit_are_rejected():
    duplicate = model(style()); duplicate['attrs']['styles'].append({**style(), 'id': 'other', 'name': 'CALLOUT'})
    too_many = model(style()); too_many['attrs']['styles'] = [{**style(), 'id': f'style-{i}', 'name': f'Style {i}'} for i in range(33)]
    unknown = model(style(), paragraph('Unknown', styleId='missing'))
    builtin = model({**deepcopy(DEFAULT_STYLES[0]), 'name': 'Renamed Normal'})
    for source in (duplicate, too_many, unknown, builtin):
        with pytest.raises(ValueError): export_docx(source, DEFAULT_LAYOUT)
