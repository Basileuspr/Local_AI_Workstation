import pytest

from services import file_parser
original_graphics_probe = file_parser._pdf_graphical_pages


@pytest.fixture(autouse=True)
def no_graphics_probe(monkeypatch):
    monkeypatch.setattr(file_parser, '_pdf_graphical_pages', lambda _data: [])


def test_scanned_pdf_ocrs_only_pages_without_a_useful_text_layer(monkeypatch):
    monkeypatch.setattr(
        file_parser,
        "_extract_pdf_page_texts",
        lambda _data: [
            "Native text remains authoritative on this page.",
            "",
            "tiny",
        ],
    )
    seen = []

    def fake_ocr(_data, page_indices):
        seen.extend(page_indices)
        return {
            1: "OCR TRANSCRIPTION PAGE TWO",
            2: "OCR TRANSCRIPTION PAGE THREE",
        }

    monkeypatch.setattr(file_parser, "_ocr_pdf_pages", fake_ocr)

    text, metadata = file_parser._parse_pdf(b"fake-pdf")

    assert seen == [1, 2]
    assert "--- Page 1 ---\nNative text remains authoritative" in text
    assert "--- Page 2 ---\nOCR TRANSCRIPTION PAGE TWO" in text
    assert "--- Page 3 ---\nOCR TRANSCRIPTION PAGE THREE" in text
    assert {key: metadata[key] for key in ('page_count', 'text_layer_pages', 'ocr_pages', 'ocr_model')} == {
        "page_count": 3,
        "text_layer_pages": [1],
        "ocr_pages": [2, 3],
        "ocr_model": file_parser.settings.ocr_model,
    }


def test_text_layer_pdf_does_not_start_the_ocr_runtime(monkeypatch):
    monkeypatch.setattr(
        file_parser,
        "_extract_pdf_page_texts",
        lambda _data: ["First native page with enough text.", "Second native page with enough text."],
    )
    monkeypatch.setattr(
        file_parser,
        "_ocr_pdf_pages",
        lambda *_args: pytest.fail("OCR must not run for text-layer pages"),
    )

    text, metadata = file_parser._parse_pdf(b"fake-pdf")

    assert "First native page" in text
    assert "Second native page" in text
    assert metadata["ocr_pages"] == []
    assert metadata["text_layer_pages"] == [1, 2]


def test_parse_file_exposes_scanned_pdf_metadata(monkeypatch):
    monkeypatch.setattr(
        file_parser,
        "_parse_pdf",
        lambda _data: (
            "--- Page 1 ---\nSCANNED SENTINEL 731",
            {
                "page_count": 1,
                "text_layer_pages": [],
                "ocr_pages": [1],
                "ocr_model": "vision-test",
            },
        ),
    )

    parsed = file_parser.parse_file(b"fake-pdf", "scan.pdf")

    assert parsed["error"] is None
    assert parsed["ocr_pages"] == [1]
    assert parsed["page_count"] == 1
    assert parsed["ocr_model"] == "vision-test"
    assert parsed["char_count"] == len(parsed["text"])


def test_ocr_requires_visible_transcription_instead_of_exposing_thinking(monkeypatch):
    class Response:
        def raise_for_status(self):
            return None

        def json(self):
            return {"message": {"content": "", "thinking": "private reasoning"}}

    monkeypatch.setattr(file_parser.httpx, "post", lambda *_args, **_kwargs: Response())

    with pytest.raises(RuntimeError, match="no visible transcription"):
        file_parser._transcribe_page_png(b"png", page_number=1, keep_alive=0)


def test_ocr_request_is_local_deterministic_and_unloads_after_last_page(monkeypatch):
    captured = {}

    class Response:
        def raise_for_status(self):
            return None

        def json(self):
            return {"message": {"content": "```text\nSCANNED SENTINEL 731\n```"}}

    def fake_post(url, **kwargs):
        captured.update(url=url, **kwargs)
        return Response()

    monkeypatch.setattr(file_parser.httpx, "post", fake_post)

    text = file_parser._transcribe_page_png(b"png", page_number=1, keep_alive=0)

    assert text == "SCANNED SENTINEL 731"
    assert captured["url"].endswith("/api/chat")
    assert captured["json"]["model"] == file_parser.settings.ocr_model
    assert captured["json"]["stream"] is False
    assert captured["json"]["think"] is False
    assert captured["json"]["keep_alive"] == 0
    assert captured["json"]["messages"][0]["images"]


def test_graphics_on_text_rich_pdf_are_described_without_replacing_text(monkeypatch):
    monkeypatch.setattr(file_parser, '_extract_pdf_page_texts', lambda _: ['Native body text with enough content.'])
    monkeypatch.setattr(file_parser, '_pdf_graphical_pages', lambda _: [0])
    monkeypatch.setattr(file_parser, '_render_pdf_page_png', lambda *_: b'png')
    monkeypatch.setattr(file_parser, '_describe_images', lambda images: {index:'Chart: A = 12, B = 25.' for index, _ in images})
    text, metadata = file_parser._parse_pdf(b'pdf')
    assert 'Native body text' in text and 'A = 12' in text
    assert metadata['ocr_pages'] == []
    assert metadata['visual_pages'] == [1]
    assert metadata['graphical_pages'] == [1]


def test_docx_preserves_paragraph_table_image_order_and_exposes_coverage(monkeypatch):
    import io
    from docx import Document
    from PIL import Image
    document = Document()
    document.add_paragraph('Before table')
    table = document.add_table(rows=2, cols=2)
    table.cell(0,0).text = 'Name'; table.cell(0,1).text = 'Value'
    table.cell(1,0).text = 'Alpha'; table.cell(1,1).text = '42'
    png = io.BytesIO(); Image.new('RGB',(32,32),'blue').save(png,format='PNG'); png.seek(0)
    document.add_picture(png)
    document.add_paragraph('After image')
    data = io.BytesIO(); document.save(data)
    monkeypatch.setattr(file_parser, '_describe_images', lambda images: {0:'A blue square.'})
    parsed = file_parser.parse_file(data.getvalue(),'example.docx')
    assert parsed['error'] is None
    text = parsed['text']
    assert text.index('Before table') < text.index('Name | Value') < text.index('Alpha | 42') < text.index('A blue square') < text.index('After image')
    assert parsed['table_count'] == 1 and parsed['image_count'] == 1
    assert parsed['visual_images'] == [1]


def test_visual_failure_keeps_native_text_and_reports_missing_coverage(monkeypatch):
    monkeypatch.setattr(file_parser, '_extract_pdf_page_texts', lambda _: ['A useful native text page remains readable.'])
    monkeypatch.setattr(file_parser, '_pdf_graphical_pages', lambda _: [0])
    monkeypatch.setattr(file_parser, '_render_pdf_page_png', lambda *_: b'png')
    def unavailable(_): raise RuntimeError('offline')
    monkeypatch.setattr(file_parser, '_describe_images', unavailable)
    parsed = file_parser.parse_file(b'pdf','diagram.pdf')
    assert parsed['error'] is None and 'useful native text' in parsed['text']
    assert parsed['visual_pages'] == [] and parsed['warnings']


def test_real_pdf_vector_chart_is_detected_on_a_text_rich_page():
    import io
    from PyPDF2 import PdfWriter
    from PyPDF2.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    page = writer.pages[0]
    font = DictionaryObject({NameObject('/Type'):NameObject('/Font'), NameObject('/Subtype'):NameObject('/Type1'),
                             NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 20 260 Td (Native text with a visible single-path chart.) Tj ET 20 20 m 80 90 l 150 30 l S')
    page[NameObject('/Contents')] = writer._add_object(stream)
    output = io.BytesIO(); writer.write(output)
    data = output.getvalue()
    assert len(file_parser._extract_pdf_page_texts(data)[0]) >= file_parser.settings.ocr_min_page_chars
    assert original_graphics_probe(data) == [0]
    assert file_parser._render_pdf_page_png(data, 0).startswith(b'\x89PNG')

