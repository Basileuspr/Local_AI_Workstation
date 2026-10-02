import base64
from io import BytesIO
import zipfile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi import FastAPI
from fastapi.testclient import TestClient
from PIL import Image

from services.document_editor import DEFAULT_LAYOUT, export_docx, import_docx, validate_model
from routes.document_editor import router


def p(text='', **attrs):
    return {'type': 'paragraph', 'attrs': attrs, 'content': [{'type': 'text', 'text': text}] if text else []}


def doc(*content):
    return {'type': 'doc', 'content': list(content)}


def collect(node, kind):
    return ([node] if node.get('type') == kind else []) + [found for child in node.get('content', []) for found in collect(child, kind)]


def png():
    buffer = BytesIO(); Image.new('RGB', (40, 30), 'blue').save(buffer, 'PNG')
    return 'data:image/png;base64,' + base64.b64encode(buffer.getvalue()).decode()


def test_text_formatting_and_paragraph_settings_round_trip():
    marks = [{'type': n} for n in ['bold', 'italic', 'underline', 'strike', 'superscript']]
    marks += [{'type': 'textStyle', 'attrs': {'fontFamily': 'Georgia', 'fontSize': '18pt', 'color': 'rgb(17, 34, 51)'}},
              {'type': 'highlight', 'attrs': {'color': '#FFFF00'}}]
    node = p('A formatted sentence', textAlign='right', indent=2, spaceBefore=12, spaceAfter=18, lineSpacing=1.5)
    node['content'][0]['marks'] = marks
    raw = export_docx(doc(node), DEFAULT_LAYOUT)
    word = Document(BytesIO(raw)); run = word.paragraphs[0].runs[0]
    assert run.bold and run.italic and run.underline and run.font.strike and run.font.superscript
    assert run.font.name == 'Georgia' and run.font.size.pt == 18
    assert str(run.font.color.rgb) == '112233'
    assert word.paragraphs[0].paragraph_format.left_indent.inches == .5
    imported = import_docx(raw)['document']['content'][0]
    assert imported['attrs']['lineSpacing'] == 1.5
    assert imported['attrs']['spaceAfter'] == 18
    assert imported['content'][0]['text'] == 'A formatted sentence'
    assert {m['type'] for m in imported['content'][0]['marks']} >= {'bold', 'textStyle', 'highlight'}


def test_headings_links_and_manual_breaks_survive():
    model = doc({'type': 'heading', 'attrs': {'level': 2}, 'content': [{'type': 'text', 'text': 'Chapter'}]},
                {'type': 'paragraph', 'content': [{'type': 'text', 'text': 'Link', 'marks': [{'type': 'link', 'attrs': {'href': 'https://example.com'}}]}, {'type': 'hardBreak'}, {'type': 'text', 'text': 'Second line'}]},
                {'type': 'pageBreak'}, p('Next page'))
    result = import_docx(export_docx(model, DEFAULT_LAYOUT))['document']
    assert result['content'][0]['type'] == 'heading' and result['content'][0]['attrs']['level'] == 2
    assert len(collect(result, 'pageBreak')) == 1 and len(collect(result, 'hardBreak')) == 1
    assert any(mark['attrs']['href'] == 'https://example.com' for text in collect(result, 'text') for mark in text.get('marks', []) if mark['type'] == 'link')


def test_lists_restart_and_nested_mixed_lists_round_trip():
    inner = {'type': 'bulletList', 'content': [{'type': 'listItem', 'content': [p('Nested bullet')]}]}
    first = {'type': 'orderedList', 'attrs': {'start': 3}, 'content': [{'type': 'listItem', 'content': [p('Three'), inner]}, {'type': 'listItem', 'content': [p('Four')]}]}
    second = {'type': 'orderedList', 'attrs': {'start': 1}, 'content': [{'type': 'listItem', 'content': [p('Restart')]}]}
    result = import_docx(export_docx(doc(first, p('Separator'), second), DEFAULT_LAYOUT))['document']
    assert [node['attrs']['start'] for node in collect(result, 'orderedList')] == [3, 1]
    assert collect(result, 'bulletList')[0]['content'][0]['content'][0]['content'][0]['text'] == 'Nested bullet'
    assert len(result['content'][0]['content']) == 2


def test_table_picture_alt_text_and_layout_round_trip():
    table = {'type': 'table', 'content': [{'type': 'tableRow', 'content': [{'type': 'tableHeader' if row == 0 else 'tableCell', 'content': [p(f'{row}:{column}')]} for column in range(2)]} for row in range(3)]}
    model = doc(table, {'type': 'image', 'attrs': {'src': png(), 'alt': 'Blue rectangle', 'width': 96}})
    layout = {**DEFAULT_LAYOUT, 'paper': 'A4', 'orientation': 'landscape', 'left': .5, 'right': .75}
    raw = export_docx(model, layout); original = Document(BytesIO(raw))
    assert len(original.tables) == 1 and len(original.tables[0].rows) == 3
    assert original.tables[0].cell(2, 1).text == '2:1'
    assert len(original.inline_shapes) == 1 and original.inline_shapes[0].width.inches == 1
    result = import_docx(raw)
    assert result['layout']['paper'] == 'A4' and result['layout']['orientation'] == 'landscape'
    assert result['layout']['left'] == .5
    assert collect(result['document'], 'image')[0]['attrs']['alt'] == 'Blue rectangle'
    assert len(collect(result['document'], 'tableRow')) == 3


def test_import_warns_about_omitted_sections_comments_and_fields():
    source = Document(); source.add_paragraph('Visible body')
    source.sections[0].header.paragraphs[0].text = 'Header not editable'
    source.add_section()
    run = source.add_paragraph().add_run(); field = OxmlElement('w:fldChar'); field.set(qn('w:fldCharType'), 'begin'); run._r.append(field)
    buffer = BytesIO(); source.save(buffer); raw = buffer.getvalue()
    result = import_docx(raw)
    assert any('Headers' in w for w in result['warnings'])
    assert any('section' in w for w in result['warnings'])
    assert any('fields' in w for w in result['warnings'])
    assert raw == buffer.getvalue()  # conversion only reads the source bytes


def test_merged_cells_are_explicitly_flattened_without_duplicating_text():
    source = Document(); table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = 'Merged text'
    buffer = BytesIO(); source.save(buffer)
    result = import_docx(buffer.getvalue())
    assert any('Merged' in w for w in result['warnings'])
    assert sum(node['text'].count('Merged text') for node in collect(result['document'], 'text')) == 1


@pytest.mark.parametrize('href', ['javascript:alert(1)', 'file:///C:/secret', 'data:text/html,abc', 'https://example.com\nsecret'])
def test_export_rejects_unsafe_links(href):
    node = p('Click'); node['content'][0]['marks'] = [{'type': 'link', 'attrs': {'href': href}}]
    with pytest.raises(ValueError, match='Links'): export_docx(doc(node), DEFAULT_LAYOUT)


@pytest.mark.parametrize('src', ['https://example.com/x.png', 'file:///C:/x.png', 'data:image/svg+xml;base64,PHN2Zz4='])
def test_export_never_fetches_external_or_active_images(src):
    with pytest.raises(ValueError): export_docx(doc({'type': 'image', 'attrs': {'src': src}}), DEFAULT_LAYOUT)


@pytest.mark.parametrize('node', [
    {'type': 'script'}, {'type': 'paragraph', 'content': [{'type': 'table'}]},
    {'type': 'table', 'content': [{'type': 'tableRow', 'content': []}]},
    {'type': 'paragraph', 'attrs': {'lineSpacing': float('nan')}},
    {'type': 'paragraph', 'content': [{'type': 'text', 'text': '\x00'}]},
])
def test_invalid_editor_models_fail_without_output(node):
    with pytest.raises(ValueError): validate_model(doc(node))


def test_rejects_hostile_zip_xml_and_macros():
    raw = export_docx(doc(p('safe')), DEFAULT_LAYOUT)
    with zipfile.ZipFile(BytesIO(raw)) as source:
        members = {name: source.read(name) for name in source.namelist()}
    for change in [{'word/vbaProject.bin': b'code'}, {'word/document.xml': b'<!DOCTYPE x [<!ENTITY e SYSTEM "file:///C:/secret">]><x>&e;</x>'}]:
        target = BytesIO()
        with zipfile.ZipFile(target, 'w') as package:
            for name, content in {**members, **change}.items(): package.writestr(name, content)
        with pytest.raises(ValueError): import_docx(target.getvalue())


def test_conversion_routes_return_real_docx_and_validation_errors():
    app = FastAPI(); app.include_router(router); client = TestClient(app)
    response = client.post('/document-editor/export', json={'document': doc(p('Route round trip')), 'layout': DEFAULT_LAYOUT})
    assert response.status_code == 200 and response.content.startswith(b'PK')
    assert response.headers['cache-control'] == 'no-store'
    imported = client.post('/document-editor/import', content=response.content)
    assert imported.status_code == 200
    assert imported.json()['document']['content'][0]['content'][0]['text'] == 'Route round trip'
    assert client.post('/document-editor/export', json={'document': {'type': 'bad'}}).status_code == 400
    assert client.post('/document-editor/import', content=b'not a docx').status_code == 400


def test_conversion_routes_obey_the_app_session_guard():
    from services.session_guard import SessionGuard
    app = FastAPI(); app.include_router(router); app.add_middleware(SessionGuard)
    client = TestClient(app, base_url='http://127.0.0.1:8000')
    payload = {'document': doc(p('Guarded conversion'))}
    assert client.post('/document-editor/export', json=payload).status_code == 403
    assert client.post('/document-editor/export', headers={'X-LAW-Session': 'test-session-token'}, json=payload).status_code == 200
    assert client.post('/document-editor/import', headers={'Origin': 'https://untrusted.example', 'X-LAW-Session': 'test-session-token'}, content=b'anything').status_code == 403


@pytest.mark.parametrize('payload', [
    {'document': {'type': []}}, {'document': doc(p('A')), 'layout': {'paper': []}},
    {'document': doc({'type': 'paragraph', 'content': [{'type': []}]})},
])
def test_malformed_types_receive_a_clear_client_error(payload):
    app = FastAPI(); app.include_router(router)
    assert TestClient(app).post('/document-editor/export', json=payload).status_code == 400
