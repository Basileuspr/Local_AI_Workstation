from io import BytesIO

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from services.document_editor import DEFAULT_LAYOUT, export_docx, import_docx
from services.document_page_layout import STORIES, PAGINATION, element


MODEL = {'type': 'doc', 'content': [{'type': 'paragraph', 'content': [{'type': 'text', 'text': 'Page layout sample'}]}]}


def saved(document):
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def field(paragraph, instruction, cached='7', complex=False):
    if complex:
        paragraph.add_run()._r.append(element('fldChar', fldCharType='begin'))
        # Instructions can be split across runs in real Word files.
        for part in (instruction[:2], instruction[2:]):
            code = element('instrText'); code.text = part
            paragraph.add_run()._r.append(code)
        paragraph.add_run()._r.append(element('fldChar', fldCharType='separate'))
        paragraph.add_run(cached)
        paragraph.add_run()._r.append(element('fldChar', fldCharType='end'))
    else:
        node = element('fldSimple', instr=instruction)
        run, text = element('r'), element('t'); text.text = cached
        run.append(text); node.append(run); paragraph._p.append(node)


@pytest.mark.parametrize('position', ['header', 'footer'])
@pytest.mark.parametrize('form', ['decimal', 'lowerRoman', 'upperRoman', 'lowerLetter', 'upperLetter'])
def test_variants_number_formats_and_page_totals_survive_repeated_conversion(position, form):
    layout = {**DEFAULT_LAYOUT, 'differentFirstPage': True, 'differentOddEven': True,
              'headerDistance': .35, 'footerDistance': .65, 'pageNumbers': True,
              'pageNumberStyle': 'pageOf', 'pageNumberFormat': form, 'pageNumberStart': 4,
              'pageNumberPosition': position, 'pageNumberAlignment': 'right', 'pageNumberFirstPage': False,
              **{key: f'{key} title\nSecond line\tTab' for key in STORIES},
              **{key + 'Alignment': 'center' for key in STORIES}}
    raw = export_docx(MODEL, layout)
    word = Document(BytesIO(raw))
    section = word.sections[0]
    assert section.different_first_page_header_footer
    assert word.settings.odd_and_even_pages_header_footer
    assert section.header_distance.inches == pytest.approx(.35)
    assert section.footer_distance.inches == pytest.approx(.65)
    assert section._sectPr.find(qn('w:pgNumType')).get(qn('w:fmt')) == form
    assert section._sectPr.find(qn('w:pgNumType')).get(qn('w:start')) == '4'
    assert not getattr(section, 'first_page_' + position)._element.xpath('.//w:fldSimple')
    for attribute in (position, 'even_page_' + position):
        story = getattr(section, attribute)
        assert [node.get(qn('w:instr')) for node in story._element.xpath('.//w:fldSimple')] == ['PAGE', 'NUMPAGES']
        assert story.paragraphs[-1].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    for _ in range(3):
        result = import_docx(raw)
        for key, value in layout.items():
            assert result['layout'][key] == value, key
        assert not any('cached' in warning or 'normalized' in warning for warning in result['warnings'])
        raw = export_docx(result['document'], result['layout'])


@pytest.mark.parametrize('complex', [False, True])
@pytest.mark.parametrize('display', ['number', 'page', 'pageOf'])
def test_imports_real_fields_without_copying_cached_numbers_into_text(complex, display):
    source = Document()
    source.add_paragraph('Body')
    footer = source.sections[0].footer
    footer.paragraphs[0].text = 'A separate footer\nSecond line'
    number = footer.add_paragraph()
    number.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if display != 'number': number.add_run('Page ')
    field(number, ' PAGE \\* roman \\* MERGEFORMAT ', 'vii', complex)
    if display == 'pageOf':
        number.add_run(' of ')
        field(number, ' NUMPAGES ', '15', complex)
    result = import_docx(saved(source))
    assert result['layout']['footer'] == 'A separate footer\nSecond line'
    assert result['layout']['pageNumberFormat'] == 'lowerRoman'
    assert result['layout']['pageNumberStyle'] == display
    assert result['layout']['pageNumbers']
    rebuilt = Document(BytesIO(export_docx(result['document'], result['layout'])))
    assert 'vii' not in rebuilt.sections[0].footer._element.xml
    assert len(rebuilt.sections[0].footer._element.xpath('.//w:fldSimple')) == (2 if display == 'pageOf' else 1)


def test_custom_and_unknown_fields_keep_cached_text_with_a_conversion_warning():
    source = Document(); source.add_paragraph('Body')
    footer = source.sections[0].footer.paragraphs[0]
    footer.add_run('Revision ')
    field(footer, 'PAGE', '7', True)
    footer.add_run(' / ')
    field(footer, 'DOCPROPERTY Company', 'Example company', True)
    result = import_docx(saved(source))
    assert result['layout']['footer'] == 'Revision 7 / Example company'
    assert result['layout']['pageNumbers'] is False
    assert any('cached text' in warning for warning in result['warnings'])


def test_conflicting_number_lines_are_reported_and_disabled_variant_text_is_retained():
    source = Document(); source.add_paragraph('Body')
    section = source.sections[0]
    section.first_page_header.paragraphs[0].text = 'Hidden title, retained'
    field(section.header.paragraphs[0], 'PAGE')
    footer = section.footer.paragraphs[0]; footer.add_run('Page '); field(footer, 'PAGE')
    result = import_docx(saved(source))
    assert not result['layout']['differentFirstPage']
    assert result['layout']['firstHeader'] == 'Hidden title, retained'
    assert any('normalized' in warning for warning in result['warnings'])
    again = import_docx(export_docx(result['document'], result['layout']))
    assert again['layout']['firstHeader'] == 'Hidden title, retained'


@pytest.mark.parametrize('patch', [
    {'firstHeader': 'x' * 2001}, {'evenFooter': '\x00'}, {'differentFirstPage': 1},
    {'headerDistance': -1}, {'footerDistance': float('nan')}, {'headerDistance': '0.5'},
    {'pageNumberStart': 0}, {'pageNumberStart': 1.5}, {'pageNumberStart': True}, {'pageNumberStart': 10000},
    {'pageNumberPosition': 'body'}, {'pageNumberAlignment': 'justify'}, {'pageNumberStyle': 'custom'},
    {'pageNumberFormat': 'unknown'}, {'firstHeaderAlignment': 'justify'}, {'pageNumberFirstPage': 'yes'},
])
def test_invalid_page_controls_fail_before_export(patch):
    with pytest.raises(ValueError): export_docx(MODEL, {**DEFAULT_LAYOUT, **patch})


def test_pagination_keeps_explicit_false_and_resolves_inherited_styles():
    source = Document()
    base = source.styles.add_style('Keep grouped', WD_STYLE_TYPE.PARAGRAPH)
    base.paragraph_format.keep_with_next = True
    base.paragraph_format.keep_together = True
    child = source.styles.add_style('Derived group', WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = base
    paragraph = source.add_paragraph('Custom style text', 'Derived group')
    paragraph.paragraph_format.page_break_before = True
    paragraph.paragraph_format.keep_with_next = False
    paragraph.paragraph_format.widow_control = False
    result = import_docx(saved(source))
    attrs = result['document']['content'][0]['attrs']
    assert {key: attrs[key] for key in PAGINATION} == dict(pageBreakBefore=True, keepWithNext=False, keepTogether=True, widowControl=False)
    exported = Document(BytesIO(export_docx(result['document'], result['layout'])))
    fmt = exported.paragraphs[0].paragraph_format
    assert fmt.page_break_before is True and fmt.keep_with_next is False
    assert fmt.keep_together is True and fmt.widow_control is False


@pytest.mark.parametrize('value', ['yes', 0, [], {}])
def test_invalid_pagination_flags_are_not_silently_coerced(value):
    model = {'type': 'doc', 'content': [{'type': 'paragraph', 'attrs': {'keepWithNext': value}}]}
    with pytest.raises(ValueError, match='page-break'): export_docx(model, DEFAULT_LAYOUT)
