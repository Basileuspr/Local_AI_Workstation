from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn

from services.document_editor import DEFAULT_LAYOUT, export_docx, import_docx, validate_model, safe_link
from services.document_references import element, walk, internal_link
from services.document_styles import BASE


def text(value, href=None):
    return {'type': 'text', 'text': value, **({'marks': [{'type': 'link', 'attrs': {'href': href}}]} if href else {})}


def p(*content, **attrs):
    return {'type': 'paragraph', 'attrs': attrs, 'content': list(content)}


def bookmark(name):
    return {'type': 'bookmark', 'attrs': {'name': name}}


def doc(*content):
    return {'type': 'doc', 'content': list(content)}


def toc(level=3):
    return {'type': 'tableOfContents', 'attrs': {'title': 'Contents', 'maxLevel': level}}


def save(document):
    stream = BytesIO(); document.save(stream); return stream.getvalue()


def test_native_bookmarks_internal_hyperlinks_and_live_contents_survive_three_round_trips():
    identifier = 'LAW_H' + 'a' * 32
    source = doc(toc(2), p(text('Overview'), styleId='heading-1', referenceId=identifier),
                 p(text('Before '), bookmark('Details'), text('after')), p(text('See details', '#Details'), text(' and overview', '#' + identifier)),
                 p(text('Deep heading'), styleId='heading-3'))
    for _ in range(3):
        package = export_docx(source, DEFAULT_LAYOUT)
        native = Document(BytesIO(package))
        starts = native.element.xpath('.//w:bookmarkStart')
        assert len({start.get(qn('w:id')) for start in starts}) == len(starts)
        assert len(native.element.xpath('.//w:bookmarkEnd')) == len(starts)
        assert 'Details' in [start.get(qn('w:name')) for start in starts]
        anchors = native.element.xpath('.//w:hyperlink/@w:anchor')
        assert anchors.count(identifier) == 2 and 'Details' in anchors
        assert not any(rel.reltype.endswith('/hyperlink') for rel in native.part.rels.values())
        assert not native.element.xpath('.//w:instrText[contains(text(), "TOC")]')
        entries = native.element.xpath('.//w:sdtContent/w:p/w:hyperlink/w:r/w:t/text()')
        assert entries == ['Overview']
        imported = import_docx(package)
        assert not any('missing destination' in warning.lower() for warning in imported['warnings'])
        source = imported['document']
        assert source['content'][0] == toc(2)
        assert source['content'][1]['attrs']['referenceId'] == identifier
        assert [node['attrs']['name'] for node in walk(source) if node['type'] == 'bookmark'] == ['Details']
        assert ''.join(node.get('text', '') for node in walk(source)).count('Overview') == 1


def test_contents_regenerates_text_and_order_from_custom_headings_in_lists_and_tables():
    source = doc(toc(), p(text('First'), styleId='heading-1'),
                 {'type': 'bulletList', 'content': [{'type': 'listItem', 'content': [p(text('Custom list'), styleId='chapter')]}]},
                 {'type': 'table', 'content': [{'type': 'tableRow', 'content': [{'type': 'tableCell', 'content': [p(text('In a cell'), styleId='heading-2')]}]}]})
    source['attrs'] = {'styles': [dict(BASE, id='chapter', name='Chapter', level=2)]}
    source = import_docx(export_docx(source, DEFAULT_LAYOUT))['document']
    source['content'][1]['content'][0]['text'] = 'Revised'
    native = Document(BytesIO(export_docx(source, DEFAULT_LAYOUT)))
    assert native.element.xpath('.//w:sdtContent/w:p/w:hyperlink/w:r/w:t/text()') == ['Revised', 'Custom list', 'In a cell']
    assert len(set(native.element.xpath('.//w:bookmarkStart/@w:name'))) == 3


def test_imports_external_word_anchor_and_collapses_range_to_point_without_losing_text():
    native = Document(); paragraph = native.add_paragraph()
    paragraph._p.append(element('bookmarkStart', id=10, name='_Ref123'))
    paragraph.add_run('Range text'); paragraph._p.append(element('bookmarkEnd', id=10))
    jump = native.add_paragraph(); internal_link(jump, jump.add_run('Jump'), '_Ref123')
    result = import_docx(save(native))
    assert result['document']['content'][0]['content'][0] == bookmark('_Ref123')
    assert any('ranges were converted to point' in warning for warning in result['warnings'])
    assert result['document']['content'][1]['content'][0]['marks'][-1]['attrs']['href'] == '#_Ref123'
    assert not any('missing destination' in warning.lower() for warning in result['warnings'])


def test_missing_destination_is_preserved_and_reported_instead_of_silently_unlinking():
    result = import_docx(export_docx(doc(p(text('Lost link', '#Removed'))), DEFAULT_LAYOUT))
    assert result['document']['content'][0]['content'][0]['marks'][-1]['attrs']['href'] == '#Removed'
    assert any('missing destinations' in warning for warning in result['warnings'])


def test_foreign_contents_and_other_content_controls_keep_visible_supported_text():
    native = Document(); sdt = element('sdt'); content = element('sdtContent'); sdt.append(content)
    paragraph = native.add_paragraph('Cached foreign contents entry'); content.append(paragraph._p)
    native.element.body.insert(0, sdt)
    result = import_docx(save(native))
    assert result['document']['content'][0]['content'][0]['text'] == 'Cached foreign contents entry'
    assert any('live fields are not retained' in warning for warning in result['warnings'])


def test_invalid_and_duplicate_bookmarks_do_not_prevent_importing_visible_text():
    native = Document(); paragraph = native.add_paragraph('Retained text')
    for index, name in enumerate(['Valid', 'Valid', 'bad name']):
        paragraph._p.extend([element('bookmarkStart', id=index, name=name), element('bookmarkEnd', id=index)])
    result = import_docx(save(native))
    assert [n['attrs']['name'] for n in walk(result['document']) if n['type'] == 'bookmark'] == ['Valid']
    assert any('Invalid or duplicate' in warning for warning in result['warnings'])


@pytest.mark.parametrize('name', ['', 'bad name', '1First', 'a' * 41, 'a/b', '<tag>', 'LAW_H' + 'a' * 32, ['not a string']])
def test_rejects_invalid_or_reserved_bookmark_names(name):
    with pytest.raises(ValueError, match='bookmark name'): validate_model(doc(p(bookmark(name))))


@pytest.mark.parametrize('attrs', [{'title': ''}, {'title': 'a' * 101}, {'title': 'bad\x00'}, {'maxLevel': True}, {'maxLevel': 4}, {'maxLevel': '2'}])
def test_rejects_invalid_contents_settings(attrs):
    node = toc(); node['attrs'].update(attrs)
    with pytest.raises(ValueError, match='contents'): validate_model(doc(node))


def test_rejects_duplicate_names_counts_and_nested_contents():
    cases = [doc(p(bookmark('Same'), bookmark('same'))), doc(toc(), toc()),
             doc(p(*[bookmark(f'Mark{i}') for i in range(101)])),
             doc({'type': 'bulletList', 'content': [{'type': 'listItem', 'content': [p(text('A')), toc()]}]}),
             doc(p(text('A'), referenceId='not-a-heading-id'))]
    for source in cases:
        with pytest.raises(ValueError): validate_model(source)


def test_export_does_not_mutate_the_supplied_model():
    source = doc(toc(), p(text('Title'), styleId='heading-1')); before = deepcopy(source)
    export_docx(source, DEFAULT_LAYOUT)
    assert source == before


@pytest.mark.parametrize('href', ['#bad name', '#', '#a/b', 'javascript:alert(1)', 'file:///C:/test', 'https://bad\x00value'])
def test_unsafe_link_schemes_and_fragments_remain_rejected(href):
    assert not safe_link(href)
    with pytest.raises(ValueError, match='Links'): validate_model(doc(p(text('Link', href))))
