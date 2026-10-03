import asyncio
from pathlib import Path
import json
import threading
import time

import httpx
from PIL import Image
from PyPDF2 import PdfWriter
from PyPDF2.generic import DictionaryObject, NameObject, DecodedStreamObject
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from services import folder_review as service
from services import folder_review_reader as reader
from services.request_queue import RequestQueue
from services.session_guard import SessionGuard
from routes.folder_review import router

OPTIONS = {'recursive': True, 'max_entries': 2000, 'max_bytes': 16 * 1024 * 1024, 'max_chars': 200000, 'batch_chars': 500}


def pdf(path, text_pages):
    writer = PdfWriter()
    for text in text_pages:
        writer.add_blank_page(width=300, height=300)
        page = writer.pages[-1]
        if text:
            font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
            stream = DecodedStreamObject()
            stream.set_data(b'BT /F1 12 Tf 20 200 Td (' + text.encode('ascii') + b') Tj ET')
            page[NameObject('/Contents')] = writer._add_object(stream)
    with path.open('wb') as handle:
        writer.write(handle)


def test_native_pdf_pages_only_with_explicit_scanned_page_gaps(tmp_path, monkeypatch):
    from services import file_parser
    monkeypatch.setattr(file_parser, '_ocr_pdf_pages', lambda *args: pytest.fail('OCR was requested'))
    file = tmp_path / 'mixed.pdf'; pdf(file, ['Readable project requirements', '', 'Additional implementation notes'])
    result = reader.extract(file.read_bytes(), file.name, 200000)
    assert result['status'] == 'readable'
    assert 'Readable project requirements' in result['text'] and '--- Page 3 ---' in result['text']
    assert result['metadata']['pages_without_extractable_text'] == [2]
    assert result['metadata']['ocr_attempted'] is False
    assert result['coverage']['partial']
    pdf(file, ['', ''])
    result = reader.extract(file.read_bytes(), file.name, 200000)
    assert result['status'] == 'unreadable' and not result['text']
    assert result['metadata']['pages_without_extractable_text'] == [1, 2]


def test_image_metadata_does_not_decode_pixels_or_return_image_text(tmp_path, monkeypatch):
    path = tmp_path / 'photo.png'; Image.new('RGB', (37, 19)).save(path)
    monkeypatch.setattr(Image.Image, 'load', lambda *args: pytest.fail('Pixels were decoded'))
    result = reader.extract(path.read_bytes(), path.name, 200000)
    assert result['status'] == 'metadata_only' and result['text'] == ''
    assert result['metadata']['width'] == 37 and result['metadata']['height'] == 19
    assert result['metadata']['format'] == 'PNG'


def test_svg_text_is_not_treated_as_readable_image_content():
    raw = b'<svg width="100" height="50"><text>Do not analyze this image text</text></svg>'
    result = reader.extract(raw, 'image.svg', 200000)
    assert not result['text'] and result['metadata']['width'] == '100'


def test_text_batches_cover_every_character_and_report_line_ranges():
    text = ''.join(f'Line {index}: ' + 'x' * 73 + '\n' for index in range(100)) + 'Final marker'
    batches = list(reader.text_batches(text, 500))
    assert ''.join(item['text'] for item in batches) == text
    assert all(len(item['text']) <= 500 for item in batches)
    assert batches[0]['first_line'] == 1 and batches[-1]['last_line'] == 101


def test_text_limits_encoding_and_static_python_structure():
    text = 'raise RuntimeError("Do not execute this source")\n\ndef double(value):\n    return value * 2\n'
    result = reader.extract(text.encode(), 'tools.py', 1000)
    assert result['metadata']['functions'] == [{'name': 'double', 'line': 3}]
    assert result['text'] == text
    result = reader.extract(('One two\n' * 1000).encode('utf-16'), 'notes.txt', 1000)
    assert result['coverage']['partial'] and result['coverage']['original_characters'] == 8000
    assert len(result['text']) == 1000
    with pytest.raises(ValueError, match='Binary'):
        reader.extract(b'a\0b', 'binary.txt', 1000)


def test_source_change_and_byte_limits_are_reported(tmp_path):
    path = tmp_path / 'growing.txt'; path.write_bytes(b'abc')
    def change():
        path.write_bytes(b'changed source')
    with pytest.raises(ValueError, match='changed|grew'):
        reader.read_source(path, tmp_path.resolve(), 1000, change)
    with pytest.raises(ValueError, match='byte limit'):
        reader.read_source(path, tmp_path.resolve(), 2, lambda: None)


def test_sources_read_only_and_all_text_batches_reviewed_with_images_excluded(tmp_path):
    root = tmp_path / 'source'; root.mkdir()
    (root / 'code.py').write_text('"""Project utilities."""\n' + '\ndef double(x):\n    return x * 2\n')
    text = ''.join('Fact %d\n' % index for index in range(160)) + 'THE FINAL FACT'
    (root / 'facts.txt').write_text(text)
    (root / 'guide.md').write_text('# Setup\nInstall the project, then run the local app.\n')
    Image.new('RGB', (20, 10)).save(root / 'cat.png')
    pdf(root / 'scan.pdf', [''])
    (root / '.env').write_text('PRIVATE=do-not-read')
    (root / 'node_modules').mkdir(); (root / 'node_modules/ignored.py').write_text('not reviewed')
    (root / 'nested').mkdir(); (root / 'nested/detail.txt').write_text('Nested information')
    before = {path.relative_to(root): path.read_bytes() for path in root.rglob('*') if path.is_file()}
    prompts = []
    async def summarize(prompt):
        prompts.append(prompt)
        return 'Observed source facts with file references.'
    manager = service.FolderReview(tmp_path / 'reports/reviews.sqlite3', summarize)
    async def execute():
        started = await manager.start(str(root), 'test-model', OPTIONS)
        await manager._task
        return manager.get(started['id']), manager.entries(started['id'], 0, 100)['items']
    review, items = asyncio.run(execute())
    assert review['status'] == 'completed_with_gaps'
    by_path = {item['path']: item for item in items}
    assert by_path['cat.png']['status'] == 'metadata_only'
    assert by_path['scan.pdf']['status'] == 'unreadable'
    assert by_path['.env']['status'] == 'excluded'
    assert by_path['node_modules']['status'] == 'excluded'
    assert by_path['nested/detail.txt']['status'] == 'reviewed'
    assert 'cat.png' not in '\n'.join(prompts) and 'PRIVATE=' not in '\n'.join(prompts)
    source_calls = [prompt for prompt in prompts if prompt.startswith('Review source file "facts.txt"')]
    slices = [json.loads(prompt.split('\n', 1)[1])['source'] for prompt in source_calls]
    assert ''.join(slices) == text
    assert by_path['facts.txt']['coverage']['batches_completed'] == len(slices)
    assert '## Per-file findings' in review['report'] and 'cat.png' in review['report']
    assert {path.relative_to(root): path.read_bytes() for path in root.rglob('*') if path.is_file()} == before
    reloaded = service.FolderReview(manager.database)
    assert reloaded.get(review['id'])['report'] == review['report']


def test_images_only_folder_never_calls_model(tmp_path):
    root = tmp_path / 'pictures'; root.mkdir(); Image.new('RGB', (10, 10)).save(root / 'image.png')
    async def forbidden(prompt):
        pytest.fail('Image metadata was sent to a model')
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3', forbidden)
    async def execute():
        started = await manager.start(str(root), 'text-model', OPTIONS); await manager._task
        return manager.get(started['id'])
    result = asyncio.run(execute())
    assert result['status'] == 'completed' and 'No readable text' in result['report']


def test_non_recursive_and_entry_caps_preserve_explicit_coverage(tmp_path):
    root = tmp_path / 'source'; root.mkdir(); (root / 'nested').mkdir(); (root / 'nested/child.txt').write_text('Child')
    (root / 'one.txt').write_text('One'); (root / 'two.txt').write_text('Two')
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    async def execute():
        first = await manager.start(str(root), '', {**OPTIONS, 'recursive': False}); await manager._task
        second = await manager.start(str(root), '', {**OPTIONS, 'max_entries': 1}); await manager._task
        return manager.get(first['id']), manager.entries(first['id'])['items'], manager.get(second['id'])
    first, entries, second = asyncio.run(execute())
    assert 'nested/child.txt' not in {item['path'] for item in entries}
    assert first['counts']['excluded'] == 1
    assert second['total'] == 1 and not second['inventory_complete']
    assert second['status'] == 'completed_with_gaps'


def test_failed_model_keeps_completed_findings_and_marks_unread_files(tmp_path):
    root = tmp_path / 'source'; root.mkdir(); (root / 'a.txt').write_text('First'); (root / 'b.txt').write_text('Second')
    count = 0
    async def failure(prompt):
        nonlocal count
        count += 1
        if count == 2:
            raise RuntimeError('Provider unavailable')
        return 'First file finding.'
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3', failure)
    async def execute():
        started = await manager.start(str(root), 'test-model', OPTIONS); await manager._task
        return manager.get(started['id']), manager.entries(started['id'])['items']
    result, entries = asyncio.run(execute())
    assert result['status'] == 'failed' and result['processed'] == 1
    assert entries[0]['status'] == 'reviewed' and entries[1]['status'] == 'not_reviewed'
    assert 'First file finding.' in result['report'] and 'Provider unavailable' in result['error']


def test_cancellation_closes_provider_before_releasing_queue_and_keeps_results(tmp_path, monkeypatch):
    class Coordinator:
        def reserve(self, owner): return True
        def release(self, owner): pass
    local_queue = RequestQueue(Coordinator())
    monkeypatch.setattr(service, 'queue', local_queue)
    async def prepare(kind): pass
    async def handoff(*args, **kwargs):
        yield {'detail': 'Model ready'}
    monkeypatch.setattr(service, 'prepare_runtime', prepare)
    monkeypatch.setattr(service, 'prepare_chat_model', handoff)
    async def context_limit(model): return service.settings.num_ctx
    monkeypatch.setattr(service, 'model_limit', context_limit)
    root = tmp_path / 'source'; root.mkdir(); (root / 'a.txt').write_text('First'); (root / 'b.txt').write_text('Second')
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    async def execute():
        running, exited = asyncio.Event(), asyncio.Event()
        calls = 0
        async def model(*args):
            nonlocal calls
            calls += 1
            if calls == 1: return 'First finding.'
            running.set()
            try: await asyncio.Future()
            finally: exited.set()
        monkeypatch.setattr(manager, 'call_model', model)
        async def offload(key, model):
            assert exited.is_set() and local_queue.active is manager._queue_job
            manager.record_processing(key, model_released=True)
        monkeypatch.setattr(manager, 'release_model', offload)
        started = await manager.start(str(root), 'test-model', OPTIONS)
        await asyncio.wait_for(running.wait(), 5)
        with pytest.raises(RuntimeError, match='already running'):
            await manager.start(str(root), 'other-model', OPTIONS)
        await manager.cancel(started['id']); await manager._task
        assert exited.is_set() and local_queue.active is None
        assert all(job.status in {'completed', 'cancelled'} for job in local_queue.jobs)
        return manager.get(started['id'])
    result = asyncio.run(execute())
    assert result['status'] == 'cancelled' and result['counts']['reviewed'] == 1


def test_backend_restart_marks_unfinished_review_interrupted(tmp_path):
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    with manager.connect() as db:
        db.execute("INSERT INTO reviews(id,root,model,options,status,phase,started_at) VALUES('unfinished','source','model','{}','running','Reading','date')")
    reloaded = service.FolderReview(manager.database)
    assert reloaded.get('unfinished')['status'] == 'interrupted'


def test_model_payload_is_text_only_and_empty_or_incomplete_output_fails(tmp_path, monkeypatch):
    original_client = httpx.AsyncClient
    payloads = []
    def transport(request):
        payloads.append(json.loads(request.content))
        return httpx.Response(200, content=b'{"message":{"content":"Source summary"}}\n{"done":true,"eval_count":3}\n')
    monkeypatch.setattr(service.httpx, 'AsyncClient', lambda **kwargs: original_client(transport=httpx.MockTransport(transport), **kwargs))
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    assert asyncio.run(manager.call_model('model', 'Review this code', 900)) == 'Source summary'
    assert 'images' not in payloads[0] and all('images' not in message for message in payloads[0]['messages'])
    def empty(request): return httpx.Response(200, content=b'{"done":true}\n')
    monkeypatch.setattr(service.httpx, 'AsyncClient', lambda **kwargs: original_client(transport=httpx.MockTransport(empty), **kwargs))
    with pytest.raises(RuntimeError, match='empty or incomplete'):
        asyncio.run(manager.call_model('model', 'Review source', 900))


def test_authenticated_api_background_work_and_report_download(tmp_path, monkeypatch):
    root = tmp_path / 'source'; root.mkdir(); (root / 'guide.md').write_text('# Project\nThis project performs calculations.\n')
    async def summarize(prompt): return 'The project performs calculations, as recorded in guide.md.'
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3', summarize)
    monkeypatch.setattr(service, '_manager', manager)
    app = FastAPI(); app.include_router(router); app.add_middleware(SessionGuard)
    headers = {'X-LAW-Session': 'test-session-token'}
    with TestClient(app, base_url='http://127.0.0.1:8000') as client:
        assert client.post('/folder-review/reviews', json={'root': str(root), 'model': 'test-model'}).status_code == 403
        assert not manager.status()['reviews']
        assert client.post('/folder-review/reviews', headers=headers, json={'root': str(root), 'model': 'test-model', 'ocr': True}).status_code == 422
        response = client.post('/folder-review/reviews', headers=headers, json={'root': str(root), 'model': 'test-model'})
        assert response.status_code == 202
        key = response.json()['id']
        for _ in range(100):
            review = client.get('/folder-review/reviews/' + key, headers=headers).json()
            if review['status'] in service.TERMINAL: break
            time.sleep(0.01)
        assert review['status'] == 'completed' and review['report_ready']
        status = client.get('/folder-review/status', headers=headers).json()
        assert 'report' not in status['reviews'][0] and status['reviews'][0]['extensions'] == {'.md': 1}
        files = client.get(f'/folder-review/reviews/{key}/files', headers=headers).json()
        assert files['items'][0]['status'] == 'reviewed'
        report = client.get(f'/folder-review/reviews/{key}/export', headers=headers)
        assert report.status_code == 200 and 'guide.md' in report.text and 'State: completed' in report.text
        assert report.headers['cache-control'] == 'no-store'
        assert client.get('/folder-review/reviews/not-a-review', headers=headers).status_code == 422
        assert client.get('/folder-review/reviews/' + '0' * 32, headers=headers).status_code == 404


def test_directory_links_do_not_escape_scope(tmp_path):
    root = tmp_path / 'source'; root.mkdir()
    outside = tmp_path / 'outside'; outside.mkdir(); (outside / 'private.txt').write_text('Outside selected scope')
    try:
        (root / 'external').symlink_to(outside, target_is_directory=True)
    except OSError:
        pytest.skip('Directory symlinks are not available to this Windows test account')
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    async def execute():
        started = await manager.start(str(root), '', OPTIONS); await manager._task
        return manager.entries(started['id'])['items']
    entries = asyncio.run(execute())
    assert len(entries) == 1 and entries[0]['status'] == 'excluded'
    assert not entries[0]['analysis']


def test_folder_review_blocks_maintenance_even_between_inference_requests(tmp_path, monkeypatch):
    from services import maintenance_gate
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    monkeypatch.setattr(service, '_manager', manager)
    class Running:
        def done(self): return False
    manager._task = Running()
    assert maintenance_gate.workers_busy()


def review_provider(monkeypatch, *, ignore_release=False):
    """Run real admission, runtime preparation and streaming against a local fake provider."""
    from services import chat_model_runtime, context_awareness, image_generation
    from services.gpu_coordination import GpuCoordinator
    local_queue = RequestQueue(GpuCoordinator())
    monkeypatch.setattr(service, 'queue', local_queue)
    monkeypatch.setattr(chat_model_runtime, 'queue', local_queue)
    monkeypatch.setattr(context_awareness, '_limits', {})
    loaded, payloads, offloads, parks = ['other:latest'], [], [], []
    def park():
        assert local_queue.active is not None
        parks.append(local_queue.active.kind)
    monkeypatch.setattr(image_generation.manager, 'park_for_chat', park)
    def handle(request):
        assert local_queue.active is not None or request.url.path == '/api/show'
        if request.url.path == '/api/show':
            return httpx.Response(200, json={'model_info': {'fixture.context_length': 4096}})
        if request.url.path == '/api/ps':
            return httpx.Response(200, json={'models': [{'name': name} for name in loaded]})
        body = json.loads(request.content)
        if request.url.path == '/api/generate':
            name = chat_model_runtime.canonical_model(body['model'])
            if body.get('keep_alive') == 0:
                offloads.append(body)
                if name in loaded and not (ignore_release and name == 'fixture-review:latest'):
                    loaded.remove(name)
            else:
                loaded.append(name)
            return httpx.Response(200, json={'done': True})
        assert request.url.path == '/api/chat'
        payloads.append(body)
        return httpx.Response(200, content=b'{"message":{"content":"Observed fixture facts."}}\n{"done":true,"eval_count":4}\n')
    original = httpx.AsyncClient
    monkeypatch.setattr(service.httpx, 'AsyncClient', lambda **kwargs: original(transport=httpx.MockTransport(handle), **kwargs))
    return local_queue, loaded, payloads, offloads, parks


def test_full_pipeline_offloads_batches_and_compacts_with_model_context_budget(tmp_path, monkeypatch):
    local_queue, loaded, payloads, offloads, parks = review_provider(monkeypatch)
    root = tmp_path / 'source'; root.mkdir()
    source = ('Fact 漢字 😀 "\\\x01"\n' * 500) + 'FINAL SOURCE FACT'
    (root / 'facts.txt').write_text(source, encoding='utf-8')
    (root / 'second.md').write_text('# Fixture\nAdditional fixture facts.')
    before = {path.name: path.read_bytes() for path in root.iterdir()}
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    async def execute():
        review = await manager.start(str(root), 'fixture-review', {**OPTIONS, 'batch_chars': 6000})
        await manager._task
        return manager.get(review['id']), manager.entries(review['id'])['items']
    review, entries = asyncio.run(execute())
    assert review['status'] == 'completed'
    source_prompts = [payload['messages'][1]['content'] for payload in payloads
                      if payload['messages'][1]['content'].startswith('Review source file "facts.txt"')]
    assert len(source_prompts) > len(list(reader.text_batches(source, 6000)))
    assert ''.join(json.loads(prompt.split('\n', 1)[1])['source'] for prompt in source_prompts) == source
    for payload in payloads:
        assert payload['options']['num_ctx'] == 4096
        upper_bound = sum(len(message['content'].encode('utf-8')) for message in payload['messages'])
        assert upper_bound + payload['options']['num_predict'] + 192 <= 4096
        assert payload['messages'][0]['content'] == service.SYSTEM
    progress = review['processing']
    assert progress['context_limit'] == 4096
    assert progress['text_batches_completed'] == progress['text_batches_total'] > 2
    assert progress['compactions'] >= 2
    assert progress['offload_preparations'] == len(payloads) == len(parks)
    assert 'compact' in parks and 'compact' in {job.kind for job in local_queue.jobs}
    assert [call['model'] for call in offloads] == ['other:latest', 'fixture-review']
    assert progress['model_released'] and not loaded and local_queue.active is None
    assert all(item['status'] == 'reviewed' for item in entries)
    assert {path.name: path.read_bytes() for path in root.iterdir()} == before
    assert service.FolderReview(manager.database).get(review['id'])['processing'] == progress


def test_offload_verification_failure_is_recorded_without_losing_findings(tmp_path, monkeypatch):
    local_queue, _, _, _, _ = review_provider(monkeypatch, ignore_release=True)
    root = tmp_path / 'source'; root.mkdir(); (root / 'facts.txt').write_text('A neutral fixture fact.')
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    async def execute():
        review = await manager.start(str(root), 'fixture-review', OPTIONS)
        await manager._task
        return manager.get(review['id'])
    result = asyncio.run(execute())
    assert result['status'] == 'failed' and result['counts']['reviewed'] == 1
    assert 'Observed fixture facts.' in result['report']
    assert not result['processing']['model_released']
    assert 'still reports' in result['processing']['model_release_error']
    assert local_queue.active is None


def test_between_batch_cleanup_does_not_unload_another_running_job(tmp_path, monkeypatch):
    from services.gpu_coordination import GpuCoordinator
    local_queue = RequestQueue(GpuCoordinator())
    monkeypatch.setattr(service, 'queue', local_queue)
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    key = 'fixture'
    with manager.connect() as db:
        db.execute("INSERT INTO reviews(id,root,model,options,status,phase,started_at) VALUES(?,?,?,?,?,?,?)",
                   (key, str(tmp_path), 'fixture-review', '{}', 'running', 'fixture', service.now()))
    manager._processing = {'offload_preparations': 1, 'model_released': False}
    async def offload(key, model):
        assert local_queue.active.kind == 'analysis'
        manager.record_processing(key, model_released=True, model_release_deferred=False)
    monkeypatch.setattr(manager, 'release_model', offload)
    async def execute():
        other = local_queue.enqueue('chat', 'Other fixture job', model='other')
        assert local_queue.try_start(other)
        await manager.release_if_idle(key, 'fixture-review')
        assert local_queue.active is other
        assert manager.get(key)['processing']['model_release_deferred']
        assert not manager.get(key)['processing']['model_released']
        local_queue.finish(other)
        await manager.release_if_idle(key, 'fixture-review')
        assert manager.get(key)['processing']['model_released']
        assert local_queue.active is None
    asyncio.run(execute())


def test_too_small_context_rejects_instructions_without_silently_truncating(tmp_path):
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    manager._context_limit = 512
    with pytest.raises(RuntimeError, match='cannot fit'):
        manager.bounded_batches('Fixture source', 6000, lambda batch: manager.source_prompt('facts.txt', 'text', batch))


def test_model_offload_requires_review_queue_ownership(tmp_path):
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    with pytest.raises(ValueError, match='active queue admission'):
        asyncio.run(manager.release_model('fixture', 'fixture-review'))


def test_existing_saved_reports_gain_processing_column_without_data_loss(tmp_path):
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    with manager.connect() as db:
        db.execute("INSERT INTO reviews(id,root,model,options,status,phase,started_at,report) VALUES(?,?,?,?,?,?,?,?)",
                   ('saved', str(tmp_path), '', '{}', 'completed', 'Review complete', service.now(), 'Existing fixture report.'))
        db.execute('ALTER TABLE reviews DROP COLUMN processing')
        db.execute('ALTER TABLE reviews DROP COLUMN overview')
    restored = service.FolderReview(manager.database).get('saved')
    assert restored['report'] == 'Existing fixture report.' and restored['processing'] == {}
    assert restored['overview'] == ''


@pytest.mark.parametrize('failure_stage', ['source', 'compaction'])
def test_partial_batch_findings_survive_failure_and_restart(tmp_path, failure_stage):
    root = tmp_path / 'source'; root.mkdir()
    path = root / 'facts.txt'; path.write_text('A neutral fixture line.\n' * 50)
    before = path.read_bytes()
    calls = 0
    async def summarize(prompt):
        nonlocal calls
        calls += 1
        if (failure_stage == 'source' and calls == 2) or prompt.startswith('Combine'):
            raise RuntimeError('Fixture provider stopped')
        return f'Saved finding {calls}.'
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3', summarize)
    async def execute():
        started = await manager.start(str(root), 'fixture-model', OPTIONS)
        await asyncio.wait_for(manager._task, 5)
        return started['id']
    key = asyncio.run(execute())
    result = manager.get(key)
    item = manager.entries(key)['items'][0]
    assert result['status'] == 'failed' and result['counts']['not_reviewed'] == 1
    assert item['coverage']['partial'] and not item['coverage']['model_analysis']
    assert item['coverage']['batches_completed'] == (1 if failure_stage == 'source' else 3)
    assert 'Saved finding 1.' in result['report'] and 'Source lines 1-' in item['analysis']
    assert 'file review did not finish' in item['analysis']
    assert item['metadata']['sha256'] == reader.hashlib.sha256(before).hexdigest()
    # Simulate a process exit before it could commit terminal state/report.
    manager.update(key, status='running', report='')
    restored = service.FolderReview(manager.database)
    recovered = restored.get(key)
    assert recovered['status'] == 'interrupted' and recovered['report_ready']
    assert 'Saved finding 1.' in recovered['report'] and 'no inference was restarted' in recovered['report']
    assert path.read_bytes() == before and restored.status()['active'] is None


def test_cancellation_during_model_loading_joins_provider_before_offload(tmp_path, monkeypatch):
    local_queue, _, payloads, _, _ = review_provider(monkeypatch)
    root = tmp_path / 'source'; root.mkdir(); (root / 'facts.txt').write_text('Fixture fact.')
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    async def execute():
        entered, exited = asyncio.Event(), asyncio.Event()
        async def prepare(*args, **kwargs):
            entered.set()
            try:
                await asyncio.Future()
                yield {'detail': 'Never reached'}
            finally:
                exited.set()
        original = manager.release_model
        async def release(key, model):
            assert exited.is_set() and local_queue.active is manager._queue_job
            await original(key, model)
        monkeypatch.setattr(service, 'prepare_chat_model', prepare)
        monkeypatch.setattr(manager, 'release_model', release)
        started = await manager.start(str(root), 'fixture-review', OPTIONS)
        await asyncio.wait_for(entered.wait(), 5)
        await manager.cancel(started['id'])
        await asyncio.wait_for(manager._task, 5)
        assert exited.is_set() and local_queue.active is None and not payloads
        return manager.get(started['id'])
    result = asyncio.run(execute())
    assert result['status'] == 'cancelled' and result['processing']['model_released']


def test_task_cancellation_joins_inventory_worker_before_readmission(tmp_path, monkeypatch):
    root = tmp_path / 'source'; root.mkdir(); (root / 'facts.txt').write_text('Fixture fact.')
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    entered, release, exited = threading.Event(), threading.Event(), threading.Event()
    original = manager.inventory
    def inventory(*args):
        entered.set()
        try:
            assert release.wait(5)
            manager.check()
            original(*args)
        finally:
            exited.set()
    monkeypatch.setattr(manager, 'inventory', inventory)
    async def execute():
        started = await manager.start(str(root), '', OPTIONS)
        try:
            assert await asyncio.to_thread(entered.wait, 5)
            manager._task.cancel()
            await asyncio.sleep(0)
            assert manager.busy() and not exited.is_set()
            with pytest.raises(RuntimeError, match='already running'):
                await manager.start(str(root), '', OPTIONS)
        finally:
            release.set()
        await asyncio.wait_for(manager._task, 5)
        assert exited.is_set() and manager._active_id is None
        assert manager.get(started['id'])['status'] == 'cancelled'
        monkeypatch.setattr(manager, 'inventory', original)
        second = await manager.start(str(root), '', OPTIONS)
        await asyncio.wait_for(manager._task, 5)
        assert manager.get(second['id'])['status'] == 'completed'
    asyncio.run(execute())


@pytest.mark.parametrize('fault', ['compose', 'commit'])
def test_finalization_failure_releases_busy_state_and_recovers_report(tmp_path, monkeypatch, fault):
    root = tmp_path / 'source'; root.mkdir(); (root / 'facts.txt').write_text('Retained fixture fact.')
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    original_compose, original_update = manager.compose_report, manager.update
    if fault == 'compose':
        monkeypatch.setattr(manager, 'compose_report', lambda *args: (_ for _ in ()).throw(RuntimeError('Fixture report error')))
    else:
        def update(key, **values):
            if 'report' in values:
                raise service.sqlite3.OperationalError('Fixture report write failure')
            return original_update(key, **values)
        monkeypatch.setattr(manager, 'update', update)
    async def execute():
        started = await manager.start(str(root), '', OPTIONS)
        await asyncio.wait_for(manager._task, 5)
        return started['id']
    key = asyncio.run(execute())
    assert manager._active_id is None and not manager.busy() and manager._queue_job is None
    state = manager.get(key, False)
    assert state['status'] == 'failed' and state['report_ready'] and 'finalize' in state['error']
    monkeypatch.setattr(manager, 'compose_report', original_compose)
    monkeypatch.setattr(manager, 'update', original_update)
    assert 'Retained fixture fact.' in manager.get(key)['report']
    assert 'Retained fixture fact.' in service.FolderReview(manager.database).get(key)['report']


def test_bad_saved_json_does_not_block_saved_findings_or_status(tmp_path):
    root = tmp_path / 'source'; root.mkdir(); (root / 'facts.txt').write_text('Retained fixture fact.')
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    async def execute():
        started = await manager.start(str(root), '', OPTIONS)
        await manager._task
        return started['id']
    key = asyncio.run(execute())
    manager.update(key, options='not JSON', processing='{"context_limit":{},"model_release_error":[]}', status='running', report='')
    with manager.connect() as db:
        db.execute('UPDATE entries SET metadata=?,coverage=? WHERE review_id=?',
                   ('[]', '{"batches_total":{},"reason":{},"partial":[]}', key))
    restored = service.FolderReview(manager.database)
    result, entry = restored.get(key), restored.entries(key)['items'][0]
    assert result['status'] == 'interrupted' and result['options'] == {} and result['data_warnings']
    assert 'context_limit' not in result['processing']
    assert entry['metadata'] == {} and entry['data_warnings'] and 'batches_total' not in entry['coverage']
    assert 'Retained fixture fact.' in result['report'] and 'Integrity notice:' in result['report']
    assert restored.status()['reviews'][0]['id'] == key


def test_changed_inventory_snapshot_is_not_reviewed(tmp_path, monkeypatch):
    root = tmp_path / 'source'; root.mkdir(); path = root / 'facts.txt'; path.write_text('Initial fact.')
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    original = manager.inventory
    def inventory(*args):
        original(*args)
        path.write_text('Changed fixture facts with a new length.')
    monkeypatch.setattr(manager, 'inventory', inventory)
    async def execute():
        started = await manager.start(str(root), '', OPTIONS)
        await manager._task
        return manager.get(started['id']), manager.entries(started['id'])['items'][0]
    result, entry = asyncio.run(execute())
    assert result['status'] == 'completed_with_gaps' and entry['status'] == 'unreadable'
    assert 'changed since the folder inventory' in entry['error'] and not entry['analysis']
    assert path.read_text() == 'Changed fixture facts with a new length.'


@pytest.mark.parametrize('temporary', [True, False])
def test_provider_transport_retry_is_bounded_and_preserves_request(tmp_path, monkeypatch, temporary):
    original_client, payloads = httpx.AsyncClient, []
    def handle(request):
        payloads.append(json.loads(request.content))
        if len(payloads) == 1:
            if temporary:
                return httpx.Response(503, json={'error': 'Fixture unavailable'})
            return httpx.Response(400, json={'error': 'Invalid fixture request'})
        return httpx.Response(200, content=b'{"message":{"content":"Complete fixture summary"},"done":true}\n')
    monkeypatch.setattr(service.httpx, 'AsyncClient', lambda **kwargs: original_client(transport=httpx.MockTransport(handle), **kwargs))
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    if temporary:
        assert asyncio.run(manager.call_model('fixture', 'Entire source', 900)) == 'Complete fixture summary'
        assert len(payloads) == 2 and payloads[0] == payloads[1]
    else:
        with pytest.raises(httpx.HTTPStatusError):
            asyncio.run(manager.call_model('fixture', 'Entire source', 900))
        assert len(payloads) == 1


@pytest.mark.parametrize('body,match', [
    (b'{broken}\n', 'invalid JSON'),
    (b'[]\n', 'non-object'),
    (b'{"message":{"content":[]},"done":true}\n', 'invalid message'),
    (b'{"message":{"content":"partial"}}\n{"message":[],"done":true}\n', 'invalid message'),
    (b'{"message":{"content":"text"},"done":"true"}\n', 'completion flag'),
    (b'{"message":{"content":"partial"}}\n', 'incomplete'),
    (b'{"message":{"content":"limit"},"done":true,"done_reason":"length"}\n', 'output limit'),
    (b'{"message":{"content":"text"},"done":true,"eval_count":[]}\n', 'usage counts'),
    (b'{"message":{"content":"' + b'x' * 32001 + b'"},"done":true}\n', 'response limit'),
    (b'x' * (256 * 1024 + 1), 'line exceeded'),
], ids=['bad-json', 'non-object', 'bad-content', 'bad-message', 'bad-done', 'incomplete', 'output-limit', 'bad-count', 'long-summary', 'long-line'])
def test_malformed_or_oversized_provider_output_is_not_accepted(tmp_path, monkeypatch, body, match):
    original_client, calls = httpx.AsyncClient, []
    def handle(request):
        calls.append(request.url.path)
        return httpx.Response(200, content=body)
    monkeypatch.setattr(service.httpx, 'AsyncClient', lambda **kwargs: original_client(transport=httpx.MockTransport(handle), **kwargs))
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    with pytest.raises(RuntimeError, match=match):
        asyncio.run(manager.call_model('fixture', 'Entire source', 900))
    assert calls == ['/api/chat']


def test_docx_expansion_is_bounded_before_parser_and_text_cap_reports_coverage(tmp_path, monkeypatch):
    import docx
    import io
    import zipfile
    package = io.BytesIO()
    with zipfile.ZipFile(package, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('word/document.xml', 'x' * 10000)
    with monkeypatch.context() as patch:
        patch.setattr(reader, 'DOCX_MAX_EXPANDED_BYTES', 1000)
        patch.setattr(docx, 'Document', lambda *args: pytest.fail('Oversized package reached XML parser'))
        with pytest.raises(ValueError, match='expanded size'):
            reader.extract(package.getvalue(), 'fixture.docx', 1000)
    document = docx.Document(); document.add_paragraph('First fixture paragraph.')
    document.add_paragraph('Second fixture paragraph.')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Left'; table.cell(0, 1).text = 'Right'
    package = io.BytesIO(); document.save(package)
    full = reader.extract(package.getvalue(), 'fixture.docx', 1000)
    partial = reader.extract(package.getvalue(), 'fixture.docx', 10)
    assert full['text'] == 'First fixture paragraph.\nSecond fixture paragraph.\nLeft | Right'
    assert partial['text'] == full['text'][:10] and partial['coverage']['partial']
    assert partial['coverage']['original_characters'] == len(full['text'])


def test_invalid_loaded_model_list_cannot_claim_verified_offload(tmp_path, monkeypatch):
    original_client = httpx.AsyncClient
    local_queue, _, _, _, _ = review_provider(monkeypatch)
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    monkeypatch.setattr(service.httpx, 'AsyncClient', lambda **kwargs: original_client(transport=httpx.MockTransport(
        lambda request: httpx.Response(200, json={})), **kwargs))
    async def execute():
        job = local_queue.enqueue('analysis', 'Fixture release')
        assert local_queue.try_start(job)
        manager._queue_job = job
        try:
            with pytest.raises(RuntimeError, match='loaded-model list'):
                await manager.release_model('fixture', 'fixture-review')
            assert not manager._processing.get('model_released')
        finally:
            local_queue.finish(job)
    asyncio.run(execute())


def test_task_cancellation_joins_runtime_offloading_before_releasing_queue(tmp_path, monkeypatch):
    local_queue, _, payloads, _, _ = review_provider(monkeypatch)
    root = tmp_path / 'source'; root.mkdir(); (root / 'facts.txt').write_text('Fixture fact.')
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    entered, release, exited = threading.Event(), threading.Event(), threading.Event()
    def park():
        entered.set()
        try:
            assert release.wait(5)
        finally:
            exited.set()
    async def prepare(kind):
        await asyncio.to_thread(park)
    monkeypatch.setattr(service, 'prepare_runtime', prepare)
    async def execute():
        started = await manager.start(str(root), 'fixture-review', OPTIONS)
        try:
            assert await asyncio.to_thread(entered.wait, 5)
            manager._task.cancel()
            await asyncio.sleep(0)
            assert manager.busy() and not exited.is_set() and local_queue.active is manager._queue_job
        finally:
            release.set()
        await asyncio.wait_for(manager._task, 5)
        assert exited.is_set() and local_queue.active is None and not payloads
        assert manager.get(started['id'])['status'] == 'cancelled'
    asyncio.run(execute())


def test_simultaneous_first_requests_share_one_review_owner(tmp_path, monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    from dataclasses import replace
    monkeypatch.setattr(service, '_manager', None)
    monkeypatch.setattr(service, 'settings', replace(service.settings, data_dir=tmp_path))
    created, barrier = [], threading.Barrier(4)
    def construct(database):
        time.sleep(0.02)
        created.append(object())
        return created[-1]
    monkeypatch.setattr(service, 'FolderReview', construct)
    def read():
        barrier.wait(timeout=5)
        return service.manager()
    with ThreadPoolExecutor(max_workers=4) as workers:
        owners = list(workers.map(lambda _: read(), range(4)))
    assert len(created) == 1 and all(owner is created[0] for owner in owners)


def test_failed_admission_does_not_create_busy_or_orphaned_review(tmp_path):
    root = tmp_path / 'source'; root.mkdir(); (root / 'facts.txt').write_text('Fixture fact.')
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    with manager.connect() as db:
        db.execute("CREATE TRIGGER block_start BEFORE INSERT ON reviews BEGIN SELECT RAISE(ABORT, 'Fixture admission failure'); END")
    async def execute():
        with pytest.raises(service.sqlite3.IntegrityError, match='admission failure'):
            await manager.start(str(root), '', OPTIONS)
        assert manager._active_id is None and not manager.busy() and manager.status()['reviews'] == []
        with manager.connect() as db:
            db.execute('DROP TRIGGER block_start')
        started = await manager.start(str(root), '', OPTIONS)
        await manager._task
        assert manager.get(started['id'])['status'] == 'completed'
    asyncio.run(execute())


def test_transport_retry_stops_after_two_attempts_and_cancel_prevents_retry(tmp_path, monkeypatch):
    original_client, calls = httpx.AsyncClient, []
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    def handle(request):
        calls.append(request)
        raise httpx.ReadError('Fixture connection closed', request=request)
    monkeypatch.setattr(service.httpx, 'AsyncClient', lambda **kwargs: original_client(transport=httpx.MockTransport(handle), **kwargs))
    with pytest.raises(httpx.ReadError):
        asyncio.run(manager.call_model('fixture', 'Entire source', 900))
    assert len(calls) == 2
    calls.clear()
    async def cancelled(model, prompt, tokens):
        manager._cancel.set()
        raise httpx.ReadError('Fixture connection closed')
    monkeypatch.setattr(manager, 'stream_model', cancelled)
    with pytest.raises(service.QueueCancelled):
        asyncio.run(manager.call_model('fixture', 'Entire source', 900))
    assert not calls


def test_model_stream_total_byte_limit_is_enforced(tmp_path):
    manager = service.FolderReview(tmp_path / 'reports/db.sqlite3')
    # Small individual lines still cannot exceed the whole-response cap.
    line = b'{"padding":"' + b'x' * 2000 + b'"}\n'
    response = httpx.Response(200, content=line * (4 * 1024 ** 2 // len(line) + 1))
    async def read():
        async for _ in manager.response_lines(response):
            pass
    with pytest.raises(RuntimeError, match='stream exceeded its byte limit'):
        asyncio.run(read())
