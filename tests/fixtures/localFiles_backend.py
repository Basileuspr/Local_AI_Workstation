"""Disposable native QA server. All source fixtures are created under the supplied temp folder."""
import sys
import os
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[2]/'backend'))
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'backend'))
import sqlite3
from contextlib import closing
from docx import Document
from fastapi import FastAPI
from fastapi.middleware.cors import CORSMiddleware
from routes.local_files import router
from services.session_guard import SessionGuard
from test_local_files import make_video
import uvicorn

root=Path(sys.argv[1]);port=int(sys.argv[2])
doc=Document();doc.add_heading('Local file editing',1);doc.add_paragraph('Original paragraph');doc.add_paragraph('Keep this paragraph')
table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='Item';table.cell(0,1).text='Quantity';table.cell(1,0).text='Tray';table.cell(1,1).text='6';doc.save(root/'sample.docx')
with closing(sqlite3.connect(root/'sample.db')) as db:
    db.executescript('CREATE TABLE inventory(id INTEGER PRIMARY KEY,name TEXT,quantity INTEGER); CREATE VIEW available AS SELECT * FROM inventory WHERE quantity>0;')
    db.executemany('INSERT INTO inventory VALUES(?,?,?)',[(i,f'Item {i}',i%10) for i in range(250)]);db.commit()
make_video(root/'sample.mp4',frames=30,audio=True,width=640,height=360)
if os.environ.get('LAW_QA_VIDEO_DESCRIPTIONS') == '1':
    from services import local_video
    from services.image_workflows import adapters
    adapters.vision_models=lambda:[{'id':'fixture-vision','name':'Controlled QA vision responses'}]
    async def controlled_analysis(frames,directory,model,limit,cancel,report,focus,transcript,analysis):
        observations=[{'id':frame['id'],'time':frame['time'],'text':'A blue field fills the frame. Its brightness differs across the sampled moments.'} for frame in frames[:limit]]
        analysis.update(status='complete',model=model,observations=observations,summary='Across the sampled moments, a blue field becomes brighter. These controlled QA responses verify presentation, not model accuracy.',note='Controlled fixture output.')
        return analysis
    local_video.vision=controlled_analysis
app=FastAPI();app.include_router(router)
app.add_middleware(SessionGuard)
app.add_middleware(CORSMiddleware,allow_origins=['app://local'],allow_methods=['*'],allow_headers=['*'])
uvicorn.run(app,host='127.0.0.1',port=port,log_level='error')
